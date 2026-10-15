# -*- coding: utf-8 -*-
"""Thesis-defense decks (10-15): thesis final v84_ko, told in the story order of the reference talk the user shared
(5 chapters: introduction / theory / research questions and methods / results / conclusion; a question opening the
talk, the research question stated after its rationale, summary slides, one take-away line per slide).
Two versions (user, 10-15): 20-minute (29 slides) and 15-minute (19 slides), each followed by backup slides.
All numbers are taken from manuscript_v84_ko (audit_numbers.py). Speaker script in the notes.
Usage: python3 build_defense.py 20|15 OUT.pptx"""
import sys, json, math
import kit as K
from kit import *
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION

VER = sys.argv[1]
OUT = sys.argv[2]
K.setup('assets', 'icons')
D = json.load(open('assets/data_v84.json'))

TITLE_KO = ('반도체 하이브리드 본딩의', '공정 단계 간 기술융합', ': 특허 기반 종단 분석')
TITLE_EN = 'Cross-Stage Technology Convergence in Semiconductor Hybrid Bonding: A Patent-Based Longitudinal Analysis'
RQ = {1: '하이브리드 본딩 기술에서 제조공정과 접합 기술 간의 융합이 어느 정도로, 어떠한 형태로 나타나는가?',
      2: '공정 단계 간 기술융합은 시간에 따라 어떻게 변화해 왔는가?',
      3: '공정 단계 간 기술융합의 중심이 핵심 기술과 주변 기술 사이에서 시기에 따라 어떻게 재편되는가?'}
RQ_SHORT = {1: '융합의 수준과 유형', 2: '융합의 시간적 변화', 3: '융합의 구조적 중심 변화'}
CH = {1: ('서론', 'Introduction'), 2: ('이론적 배경', 'Theoretical Background'),
      3: ('연구 질문 및 연구 방법', 'Research Questions & Methods'), 4: ('연구 결과', 'Results'), 5: ('결론', 'Conclusion')}


SCRIPT = None      # script override for the slide being built (15-minute version)
APPX = None        # when set, the slide is built as backup slide number APPX


def CS(num, title, section, script=None, take=None, **kw):
    """content slide; honours the version-specific script and the backup-slide numbering"""
    if APPX is not None:
        num, section = f'부록 {APPX}', '부록'
        if title == '서론':
            title = '반도체 분업 구조와 첨단 패키징'
    return content_slide(num, title, section, script=SCRIPT if SCRIPT is not None else script, take=take, **kw)


def run(fn, script=None, appx=None, **kw):
    global SCRIPT, APPX
    SCRIPT, APPX = script, appx
    fn(**kw)
    SCRIPT = APPX = None


def rq_band(n):
    """research question stated at the foot of its rationale slide (replaces the take-away line)"""
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, X0, 15.85, CW, 2.05, fill=NAVY, radius=0.18, name='Research question')
    chip(X0 + 0.45, 16.43, 2.4, 0.9, f'RQ{n}', fill='D4A72C', color=NAVY, size=14)
    text(X0 + 3.2, 15.85, CW - 3.6, 2.05, [P((RQ[n], 16, WHITE, True), line=1.1)], anchor='m')


def dark_overlay(alpha=0.55, x=0, y=0, w=SW, h=SH, col='05070B'):
    return shape(MSO_SHAPE.RECTANGLE, x, y, w, h, fill=col, alpha=alpha)


def die_stack(x, y, w, hybrid, top='8A97AD', bot='5E6C85', bump='A9AFBA', oxide='BFE3E6', die_h=0.9):
    """one die on another: solder bumps or hybrid-bonded Cu pads in oxide (as on the poster)"""
    gap = 0.1 if hybrid else 0.8
    shape(MSO_SHAPE.RECTANGLE, x, y, w, die_h, fill=top)
    by = y + die_h + gap
    if hybrid:
        oy = y + die_h - 0.38
        shape(MSO_SHAPE.RECTANGLE, x, oy, w, 0.38, fill=oxide)
        shape(MSO_SHAPE.RECTANGLE, x, by, w, 0.38, fill=oxide)
        n = 6
        for k in range(n):
            px = x + 0.45 + k * (w - 0.9 - 0.7) / (n - 1)
            shape(MSO_SHAPE.RECTANGLE, px, oy, 0.7, 0.38, fill=BOND)
            shape(MSO_SHAPE.RECTANGLE, px, by, 0.7, 0.38, fill=BOND)
        shape(MSO_SHAPE.RECTANGLE, x, by - gap, w, gap, fill=INTEG)
        by += 0.38
    else:
        n = 5
        for k in range(n):
            px = x + 0.55 + k * (w - 1.1 - 0.85) / (n - 1)
            shape(MSO_SHAPE.OVAL, px, y + die_h - 0.02, 0.85, gap + 0.04, fill=bump, line='7D8594', lw=0.75)
    shape(MSO_SHAPE.RECTANGLE, x, by, w, die_h, fill=bot)
    return by + die_h


def rq_panel(x, y, w, n, answer_runs, h=None, qh=1.6):
    """right-hand panel of a result slide (reference: the hypothesis beside the table)"""
    hh = h or 6.0
    card(x, y, w, hh, fill=PANEL)
    text(x + 0.5, y + 0.35, w - 1.0, 0.8, [P((f'Research Question {n}', 18 if w > 9 else 15, NAVY, True))], anchor='m')
    text(x + 0.5, y + 1.2, w - 1.0, qh, [P((RQ[n], 12.5, MUTED), line=1.05)])
    line(x + 0.5, y + 1.35 + qh, x + w - 0.5, y + 1.35 + qh, RULE, 1.0)
    text(x + 0.5, y + 1.55 + qh, w - 1.0, hh - 1.8 - qh, answer_runs)


def cut(t, n=44):
    return t


# ================================================================ 0 cover, contents
def s_cover():
    image_slide('wafer.png', '표지', '표지', script=SCRIPT or (
        '안녕하십니까. 기술경영학과 이형규입니다. 지금부터 "반도체 하이브리드 본딩의 공정 단계 간 기술융합: 특허 기반 종단 분석"이라는 '
        '주제로 석사학위논문 심사 발표를 진행하겠습니다. 잘 전달되었으면 좋겠습니다.'))
    shape(MSO_SHAPE.RECTANGLE, 13.2, 0, SW - 13.2, SH, fill='060810', alpha=0.80)
    chip(15.0, 2.2, 5.6, 0.95, '석사학위논문 심사 발표', fill=None, color=GOLDB, size=13, line_=GOLDB)
    text(15.0, 3.7, 17.6, 5.4, [P((TITLE_KO[0], 32, WHITE, True), line=1.05),
                               P((TITLE_KO[1], 32, 'C5B5FF', True), line=1.05),
                               P((TITLE_KO[2], 32, WHITE, True), line=1.05)])
    text(15.0, 9.25, 17.4, 1.6, [P((TITLE_EN, 13, DIM, False, True), line=1.1)])
    line(15.0, 11.4, 31.8, 11.4, '4A5670', 1.0)
    text(15.0, 11.8, 17.4, 4.6, [P(('이형규', 22, WHITE, True), after=6),
                                P(('한양대학교 기술경영전문대학원 기술경영학과', 14, 'DCE3EE'), after=3),
                                P(('지도교수  강민구', 14, 'DCE3EE'), after=3),
                                P(('2026년 11월', 14, 'DCE3EE'))])
    picture('assets/emblem.png', SW - 2.9, SH - 2.9, 1.6, 1.6)


def s_contents():
    CS('', 'Contents', '목차', script=(
        '발표는 다섯 부분으로 구성되어 있습니다. 왜 이 연구를 시작했는지부터, 어떤 이론을 바탕으로 연구 질문을 세웠고 '
        '어떤 자료와 방법으로 분석했는지, 그리고 최종적으로 어떤 결과가 나왔으며 그것이 무엇을 의미하는지를 담았습니다.'))
    ic = ['FaLightbulb', 'FaBookOpen', 'FaMagnifyingGlass', 'FaChartLine', 'FaBullseye']
    w, g = 5.6, 0.72
    x0 = (SW - 5 * w - 4 * g) / 2
    names = {3: '연구 질문 및\n연구 방법'}
    for k in range(5):
        x = x0 + k * (w + g)
        card(x, 3.9, w, 11.4, fill=PANEL)
        text(x, 4.6, w, 1.6, [P((f'0{k + 1}', 36, NAVY, True), align=C)], anchor='m')
        icon_circle(ic[k], x + w / 2 - 1.2, 6.9, 2.4, NAVY)
        nm = names.get(k + 1, CH[k + 1][0])
        text(x + 0.3, 10.0, w - 0.6, 2.4, [P((ln_, 17, NAVY, True), align=C, line=1.05) for ln_ in nm.split('\n')], anchor='m')
        text(x + 0.3, 12.6, w - 0.6, 1.6, [P((CH[k + 1][1], 11.5, MUTED), align=C)], anchor='t')


def divider(n, script):
    image_slide('wafer.png', CH[n][0], f'0{n} {CH[n][0]}', script=script)
    dark_overlay(0.62)
    shape(MSO_SHAPE.RECTANGLE, 7.9, 5.2, SW - 15.8, 8.6, line=WHITE, lw=2.0)
    text(7.9, 5.9, SW - 15.8, 1.6, [P((f'0{n}', 30, GOLDB, True), align=C)], anchor='m')
    text(7.9, 7.6, SW - 15.8, 2.6, [P((CH[n][0], 44, WHITE, True), align=C)], anchor='m')
    text(7.9, 10.4, SW - 15.8, 1.2, [P((CH[n][1], 18, DIM), align=C)], anchor='m')


def rq_card(n, script):
    image_slide('circuit.png', CH[3][0], f'Research Question {n}', script=script)
    dark_overlay(0.35)
    text(2.0, 3.2, SW - 4.0, 2.2, [P((f'Research Question {n}', 44, WHITE, True), align=C)], anchor='m')
    text(2.0, 5.5, SW - 4.0, 1.0, [P((RQ_SHORT[n], 18, 'D4A72C', True), align=C)], anchor='m')
    dark_overlay(0.78, 0, 8.3, SW, 6.2)
    text(3.0, 8.3, SW - 6.0, 6.2, [P((RQ[n], 28, WHITE, True), align=C, line=1.15)], anchor='m')


# ================================================================ 01 서론
def s_intro_hook():
    CS('1', '서론', CH[1][0], take='더 높이, 더 촘촘히 쌓기 위한 연결 기술: 하이브리드 본딩', script=(
        '생성형 AI, AI 가속기, 그리고 HBM. 요즘 반도체 이야기에서 빠지지 않는 이 세 단어를 이어 주는 공통의 과제는 무엇일까요? '
        '생성형 AI는 막대한 데이터를 처리하는 AI 가속기를 필요로 하고, AI 가속기는 D램 다이를 여러 층 쌓은 HBM을 필요로 합니다. '
        '결국 메모리를 더 높이, 더 촘촘하게 쌓아 연결하는 것이 관건입니다. 그 연결을 위한 차세대 접합 기술로 주목받는 것이 '
        '오늘 말씀드릴 하이브리드 본딩입니다.'))
    items = [('FaRobot', '생성형 AI', 'Generative AI'), ('FaMicrochip', 'AI 가속기', 'AI Accelerator'), ('FaMemory', 'HBM', 'High Bandwidth Memory')]
    for k, (ic, ko, en) in enumerate(items):
        y = 2.7 + k * 4.65
        card(X0 + 1.0, y, 10.4, 4.1, fill=PANEL)
        icon_circle(ic, X0 + 1.6, y + 0.75, 2.6, NAVY)
        text(X0 + 4.7, y + 0.9, 6.5, 1.3, [P((ko, 22, NAVY, True))], anchor='m')
        text(X0 + 4.7, y + 2.15, 6.5, 0.9, [P((en, 12, MUTED))], anchor='m')
    shape(MSO_SHAPE.RIGHT_ARROW, 13.6, 7.6, 2.6, 3.2, fill='C9D2E0')
    # HBM stack, hybrid-bonded interfaces
    sx, sw_ = 18.6, 8.6
    card(17.0, 2.7, 15.4, 13.95, fill=PANEL)
    yy = 3.6
    for k in range(4):
        shape(MSO_SHAPE.RECTANGLE, sx, yy, sw_, 1.0, fill=('8FA6CC', '7F98C2', '6F8AB7', '607DAD')[k])
        if k == 0:
            text(sx, yy, sw_, 1.0, [P(('DRAM 다이', 13, NAVY, True), align=C)], anchor='m')
        yy += 1.0
        shape(MSO_SHAPE.RECTANGLE, sx, yy, sw_, 0.34, fill='BFE3E6')
        for m in range(8):
            shape(MSO_SHAPE.RECTANGLE, sx + 0.35 + m * (sw_ - 0.7 - 0.5) / 7, yy, 0.5, 0.34, fill=BOND)
        if k == 1: iy = yy + 0.17
        yy += 0.34
    shape(MSO_SHAPE.RECTANGLE, sx - 0.5, yy, sw_ + 1.0, 1.2, fill='3E5A8A')
    text(sx - 0.5, yy, sw_ + 1.0, 1.2, [P(('베이스(로직) 다이', 13, WHITE, True), align=C)], anchor='m')
    lx = sx + sw_ + 0.4
    line(sx + sw_ + 0.05, iy, lx + 0.4, iy, NAVY, 1.25)
    text(lx + 0.5, iy - 0.75, 3.4, 2.4, [P(('하이브리드', 13, INTEG, True)), P(('본딩 계면', 13, INTEG, True), after=3),
                                        P(('Cu 패드 +', 11, MUTED)), P(('산화막 직접 접합', 11, MUTED))])
    text(17.0, 11.6, 15.4, 1.9, [P(('Hybrid Bonding', 30, NAVY, True), align=C)], anchor='m')
    text(17.0, 13.5, 15.4, 1.0, [P(('차세대 HBM 적층의 접합 기술 (개념도)', 12, MUTED), align=C)], anchor='m')


def s_intro_division():
    CS('1', '서론', CH[1][0], take='미세화의 한계 이후, 전공정과 후공정 사이의 인터페이스가 변하고 있다', script=(
        '반도체 제조는 크게 두 단계로 나뉩니다. 웨이퍼 위에 트랜지스터를 만드는 전공정, 그리고 완성된 웨이퍼를 잘라 조립하고 패키징하는 후공정입니다. '
        '전공정은 종합반도체기업과 파운드리의 팹에서, 후공정은 조립·테스트 전문업체에서 맡는 분업이 반세기 넘게 유지되어 왔습니다. '
        '이 분업이 가능했던 것은 두 단계 사이의 인터페이스가 안정적이었기 때문입니다. 그런데 회로 선폭을 줄이는 미세화가 한계에 부딪히자, '
        '산업은 다이를 나누고 쌓아 연결하는 첨단 패키징에서 새로운 성능 향상의 기회를 찾기 시작했고, 한 단계의 성능이 다른 단계의 공정 역량에 '
        '더 직접적으로 의존하게 되었습니다.'))
    bw, bh, by = 11.6, 5.6, 2.7
    for k, (ko, en, desc, who, col, colL) in enumerate([
            ('전공정', 'Front End', '웨이퍼 위 트랜지스터 형성', '종합반도체기업(IDM) · 파운드리', PROC, PROC_L),
            ('후공정', 'Back End', '절단 · 조립 · 패키징 · 테스트', '조립·테스트 전문업체(OSAT)', BOND, BOND_L)]):
        x = X0 + 1.0 if k == 0 else X1 - 1.0 - bw
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, by, bw, bh, fill=colL, radius=0.08)
        text(x + 0.6, by + 0.45, bw - 1.2, 1.2, [P((ko, 24, col, True), (f'  {en}', 14, col))], anchor='m')
        text(x + 0.6, by + 1.85, bw - 1.2, 1.0, [P((desc, 16, TEXT, True))], anchor='m')
        chip(x + 0.6, by + 3.4, bw - 1.2, 1.2, who, fill=WHITE, color=col, size=14)
    cx = SW / 2
    line(cx, by + 0.2, cx, by + bh / 2 - 0.7, NAVY, 2.0, dash='dash')
    line(cx, by + bh / 2 + 1.55, cx, by + bh - 0.2, NAVY, 2.0, dash='dash')
    chip(cx - 2.6, by + bh / 2 - 0.55, 5.2, 1.1, '인터페이스', fill=NAVY, color=WHITE, size=14)
    text(cx - 3.2, by + bh / 2 + 0.65, 6.4, 0.8, [P(('반세기 이상의 분업', 12, MUTED), align=C)], anchor='m')
    # bottom flow
    fy, fh = 9.6, 4.9
    steps = [('미세화의 한계', '집적도 향상 둔화 · 비용 급증', 'FaHourglassHalf'),
             ('첨단 패키징', '다이 분할 · 적층 · 연결', 'FaLayerGroup'),
             ('공정 단계 간 상호의존', '한 단계의 성능이\n다른 단계의 공정 역량에 의존', 'FaLink')]
    sw2, g = 8.9, 2.0
    x0 = (SW - 3 * sw2 - 2 * g) / 2
    for k, (t1, t2, ic) in enumerate(steps):
        x = x0 + k * (sw2 + g)
        card(x, fy, sw2, fh, fill=WHITE, line_=NAVY if k == 2 else 'C9D2E0')
        icon(ic, x + sw2 / 2 - 0.75, fy + 0.5, 1.5, NAVY)
        text(x + 0.3, fy + 2.2, sw2 - 0.6, 1.0, [P((t1, 18, NAVY, True), align=C)], anchor='m')
        text(x + 0.3, fy + 3.25, sw2 - 0.6, 1.9, [P((ln_, 13, MUTED), align=C) for ln_ in t2.split('\n')], anchor='t')
        if k < 2:
            arrow(x + sw2 + 0.25, fy + fh / 2, x + sw2 + g - 0.25, fy + fh / 2, NAVY, 2.5)
    source('Kapoor (2013); Macher & Mowery (2004); Arden et al. (2010); Khan et al. (2018)', y=15.0, align=R)


def s_intro_hb():
    CS('1', '서론', CH[1][0], take='후공정의 접합 기술이지만, 그 성패는 전공정의 제조공정 역량이 좌우한다', script=(
        '하이브리드 본딩은 솔더 범프 없이 구리 패드와 주변 산화막을 직접 맞붙이는 기술입니다. 기존의 범프 접합이 두 판을 납이라는 접착제로 붙이는 '
        '방식이라면, 하이브리드 본딩은 두 표면을 거울처럼 평평하게 갈아 접착제 없이 그대로 맞붙이는 방식에 비유할 수 있습니다. 덕분에 기존 방식으로는 '
        '어려운 10마이크로미터 이하의 연결 간격을 구현합니다. 그런데 이 기술의 성패를 좌우하는 나노미터 수준의 평탄화, 표면 활성화, 파티클 제어, '
        '정밀 정렬은 후공정이 쌓아 온 역량이 아니라 전공정의 제조공정 역량입니다. 후공정 기술의 성패를 전공정 역량이 좌우하게 된 것입니다.'))
    cw2, ch_, y = 8.2, 9.6, 2.75
    for k, hyb in enumerate((False, True)):
        x = X0 + k * (cw2 + 1.6)
        card(x, y, cw2, ch_, fill=PANEL)
        text(x + 0.3, y + 0.3, cw2 - 0.6, 0.9, [P((('하이브리드 본딩' if hyb else '기존: 솔더 범프 접합'), 17, (INTEG if hyb else MUTED), True), align=C)], anchor='m')
        die_stack(x + 0.9, y + 1.55, cw2 - 1.8, hyb)
        if hyb:
            items = [('BFE3E6', '산화막'), (BOND, '구리 패드'), (INTEG, '접합 계면')]
        else:
            items = [('A9AFBA', '솔더 범프'), ('8A97AD', '다이')]
        lx = x + 0.6
        for col, lab in items:
            shape(MSO_SHAPE.RECTANGLE, lx, y + 5.0, 0.42, 0.42, fill=col)
            text(lx + 0.55, y + 4.86, 2.0, 0.7, [P((lab, 11, MUTED))], anchor='m')
            lx += 2.4
        text(x + 0.4, y + 5.9, cw2 - 0.8, 3.1, [
            P((('구리–구리 + 산화막–산화막' if hyb else '다이 사이를 솔더 범프로 연결'), 14, TEXT, True), align=C, after=4),
            P((('직접 접합' if hyb else '연결 간격을 좁히는 데 한계'), 14, TEXT, True), align=C, after=6),
            P((('10 µm 이하 연결 간격' if hyb else '비유: 접착제(솔더)로 붙이기'), 13, (INTEG if hyb else MUTED), hyb), align=C, after=2),
            P((('비유: 거울처럼 평평한 두 면을 그대로 맞붙이기' if hyb else ''), 11, MUTED), align=C)])
    arrow(X0 + cw2 + 0.2, y + 4.0, X0 + cw2 + 1.4, y + 4.0, INTEG, 3)
    # right: the front-end processes that decide the yield
    rx = X0 + 2 * cw2 + 1.6 + 1.6
    rw = X1 - rx
    text(rx, y, rw, 0.9, [P(('성패를 좌우하는 공정', 17, NAVY, True))], anchor='m')
    procs = [('CMP 평탄화', '나노미터 이하의 표면 거칠기'), ('표면 활성화', '플라즈마 처리'), ('파티클 제어', '세정 · 오염 관리'), ('정밀 정렬', '패드 간 위치 정합')]
    bw_ = rw - 3.4
    for k, (t1, t2) in enumerate(procs):
        yy = y + 1.25 + k * 2.05
        obox(rx, yy, bw_, 1.7, t1, t2, size=15, sub_size=11, color=PROC, fill=PROC_L)
    bx = rx + bw_ + 0.35
    line(bx, y + 1.35, bx + 0.45, y + 1.35, PROC, 2); line(bx + 0.45, y + 1.35, bx + 0.45, y + 9.35, PROC, 2); line(bx, y + 9.35, bx + 0.45, y + 9.35, PROC, 2)
    text(bx + 0.6, y + 3.3, 2.3, 4.2, [P(('전공정의', 13, PROC, True), align=C), P(('제조공정', 13, PROC, True), align=C),
                                      P(('역량', 13, PROC, True), align=C, after=4), P(('CPC', 11, MUTED), align=C), P(('H01L21', 11, MUTED), align=C)], anchor='m')
    text(X0, 12.85, CW, 0.9, [P(('하이브리드 본딩: 솔더 범프 없이 구리 패드와 주변 산화막을 직접 접합하여 10 µm 이하의 연결 간격을 구현 (연구 단계에서는 1 µm 이하)', 13, TEXT))], anchor='m')
    text(X0, 13.75, CW, 0.9, [P(('W2W(웨이퍼-투-웨이퍼): CMOS 이미지센서·3D NAND에서 먼저 상용화  ·  D2W(다이-투-웨이퍼): 칩렛·HBM·3D 로직에 필요', 13, TEXT))], anchor='m')
    source('Lau (2021, 2022); KnowMade (2024)', y=14.9)


def s_intro_gap():
    CS('1', '서론', CH[1][0], take='한 산업의 공정 단계 사이에서 일어나는 융합을, 개별 특허 수준에서, 시간에 따라 분석한다', script=(
        '그래서 이 현상을 기술융합 연구에서 찾아보았습니다. 기술융합 연구는 오랫동안 축적되어 왔지만, 아쉬운 점이 세 가지 있었습니다. '
        '첫째, 대부분 정보통신과 가전, 식품과 제약처럼 서로 다른 산업 사이의 융합을 다루었고, 한 산업 안에서 공정 단계를 가로지르는 융합은 드물었습니다. '
        '둘째, 융합을 분야 수준의 지수로 측정하다 보니 어떤 발명이 융합을 담고 있는지는 설명하기 어려웠습니다. 셋째, 이론은 융합을 단계적 과정으로 보지만, '
        '실증 연구는 융합의 정도가 늘었는지를 확인하는 데 머물렀습니다. 그래서 저는 한 산업의 공정 단계 사이에서 일어나는 융합을 개별 특허 수준에서, '
        '시간의 흐름에 따라 분석하였습니다.'))
    pw, py, ph = 13.3, 3.1, 13.0
    lx, rx = X0 + 0.4, X1 - 0.4 - pw
    rows = [('분석 범위', '서로 다른 산업 · 기술 분야 사이', 'ICT–가전, 식품–제약, 나노–바이오', '한 산업 안, 공정 단계 사이', '전공정 제조공정 ↔ 후공정 접합', 'FaSitemap'),
            ('측정 단위', '분야 수준의 공동분류 · 인용 지수', '어떤 발명이 융합을 담는지 설명 어려움', '개별 특허의 공동분류 유형', '어떤 발명이 경계를 넘는지 추정', 'FaFileLines'),
            ('시간 관점', '융합 정도의 증가 추세 확인', '이론(단계적 과정)과의 간극', '융합의 내용과 중심의 변화', '형성 · 확산 · 재편의 추적', 'FaClockRotateLeft')]
    for k, (x, head, col, fill) in enumerate(((lx, '기존 기술융합 연구', MUTED, PANEL), (rx, '본 연구', NAVY, PANEL2))):
        shape(MSO_SHAPE.RECTANGLE, x, py, pw, ph, line=col, lw=1.0, dash='dash')
        shape_text(shape(MSO_SHAPE.RECTANGLE, x + pw / 2 - 3.6, py - 0.5, 7.2, 1.0, fill=col), P((head, 15, WHITE, True), align=C))
        for i, r in enumerate(rows):
            yy = py + 1.0 + i * 3.9
            card(x + 0.5, yy, pw - 1.0, 3.4, fill=fill)
            icon(r[5], x + 0.95, yy + 0.35, 0.9, col)
            text(x + 2.1, yy + 0.3, pw - 3.0, 0.95, [P((r[0], 12, col, True))], anchor='m')
            t1, t2 = (r[1], r[2]) if k == 0 else (r[3], r[4])
            text(x + 0.95, yy + 1.35, pw - 1.9, 1.0, [P((t1, 16, TEXT if k == 0 else NAVY, True))], anchor='m')
            text(x + 0.95, yy + 2.35, pw - 1.9, 0.8, [P((t2, 12, MUTED))], anchor='m')
    shape(MSO_SHAPE.RIGHT_ARROW, lx + pw + 0.55, py + ph / 2 - 1.2, rx - lx - pw - 1.1, 2.4, fill=NAVY)
    source('Hacklin et al. (2009, 2010); Curran & Leker (2011); Bröring et al. (2006); No & Park (2010); Karvonen & Kässi (2012)', y=16.3)


# ================================================================ 02 이론적 배경
def s_th_convergence():
    CS('2.1', '기술융합의 개념', CH[2][0], take='공정 단계 간 융합 = 기술융합의 고전적 개념을 한 산업 내부의 공정 경계에 적용한 것', script=(
        '먼저 기술융합의 개념입니다. Rosenberg는 19세기 총기와 재봉틀을 위해 개발된 가공 기술이 공통의 생산공정을 통해 다른 산업으로 퍼져 나가고, '
        '그 공정 지식을 축적한 공작기계 산업이 독립된 산업으로 등장한 과정을 기술적 수렴이라고 불렀습니다. Kodama는 분리된 기술 분야가 결합하여 '
        '어느 한쪽만으로는 갖지 못한 역량을 만드는 과정을 기술융합으로 정의했습니다. Hacklin 등은 융합이 단계적으로 진행된다는 단계 모형을 제시했습니다. '
        '하이브리드 본딩에서도 두 공정 단계가 증착, 식각, 도금, 연마라는 공통의 공정과 장비를 공유하기 시작했고, 10마이크로미터 이하의 연결 간격은 '
        '두 영역이 결합해야만 달성됩니다. 즉, 공정 단계 간 융합은 기술융합의 고전적 개념을 한 산업 내부의 공정 경계에 적용한 것입니다.'))
    cols = [('기술적 수렴', 'Technological Convergence', 'Rosenberg (1963)', '공통의 생산공정을 통해 가공 기술이\n산업 경계를 넘어 확산\n→ 공작기계 산업의 등장', '두 공정 단계가 증착 · 식각 · 도금 · CMP\n공정과 장비를 공유'),
            ('기술융합', 'Technology Fusion', 'Kodama (1992)', '분리된 기술 분야의 결합으로\n어느 한 분야도 갖지 못한\n역량을 창출', '10 µm 이하 연결 간격은\n두 영역의 결합으로만 달성'),
            ('단계 모형', 'Stage Model of Convergence', 'Hacklin et al. (2009, 2010)', '지식 → 기술 → 응용 → 산업 융합\n단계 사이의 공진화적 순환', '특허로 관찰할 수 있는 것은\n"기술 융합" 단계')]
    w, g = 9.0, 1.94
    x0 = (SW - 3 * w - 2 * g) / 2
    for k, (ko, en, who, desc, hb) in enumerate(cols):
        x = x0 + k * (w + g)
        obox(x, 2.7, w, 2.3, ko, en, size=20, sub_size=11.5, color=NAVY, fill=WHITE, lw=2.0)
        text(x, 5.1, w, 0.75, [P((who, 13, MUTED, True), align=C)], anchor='m')
        card(x, 6.0, w, 4.1, fill=PANEL)
        text(x + 0.4, 6.0, w - 0.8, 4.1, [P((ln_, 14, TEXT), align=C) for ln_ in desc.split('\n')], anchor='m')
        arrow(x + w / 2, 10.25, x + w / 2, 11.05, NAVY, 2.0)
        card(x, 11.2, w, 3.4, fill=PANEL2)
        text(x + 0.4, 11.35, w - 0.8, 0.7, [P(('하이브리드 본딩에서는', 11.5, NAVY, True), align=C)], anchor='m')
        text(x + 0.4, 12.05, w - 0.8, 2.4, [P((ln_, 14, NAVY, True), align=C) for ln_ in hb.split('\n')], anchor='m')
        if k < 2:
            arrow(x + w + 0.25, 3.85, x + w + g - 0.25, 3.85, NAVY, 2.0)


def s_th_stage():
    CS('2.1', '융합의 단계 모형과 하이브리드 본딩', CH[2][0], take='본 연구가 특허 공동분류로 측정하는 것은 "기술 융합" 단계이다', script=(
        'Hacklin 등은 융합이 지식 융합, 기술 융합, 응용 융합, 산업 융합의 네 단계로 진행된다고 보았습니다. 하이브리드 본딩에 대입해 보면, '
        '평탄화나 표면 활성화 같은 제조공정 지식이 칩 연결 문제에 쓰이기 시작하는 것이 지식 융합, 그 지식이 접합 구조와 하나의 발명 안에서 결합되는 것이 '
        '기술 융합, 결합된 기술이 이미지센서나 HBM 같은 제품으로 구현되는 것이 응용 융합, 그리고 파운드리가 첨단 패키징을 내재화하는 것처럼 분업 구조가 '
        '바뀌는 것이 산업 융합입니다. 이 가운데 특허의 공동분류로 직접 관찰할 수 있는 것은 기술 융합 단계입니다. 응용 융합은 세 번째 연구 질문에서 '
        '간접적으로만 다루고, 산업 융합은 분석 범위 밖입니다.'))
    st = [('지식 융합', 'Knowledge', '제조공정 지식이 칩 연결 문제에 적용', '평탄화 · 표면 활성화 · 웨이퍼 직접 접합', '형성의 출발점', None),
          ('기술 융합', 'Technological', '그 지식이 접합 구조 · 연결 방법과 하나의 발명 안에서 결합', '공동분류로 관찰', '특허 공동분류로 직접 관찰', 'hl'),
          ('응용 융합', 'Applicational', '결합된 기술이 제품으로 구현', '이미지센서 · 3D NAND · HBM · 칩렛', 'RQ3에서 간접적으로', None),
          ('산업 융합', 'Industrial', '공정 단계 간 분업 구조의 재편', '파운드리의 첨단 패키징 내재화', '분석 범위 밖', 'out')]
    w, g = 7.35, 0.42
    x0 = (SW - 4 * w - 3 * g) / 2
    for k, (ko, en, d1, d2, status, kind) in enumerate(st):
        x = x0 + k * (w + g)
        hlk = kind == 'hl'
        sh = shape(MSO_SHAPE.CHEVRON if k else MSO_SHAPE.PENTAGON, x, 2.9, w + 0.35, 2.4, fill=NAVY if hlk else PANEL2)
        try: sh.adjustments[0] = 0.28
        except IndexError: pass
        shape_text(sh, [P((ko, 19, WHITE if hlk else NAVY, True), align=C), P((en, 11, 'C8D3E6' if hlk else MUTED), align=C)])
        card(x + 0.1, 5.8, w - 0.2, 6.0, fill=PANEL if not hlk else 'EAF0FA', line_=NAVY if hlk else None)
        text(x + 0.45, 6.15, w - 0.9, 2.9, [P((d1, 15, TEXT, True), align=C)], anchor='m')
        line(x + 1.2, 9.25, x + w - 1.2, 9.25, RULE, 1.0)
        text(x + 0.45, 9.45, w - 0.9, 2.1, [P((d2, 13, MUTED), align=C)], anchor='m')
        fill_, col_ = {'hl': (HL, NAVY), 'out': ('EEF0F3', MUTED), None: (WHITE, NAVY)}[kind]
        chip(x + 0.4, 12.35, w - 0.8, 1.3, status, fill=fill_, color=col_, size=13, line_=None if kind else 'C9D2E0')
    text(X0, 14.2, CW, 1.0, [P(('특허 공동분류는 서로 다른 분야의 기술이 하나의 발명 안에서 결합되는 "기술 융합" 단계를 관찰하는 데 적합하다', 13, MUTED), align=C)], anchor='m')
    source('Hacklin et al. (2010); Curran & Leker (2011); Lau (2022); KnowMade (2024)', y=15.3, align=C)


def s_th_boundary():
    CS('2.2', '수직 전문화와 지식 경계', CH[2][0], take='인터페이스가 불안정해지면, 지식의 결합은 조직보다 먼저 발명에서 나타난다', script=(
        '두 번째 이론은 수직 전문화와 지식 경계입니다. 반도체 산업은 설계, 제조, 조립이 분리되면서 팹리스, 파운드리, 조립·테스트 전문업체가 등장한 '
        '대표적인 산업입니다. 그런데 Kapoor와 Adner는 1974년부터 2005년까지 DRAM 산업의 12개 제품 세대를 분석하여, 외주화한 부품의 지식까지 보유한 '
        '비통합 기업이 기술 전환에서 성과 우위를 얻었음을 보였습니다. Brusoni 등도 구성요소 사이의 상호의존을 예측하기 어려울 때 기업은 직접 생산하는 '
        '범위보다 넓은 지식을 가져야 한다고 설명합니다. 즉, 생산의 경계와 지식의 경계는 다릅니다. 하이브리드 본딩처럼 인터페이스가 불안정해지면, '
        '그 징후는 조직이나 산업 구조보다 먼저 두 단계의 지식을 함께 담은 발명으로 나타날 것입니다.'))
    stages = [('설계', '팹리스'), ('제조', '종합반도체기업 · 파운드리'), ('조립 · 테스트', '조립·테스트 전문업체')]
    w, g, y = 7.3, 0.8, 3.5
    x0 = X0 + 0.6
    for k, (a, b) in enumerate(stages):
        x = x0 + k * (w + g)
        obox(x, y, w, 2.6, a, b, size=19, sub_size=12, color=NAVY, fill=WHITE)
        if k < 2: arrow(x + w + 0.12, y + 1.3, x + w + g - 0.12, y + 1.3, NAVY, 2)
    # production boundary (one firm = 제조), knowledge boundary wider
    px = x0 + (w + g)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, px - 0.3, y - 0.4, w + 0.6, 3.4, line=PROC, lw=2.5, radius=0.08)
    text(px - 0.3, y + 3.1, w + 0.6, 0.7, [P(('생산의 경계', 13, PROC, True), align=C)], anchor='m')
    kx0, kx1 = px - 0.55, px + 2 * w + g + 0.15 - 0.25 * w
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, kx0, y - 0.95, kx1 - kx0, 5.3, line=INTEG, lw=2.5, radius=0.06, dash='dash')
    text(kx1 - 5.0, y + 3.65, 4.8, 0.7, [P(('지식의 경계', 13, INTEG, True), align=R)], anchor='m')
    # right: statement
    rx = x0 + 3 * w + 2 * g + 0.9
    card(rx, 2.7, X1 - rx, 5.9, fill=PANEL2)
    text(rx + 0.4, 2.7, X1 - rx - 0.8, 5.9, [P(('생산의 경계', 18, NAVY, True), align=C), P(('≠', 22, NAVY, True), align=C),
                                            P(('지식의 경계', 18, INTEG, True), align=C)], anchor='m')
    cards = [('Kapoor & Adner (2012)', 'DRAM 산업 12개 제품 세대(1974–2005): 외주화한 부품의 지식까지 보유한 비통합 기업이 기술 전환에서 성과 우위'),
             ('Brusoni et al. (2001)', '구성요소 사이의 상호의존을 예측하기 어려울수록 직접 생산하는 범위보다 넓은 지식이 필요'),
             ('Jacobides et al. (2006)', '산업 아키텍처(누가 무엇을 하고 가치를 어떻게 나누는가)는 기술 변화에 따라 재편될 수 있음')]
    cw3 = (CW - 2 * 0.6) / 3
    for k, (h, t) in enumerate(cards):
        x = X0 + k * (cw3 + 0.6)
        card(x, 9.6, cw3, 5.0, fill=PANEL)
        text(x + 0.45, 9.85, cw3 - 0.9, 0.85, [P((h, 14, NAVY, True))], anchor='m')
        text(x + 0.45, 10.75, cw3 - 0.9, 3.7, [P((t, 13.5, TEXT), line=1.1)])
    source('Macher & Mowery (2004); Brown & Linden (2009); Kapoor (2013)', y=15.25, align=R)


def s_th_cycle():
    CS('2.3', '기술 순환 모형과 핵심–주변 구조', CH[2][0], take='지배적 설계가 자리 잡으면, 새로운 결합은 핵심이 아닌 주변에서 일어난다', script=(
        '세 번째는 기술 순환 모형과 핵심–주변 구조입니다. Anderson과 Tushman에 따르면 새로운 기술은 여러 설계가 경쟁하는 격동기를 거쳐 하나의 지배적 '
        '설계로 수렴하고, 그 뒤에는 정해진 설계를 다듬는 점진적 변화가 이어집니다. Murmann과 Frenken은 지배적 설계를, 다른 요소와 안정된 인터페이스 역할을 '
        '하는 핵심 요소의 설계로 보았습니다. 핵심 설계가 안정되면 이후의 기술 변화는 그 인터페이스를 전제로 주변 요소에서 이어집니다. 도로에 비유하면, '
        '규격이 정해지기 전에는 설계자와 시공자가 한 몸처럼 움직여야 하지만, 규격이 정해진 뒤에는 각자 그 규격을 전제로 일하게 되는 것과 같습니다.'))
    st = [('기술적 불연속', '새로운 기술 원리'), ('격동기', '여러 설계의 경쟁'), ('지배적 설계', '핵심 요소의 설계 안정'), ('점진적 변화', '정해진 설계의 정교화')]
    w, g, y = 6.6, 1.3, 2.8
    x0 = (SW - 4 * w - 3 * g) / 2
    for k, (a, b) in enumerate(st):
        x = x0 + k * (w + g)
        obox(x, y, w, 2.2, a, b, size=17, sub_size=12, color=NAVY, fill=NAVY if k == 2 else WHITE, tcolor=WHITE if k == 2 else NAVY,
             sub_color='C8D3E6' if k == 2 else MUTED)
        if k < 3: arrow(x + w + 0.15, y + 1.1, x + w + g - 0.15, y + 1.1, NAVY, 2)
    text(x0, y + 2.3, 4 * w + 3 * g, 0.7, [P(('기술 순환 모형 (Anderson & Tushman, 1990)', 12, MUTED), align=R)], anchor='m')
    # core-periphery diagram
    cx, cy = 9.4, 10.9
    card(X0, 6.2, 16.0, 8.9, fill=PANEL)
    for k in range(6):
        ang = math.radians(90 + k * 60)
        px, py = cx + 4.6 * math.cos(ang), cy - 3.0 * math.sin(ang)
        line(cx, cy, px, py, '9AA6B8', 1.5)
        shape_text(shape(MSO_SHAPE.OVAL, px - 1.0, py - 0.62, 2.0, 1.24, fill=WHITE, line='9AA6B8', lw=1.25), P(('주변', 11, MUTED, True), align=C), margin=0.02)
    shape_text(shape(MSO_SHAPE.OVAL, cx - 1.7, cy - 1.05, 3.4, 2.1, fill=NAVY), [P(('핵심 요소', 14, WHITE, True), align=C)], margin=0.02)
    text(X0 + 0.4, 6.35, 7.0, 0.8, [P(('핵심–주변 구조', 14, NAVY, True))], anchor='m')
    # right: two statements
    rx = X0 + 16.6
    for k, (h, t) in enumerate([('Murmann & Frenken (2006)', '지배적 설계 = 다른 요소와 안정된 인터페이스 역할을 하는 핵심 요소의 안정된 설계 → 이후의 변화는 주변 요소에서'),
                                ('Henderson & Clark (1990)', '확립된 아키텍처 지식은 조직의 의사소통 경로와 문제 해결 방식에 내재되어 당연한 전제가 됨')]):
        yy = 6.2 + k * 3.0
        card(rx, yy, X1 - rx, 2.75, fill=WHITE, line_='C9D2E0')
        text(rx + 0.45, yy + 0.2, X1 - rx - 0.9, 0.75, [P((h, 13.5, NAVY, True))], anchor='m')
        text(rx + 0.45, yy + 0.95, X1 - rx - 0.9, 1.75, [P((t, 13, TEXT), line=1.1)])
    card(rx, 12.3, X1 - rx, 2.8, fill=PANEL2)
    icon('FaRoad', rx + 0.5, 12.75, 1.3, NAVY)
    text(rx + 2.2, 12.4, X1 - rx - 2.6, 2.6, [P(('비유: 도로의 규격', 13.5, NAVY, True), after=3),
                                              P(('규격이 정해지기 전에는 설계와 시공이 한 몸처럼 움직이지만, 정해진 뒤에는 각자 규격을 전제로 일한다', 12.5, TEXT), line=1.1)], anchor='m')


def s_th_summary(num='', title='Chapter 2  Summary'):
    CS(num, title, CH[2][0], script=(
        '이론을 요약해 보면, 첫째, 공정 단계 간 융합은 기술융합의 개념을 한 산업 내부에 적용한 것이며, 특허로는 기술 융합 단계를 관찰할 수 있습니다. '
        '둘째, 생산의 경계와 지식의 경계는 다르기 때문에, 인터페이스가 불안정해지면 경계를 넘는 발명이 나타납니다. 셋째, 지배적 설계가 고착화되면 '
        '융합의 중심은 핵심에서 주변으로 이동할 수 있습니다. 이 세 가지를 기억해 주시면 감사하겠습니다.'))
    panels = [('기술융합 · 단계 모형', 'Rosenberg · Kodama · Hacklin', '공정 단계 간 융합', ('특허로 관찰하는', '"기술 융합" 단계'), '연구의 관점'),
              ('수직 전문화 · 지식 경계', 'Kapoor & Adner · Brusoni', '생산의 경계 ≠ 지식의 경계', ('불안정한 인터페이스', '경계를 넘는 발명'), 'RQ1'),
              ('기술 순환 · 핵심–주변', 'Anderson & Tushman · Murmann & Frenken', '지배적 설계의 고착화', ('핵심 기술', '주변 공정'), 'RQ2 · RQ3')]
    pw, g = 9.6, 0.9
    x0 = (SW - 3 * pw - 2 * g) / 2
    for k, (top, who, mid, (b1, b2), rq) in enumerate(panels):
        x = x0 + k * (pw + g)
        shape(MSO_SHAPE.RECTANGLE, x, 2.55, pw, 14.5, line='9AA6B8', lw=1.0, dash='dash')
        obox(x + 0.6, 3.1, pw - 1.2, 2.3, top, who, size=17, sub_size=10.5, color=NAVY, fill=NAVY, tcolor=WHITE, sub_color='C8D3E6')
        arrow(x + pw / 2, 5.5, x + pw / 2, 6.35, NAVY, 2)
        obox(x + 1.0, 6.45, pw - 2.0, 2.3, mid, None, size=16, color=NAVY, fill=WHITE)
        line(x + pw / 2, 8.85, x + pw / 2, 9.45, NAVY, 2)
        line(x + 2.35, 9.45, x + pw - 2.35, 9.45, NAVY, 2)
        for j, b in enumerate((b1, b2)):
            bx = x + 0.45 + j * (pw / 2 - 0.2)
            arrow(bx + (pw / 2 - 0.7) / 2, 9.45, bx + (pw / 2 - 0.7) / 2, 10.25, NAVY, 2)
            obox(bx, 10.35, pw / 2 - 0.7, 3.0, b, None, size=14.5, color=NAVY, fill=PANEL2, lw=1.25)
        if k == 2:
            arrow(x + 0.45 + pw / 2 - 0.7 + 0.05, 11.85, x + 0.45 + pw / 2 - 0.2 - 0.05, 11.85, INTEG, 2.5)
        chip(x + pw / 2 - 2.2, 14.55, 4.4, 1.15, rq, fill=HL, color=NAVY, size=14)


# ================================================================ 03 연구 질문 및 연구 방법
def s_rq1_why():
    CS('3.1', 'RQ1) 공정 단계 간 융합의 수준과 유형', CH[3][0], script=(
        '이제 연구 질문입니다. 첫 번째 연구 질문은 융합의 수준과 유형에 관한 것입니다. 앞에서 말씀드린 지식 경계의 논리에 따르면, 하이브리드 본딩으로 '
        '전공정과 후공정 사이의 인터페이스가 불안정해졌다면 두 영역의 기술을 함께 담은 발명이 예외적이지 않은 비중으로 나타나야 합니다. 그리고 발명을 '
        '기술 요소의 새로운 조합으로 보는 Fleming의 관점에 따르면, 두 영역에 걸친 발명은 한 영역에 머무는 발명보다 더 다양한 세부 기술을 포함하고 있을 것입니다. '
        '그래서 첫 번째 연구 질문은 "하이브리드 본딩 기술에서 제조공정과 접합 기술 간의 융합이 어느 정도로, 어떠한 형태로 나타나는가"입니다.'))
    y = 3.0
    obox(X0 + 0.4, y, 7.8, 3.0, '제조공정 기술', 'CPC H01L21', size=19, sub_size=13, color=PROC, fill=PROC_L)
    obox(X0 + 10.6, y, 7.8, 3.0, '접합·인터커넥트 기술', 'CPC H01L24', size=19, sub_size=13, color=BOND, fill=BOND_L)
    zx = X0 + 9.4
    pts = [(zx, y - 0.2), (zx - 0.45, y + 0.7), (zx + 0.45, y + 1.5), (zx - 0.45, y + 2.3), (zx, y + 3.2)]
    for a, b in zip(pts[:-1], pts[1:]):
        line(a[0], a[1], b[0], b[1], 'C0392B', 2.5)
    text(zx - 3.0, y + 3.3, 6.0, 0.8, [P(('불안정해진 인터페이스', 12.5, 'C0392B', True), align=C)], anchor='m')
    arrow(X0 + 4.3, y + 3.15, X0 + 7.0, y + 5.4, PROC, 2.5)
    arrow(X0 + 14.5, y + 3.15, X0 + 11.8, y + 5.4, BOND, 2.5)
    obox(X0 + 4.2, y + 5.6, 10.0, 3.0, '두 영역을 함께 담은 발명', '경계를 넘는 발명 (통합형)', size=19, sub_size=13, color=INTEG, fill=INTEG_L)
    text(X0, y + 9.0, 18.8, 2.7, [P(('발명 = 기존 지식 요소의 새로운 조합 (Fleming, 2001)', 13, MUTED), align=C, after=3),
                                 P(('특허 범위 = 포괄하는 분류의 수 (Lerner, 1994)', 13, MUTED), align=C)], anchor='t')
    rx = X0 + 20.0
    text(rx, 2.8, X1 - rx, 0.9, [P(('이론적 기대', 18, NAVY, True))], anchor='m')
    for k, (h, t) in enumerate([('기대 ①', '두 영역의 기술을 함께 포함하는 발명이 예외적이지 않은 비중으로 나타난다'),
                                ('기대 ②', '그러한 발명은 한 영역에 머무는 발명보다 더 다양한 세부 기술을 포함한다')]):
        yy = 4.0 + k * 5.0
        card(rx, yy, X1 - rx, 4.4, fill=PANEL2)
        chip(rx + 0.5, yy + 0.45, 2.6, 0.95, h, fill=NAVY, color=WHITE, size=13)
        text(rx + 0.5, yy + 1.6, X1 - rx - 1.0, 2.6, [P((t, 15.5, NAVY, True), line=1.12)], anchor='t')
    source('Kapoor & Adner (2012); Brusoni et al. (2001); Jacobides et al. (2006)', y=15.05, align=R)
    rq_band(1)


def s_rq2_why():
    CS('3.1', 'RQ2) 공정 단계 간 융합의 시간적 변화', CH[3][0], script=(
        '두 번째는 시간적 변화입니다. 단계 모형과 기술 순환 모형을 결합하면, 기술이 형성되는 시기에는 직접 접합을 가능하게 하는 공정 발명이 주를 이루고, '
        '확산되는 시기에는 공정 지식과 접합 발명의 결합이 늘며, 지배적 설계가 자리 잡은 뒤에는 결합이 응용에 필요한 주변 공정으로 옮겨 갈 것으로 기대할 수 '
        '있습니다. 실제로 하이브리드 본딩은 이미지센서와 3D NAND의 웨이퍼 단위 접합에서 시작해 HBM과 로직의 다이 단위 접합으로 확대되었고, 발명의 중심도 '
        '표면 준비나 평탄화 같은 공정에서 접합 구조와 적층으로 옮겨 갔습니다. 따라서 융합의 수준은 꾸준히 늘기보다 단계에 따라 오르내릴 수 있습니다. '
        '그래서 두 번째 연구 질문은 "공정 단계 간 기술융합은 시간에 따라 어떻게 변화해 왔는가"입니다.'))
    ph = [('기술 형성', '격동기', '직접 접합을 가능하게 하는\n공정 발명이 주를 이룸', '지식 융합', PROC, PROC_L),
          ('기술 확산', '설계의 수렴', '공정 지식이 접합 발명과\n결합하는 사례가 증가', '기술 융합', INTEG, INTEG_L),
          ('지배적 설계 이후', '점진적 변화', '두 기술의 결합이 응용에\n필요한 주변 공정으로 이동', '응용 융합', BOND, BOND_L)]
    w, g, y = 9.3, 0.9, 2.9
    x0 = (SW - 3 * w - 2 * g) / 2
    shape(MSO_SHAPE.RIGHT_ARROW, x0 - 0.2, y + 0.1, 3 * w + 2 * g + 0.4, 1.6, fill=PANEL2)
    for k, (t1, t2, d, stage, col, colL) in enumerate(ph):
        x = x0 + k * (w + g)
        text(x, y + 0.1, w, 1.6, [P((t1, 17, NAVY, True), (f'  ({t2})', 12.5, MUTED), align=C)], anchor='m')
        card(x, y + 2.1, w, 4.6, fill=colL)
        text(x + 0.3, y + 2.25, w - 0.6, 0.8, [P(('예상되는 발명', 12, col, True), align=C)], anchor='m')
        text(x + 0.3, y + 3.0, w - 0.6, 2.4, [P((ln_, 15.5, TEXT, True), align=C) for ln_ in d.split('\n')], anchor='m')
        chip(x + w / 2 - 2.4, y + 5.55, 4.8, 0.95, f'단계 모형: {stage}', fill=WHITE, color=col, size=12.5)
    card(X0, 10.1, CW, 4.6, fill=PANEL)
    text(X0 + 0.6, 10.35, 10.0, 0.8, [P(('하이브리드 본딩의 실제 전개', 14, NAVY, True))], anchor='m')
    rows = [('적용 범위', '웨이퍼 단위(W2W): CMOS 이미지센서 · 3D NAND', '다이 단위(D2W): HBM · 칩렛 · 3D 로직'),
            ('발명의 중심', '표면 준비 · 저온 어닐링 · 평탄화 (공정)', '접합 구조 · 적층')]
    for i, (h, a, b) in enumerate(rows):
        yy = 11.35 + i * 1.55
        chip(X0 + 0.6, yy, 3.3, 1.05, h, fill=NAVY, color=WHITE, size=12.5)
        text(X0 + 4.3, yy, 11.6, 1.05, [P((a, 12.5, TEXT))], anchor='m')
        arrow(X0 + 16.1, yy + 0.52, X0 + 17.6, yy + 0.52, NAVY, 2)
        text(X0 + 18.0, yy, 12.6, 1.05, [P((b, 12.5, TEXT, True))], anchor='m')
    source('Hacklin et al. (2009, 2010); Anderson & Tushman (1990); Jeong et al. (2015); Yun & Geum (2019); Lau (2021, 2022)', y=15.05, align=R)
    rq_band(2)


def s_rq3_why():
    CS('3.1', 'RQ3) 공정 단계 간 융합의 구조적 중심 변화', CH[3][0], script=(
        '세 번째는 융합의 중심 변화입니다. 하이브리드 본딩의 핵심은 범프 없이 접합 패드끼리 직접 연결하는 방식입니다. 이 인터페이스가 확정되지 않은 '
        '시기에는 핵심 기술을 발명하려면 평탄화나 웨이퍼 접합 같은 제조공정 지식을 함께 다루어야 합니다. 그러나 설계가 고착화되면 그 지식은 주어진 조건이 '
        '되고, 새로운 결합은 다이를 다루고 패키지로 완성하는 주변 공정에서 일어날 것입니다. 앞서 말씀드린 도로의 비유처럼, 규격이 정해지면 새로운 일은 '
        '그 규격 위에서 이루어집니다. 그래서 세 번째 연구 질문은 "공정 단계 간 기술융합의 중심이 핵심 기술과 주변 기술 사이에서 시기에 따라 어떻게 '
        '재편되는가"입니다. 이 세 질문에 답하기 위해 분석 대상과 자료를 선정하였습니다.'))
    pw, py, ph_ = 14.6, 2.75, 10.7
    for k, (head, sub) in enumerate([('인터페이스가 확정되기 전', '격동기'), ('핵심 설계가 고착화된 후', '지배적 설계 이후')]):
        x = X0 if k == 0 else X1 - pw
        card(x, py, pw, ph_, fill=PANEL)
        text(x + 0.5, py + 0.3, pw - 1.0, 0.9, [P((head, 17, NAVY, True), (f'  ({sub})', 12.5, MUTED))], anchor='m')
        cx_, cy_ = x + pw / 2, py + 5.6
        if k == 0:
            line(cx_ - 0.5, cy_, cx_ - 4.3, cy_ - 2.3, PROC, 6)
            obox(cx_ - 7.0, cy_ - 3.7, 5.2, 2.0, '제조공정 지식', '평탄화 · 웨이퍼 접합', size=14, sub_size=11, color=PROC, fill=PROC_L)
            text(x + 0.5, py + 8.6, pw - 1.0, 1.8, [P(('핵심 기술의 발명이 경계 반대편의 지식을', 13.5, TEXT), align=C), P(('함께 다루어야 함', 13.5, TEXT, True), align=C)], anchor='m')
        else:
            line(cx_ - 0.5, cy_, cx_ - 4.3, cy_ - 2.3, '9AA6B8', 1.25, dash='dash')
            obox(cx_ - 7.0, cy_ - 3.7, 5.2, 2.0, '제조공정 지식', '주어진 조건으로 전제', size=14, sub_size=11, color='9AA6B8', fill=WHITE, tcolor=MUTED)
            for j, (lab, dy) in enumerate((('다이 지지 · 적층', -2.6), ('봉지 · 패키지 완성', 2.4))):
                bx, by_ = cx_ + 2.6, cy_ + dy - 0.95
                line(cx_ + 1.6, cy_ + (0.3 if dy > 0 else -0.3), bx, by_ + 0.95, TEAL, 5)
                obox(bx, by_, 4.9, 1.9, lab, '주변 공정', size=13.5, sub_size=10.5, color=TEAL, fill='E2F5F3')
            text(x + 0.5, py + 8.6, pw - 1.0, 1.8, [P(('새로운 결합은 핵심의 인터페이스를 전제로', 13.5, TEXT), align=C), P(('주변 공정에서 일어남', 13.5, TEXT, True), align=C)], anchor='m')
        shape_text(shape(MSO_SHAPE.OVAL, cx_ - 1.9, cy_ - 1.2, 3.8, 2.4, fill=BOND), [P(('핵심 기술', 14, WHITE, True), align=C), P(('접합 패드 직접 연결', 10, WHITE), align=C)], margin=0.02)
    shape(MSO_SHAPE.RIGHT_ARROW, X0 + pw + 0.35, py + ph_ / 2 - 0.9, X1 - pw - X0 - pw - 0.7, 1.8, fill=NAVY)
    source('Murmann & Frenken (2006); Henderson & Clark (1990); Brusoni et al. (2001); Anderson & Tushman (1990)', y=13.75, align=R)
    text(X0, 14.3, CW, 1.3, [P(('기대: 핵심 기술을 포함한 발명이 제조공정 지식을 함께 포함하는 정도는 시기에 따라 줄고, 접합과 결합하는 제조공정의 중심은 웨이퍼 제조공정에서 주변 공정으로 이동할 것이다', 13, NAVY, True), align=C, line=1.1)], anchor='m')
    rq_band(3)


def s_object():
    CS('3.2', '분석 대상: 하이브리드 본딩 기술과 산업', CH[3][0], take='하이브리드 본딩은 기술적 · 산업적으로 두 공정 단계에 걸쳐 있다', script=(
        '분석 대상인 하이브리드 본딩을 조금 더 들여다보겠습니다. 왼쪽처럼 하이브리드 본딩은 기반 공정, 접합 형태, 소자 구조, 적용 제품의 네 층위로 '
        '정리할 수 있는데, 기반 공정은 CPC의 제조공정 클래스인 H01L21에, 접합 형태는 접합·인터커넥트 클래스인 H01L24에 주로 대응합니다. 오른쪽은 '
        '다이-투-웨이퍼 하이브리드 본딩의 실제 공정 흐름입니다. 다이를 자르기 전까지 절연막 증착, 다마신 식각, 구리 증착과 도금, 화학적 기계연마가 이어지는데, '
        '모두 전공정에서 쓰이던 공정과 장비입니다. 특허 활동을 주도하는 출원인에도 TSMC, Adeia, YMTC, Intel, Samsung처럼 라이선싱 기업, 파운드리, '
        '메모리 기업이 함께 들어 있습니다.'))
    layers = [('기반 공정', 'H01L21', ['CMP 평탄화', '세정·오염 제어', '표면 활성화', '정렬', '저온 어닐링'], PROC, PROC_L),
              ('접합 형태', 'H01L24', ['W2W', 'D2W', 'D2D'], BOND, BOND_L),
              ('소자 구조', '', ['적층 다이', '칩렛 집적', '메모리-온-로직'], NAVY, PANEL2),
              ('적용 제품', '', ['CIS', '3D NAND', 'HBM', '로직 적층', 'MicroLED'], NAVY, PANEL2)]
    lw_, y0 = 14.9, 2.6
    for k, (t, cpc, chips_, col, colL) in enumerate(layers):
        yy = y0 + k * 3.35
        card(X0, yy, lw_, 2.75, fill=colL)
        text(X0 + 0.35, yy + 0.12, 6.0, 0.8, [P((t, 14.5, col, True), ((f'  CPC {cpc}' if cpc else ''), 11, MUTED))], anchor='m')
        n = len(chips_); cwid = (lw_ - 0.7 - (n - 1) * 0.2) / n
        for j, c_ in enumerate(chips_):
            chip(X0 + 0.35 + j * (cwid + 0.2), yy + 1.05, cwid, 1.35, c_, fill=WHITE, color=TEXT, size=11.5 if n == 5 else 13, bold=False, line_='C9D2E0', radius=0.2)
        if k < 3:
            text(X0, yy + 2.72, lw_, 0.65, [P(('▼', 10, MUTED), align=C)], anchor='m')
    rx = X0 + lw_ + 0.8
    rw = X1 - rx
    iw_ = rw - 0.8
    img_h = iw_ * 945 / 1477
    shape(MSO_SHAPE.RECTANGLE, rx + 0.4, 2.6, iw_, img_h, fill=WHITE, line='C9D2E0', lw=1.0)
    picture('assets/d2w_flow_amat.png', rx + 0.5, 2.7, iw_ - 0.2, img_h - 0.2)
    text(rx, 2.65 + img_h, rw, 0.55, [P(('D2W 공정 흐름과 장비: Applied Materials (2025) · 출원인: KnowMade (2024)', 10, MUTED), align=R)], anchor='m')
    yy = 3.35 + img_h
    text(rx, yy, rw, 1.35, [P(('다이 절단 전까지 절연막 증착 → 다마신 식각 → Cu 배리어·시드 증착 → Cu 도금 → CMP: ', 12, TEXT, True),
                               ('모두 전공정의 공정과 장비', 12, PROC, True), line=1.05)], anchor='m')
    text(rx, yy + 1.5, 3.0, 0.8, [P(('주요 출원인', 11.5, MUTED, True))], anchor='m')
    cx_ = rx + 3.0
    for nm in ['TSMC', 'Adeia', 'YMTC', 'Intel', 'Samsung']:
        cw_ = 2.3 if nm == 'Samsung' else 1.75
        chip(cx_, yy + 1.52, cw_, 0.75, nm, fill=PANEL2, color=NAVY, size=11)
        cx_ += cw_ + 0.15


def s_data():
    CS('3.2', '분석 데이터: PATSTAT 하이브리드 본딩 관련 출원', CH[3][0], take='1968–2024년에 출원된 하이브리드 본딩 관련 특허출원 928건을 분석', script=(
        '분석 자료는 유럽특허청의 PATSTAT에서 추출했습니다. 먼저 제조공정 분류인 H01L21이나 접합 분류인 H01L24를 가진 출원을 찾고, 그중 제목에 '
        'hybrid bonding, direct bonding, Cu–Cu 같은 키워드가 들어간 출원을 골랐습니다. 그 결과 1968년부터 2024년까지 출원된 928건이 선정되었습니다. '
        '오른쪽 그래프처럼 출원은 2010년 이전에는 드물다가 2010년대 후반에 급증하여 2022년에 176건으로 정점에 이르렀습니다. 2023년 이후의 감소는 출원 후 '
        '공개까지의 시차 때문일 가능성이 큽니다.'))
    fx, fw = X0, 12.2
    steps = [('EPO PATSTAT Global', '유럽특허청 특허 데이터베이스', fw, NAVY2),
             ('1차: CPC 분류', 'H01L21(제조공정) 또는 H01L24(접합)', fw - 1.2, NAVY2),
             ('2차: 제목 키워드', 'hybrid bonding · direct bonding · Cu–Cu ·\ncopper to copper · metal-oxide bond', fw - 2.4, NAVY2)]
    yy = 2.75
    for t, sub, w, col in steps:
        x = fx + (fw - w) / 2
        sh = shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, yy, w, 2.45, fill=PANEL2, radius=0.1)
        shape_text(sh, [P((t, 15, NAVY, True), align=C)] + [P((ln_, 11.5, MUTED), align=C) for ln_ in sub.split('\n')])
        text(fx, yy + 2.45, fw, 0.6, [P(('▼', 11, NAVY), align=C)], anchor='m')
        yy += 3.05
    sh = shape(MSO_SHAPE.ROUNDED_RECTANGLE, fx + 2.2, yy, fw - 4.4, 2.6, fill=NAVY, radius=0.1)
    shape_text(sh, [P(('928건', 30, WHITE, True), align=C), P(('1968–2024년 출원', 13, 'DCE3EE'), align=C)])
    text(fx, yy + 2.75, fw, 0.6, [P(('출원–CPC 쌍 5,277건 · 서브그룹 기호 262개', 12, MUTED), align=C)], anchor='m')
    # yearly filings (native chart)
    yc = D['year_counts']
    years = list(range(1968, 2025))
    cd = CategoryChartData()
    cd.categories = [str(y) if y % 10 == 0 else '' for y in years]
    cd.add_series('출원 건수', [yc.get(str(y), 0) for y in years])
    cx = fx + fw + 0.9
    gf = K.s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, cm(cx), cm(2.6), cm(X1 - cx), cm(13.6), cd)
    ch = gf.chart
    ch.has_title = False
    style_chart(ch, legend=False, size=12)
    pl = ch.plots[0]; pl.gap_width = 35
    ser = pl.series[0]
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(NAVY)
    for i, y in enumerate(years):
        if y >= 2023:
            pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb('9AA6B8')
    i22 = years.index(2022)
    dl = ser.points[i22].data_label
    dl.has_text_frame = True; dl.text_frame.text = '176 (2022)'; dl.position = XL_LABEL_POSITION.OUTSIDE_END
    for p in dl.text_frame.paragraphs:
        for r in p.runs: font_run(r, 12, NAVY, True)
    ca = ch.category_axis
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    cax = ca._element
    for tag, val in (('c:tickLblSkip', '1'), ('c:tickMarkSkip', '10')):
        el = etree.SubElement(cax, qn(tag)); el.set('val', val)
    ch.value_axis.maximum_scale = 200; ch.value_axis.major_unit = 50
    axis_title(ch.value_axis, '출원 건수'); axis_title(ca, '출원연도')
    chip(X1 - 5.4, 3.0, 4.4, 0.85, '2023–24: 공개 시차', fill=PANEL2, color=MUTED, size=11, bold=False)


def s_measure():
    CS('3.3', '측정: 공동분류 유형', CH[3][0], take='H01L21과 H01L24를 함께 가진 출원 = 공정 단계의 경계를 넘는 "통합형"', script=(
        '그렇다면 융합은 어떻게 측정했을까요? 특허에는 기술 내용에 따라 CPC 분류 기호가 부여됩니다. 한 특허에 서로 다른 영역의 기호가 함께 붙어 있다면 '
        '두 영역의 지식이 하나의 발명 안에서 결합되었다는 흔적으로 볼 수 있습니다. 실제 출원 세 건을 보시면, 왼쪽은 제조공정 기호만 가진 공정 단독형, '
        '가운데는 제조공정 기호와 접합 기호를 함께 가진 통합형, 오른쪽은 접합 기호만 가진 접합 단독형입니다. 가운데의 통합형이 공정 단계의 경계를 넘는 '
        '특허입니다. 다만 이것은 대리측정치이며, 법적 청구범위를 측정한 것은 아닙니다.'))
    ex = [('공정 단독형', 'Process-only', 'H01L21만 보유', 'TW · 2019', 'DIRECT BONDING OF\nHETEROGENEOUS SUBSTRATES', ['21/2007', '21/4803', '21/76251', '21/76275'], PROC, PROC_L),
          ('통합형', 'Integrated', 'H01L21 + H01L24', 'US · 2017', 'Metal-metal direct\nbonding method', ['21/185', '21/187', '24/05'], INTEG, INTEG_L),
          ('접합 단독형', 'Bonding-only', 'H01L24만 보유', 'US · 2017', 'Hybrid bond pad\nstructure', ['24/00', '24/05', '24/08', '24/80'], BOND, BOND_L)]
    w, g = 9.3, 0.95
    x0 = (SW - 3 * w - 2 * g) / 2
    for k, (ko, en, rule_, who, title, syms, col, colL) in enumerate(ex):
        x = x0 + k * (w + g)
        sh = card(x, 2.7, w, 4.0, fill='5B6475' if k != 1 else '4B3C7A', shadow=True)
        text(x + 0.4, 2.85, w - 0.8, 0.7, [P((f'실제 출원 예시  ·  {who}', 11, 'DCE3EE'), align=C)], anchor='m')
        text(x + 0.4, 3.55, w - 0.8, 2.9, [P((ln_, 16, WHITE, True), align=C) for ln_ in title.split('\n')], anchor='m')
        text(x, 6.9, w, 0.7, [P(('부여된 CPC 기호 (H01L 생략)', 11, MUTED), align=C)], anchor='m')
        widths = [0.24 * len(sym) + 0.55 for sym in syms]
        sx_ = x + (w - sum(widths) - (len(syms) - 1) * 0.12) / 2
        for sym, cwid in zip(syms, widths):
            c_ = PROC if sym.startswith('21') else BOND
            chip(sx_, 7.7, cwid, 0.95, sym, fill=WHITE, color=c_, size=11, line_=c_)
            sx_ += cwid + 0.12
        arrow(x + w / 2, 8.9, x + w / 2, 9.75, col, 2.5)
        obox(x, 9.9, w, 2.6, ko, f'{en}  ·  {rule_}', size=20, sub_size=12, color=col, fill=colL, lw=2.5 if k == 1 else 1.5)
    # concept ladder
    card(X0, 13.2, CW, 2.95, fill=PANEL)
    lad = [('이론적 개념', '공정 단계 간 기술융합'), ('관찰 현상', '두 기술이 한 출원에 함께 포함'), ('조작적 지표 (H01L21–H01L24)', '공동분류 = 대리측정치')]
    lw3 = (CW - 1.2 - 2 * 1.2) / 3
    for k, (h, t) in enumerate(lad):
        x = X0 + 0.6 + k * (lw3 + 1.2)
        text(x, 13.35, lw3, 0.7, [P((h, 11.5, MUTED, True), align=C)], anchor='m')
        text(x, 14.0, lw3, 0.9, [P((t, 12.5, NAVY, True), align=C)], anchor='m')
        if k < 2: arrow(x + lw3 + 0.15, 14.45, x + lw3 + 1.05, 14.45, NAVY, 2)
    text(X0 + 0.6, 15.05, CW - 1.2, 0.9, [P(('CPC 기호는 심사관이 기술 내용에 따라 부여 → 법적 청구범위의 측정이 아님 (Curran & Leker, 2011; Fleming, 2001)', 11.5, MUTED), align=C)], anchor='m')


def s_core():
    CS('3.3', '측정: 핵심 기술의 식별', CH[3][0], take='접합 패드에 관한 기호 4개(24/03 · 24/05 · 24/08 · 24/80)를 하이브리드 본딩의 핵심 기술로 정의', script=(
        '세 번째 연구 질문을 분석하려면 무엇이 핵심 기술인지 먼저 정해야 합니다. 기술 클러스터 분석으로 200건 이상의 출원에 공통으로 부여된 기호를 찾고, '
        '기술 네트워크 분석으로 기호 사이의 연결 구조를 확인했습니다. 그 결과 접합 패드에 관한 네 기호, 즉 24/03, 24/05, 24/08, 24/80이 가장 많이 '
        '부여되면서 서로 가장 강하게 연결된 기호로 나타났습니다. 이 네 기호를 핵심 기술로 정의했고, 접합 기호를 가진 출원 763건 가운데 81.3%가 핵심 기술을 '
        '포함하고 있었습니다.'))
    iw = 14.0; ih = iw * 1560 / 1890
    picture('assets/network_ko.png', X0, 2.45, iw, ih)
    text(X0, 2.45 + ih - 0.1, iw, 0.55, [P(('회귀 자료의 기술 네트워크 (출원 30건 이상 기호, 논문 그림 4)', 10, MUTED), align=C)], anchor='m')
    rx = X0 + iw + 0.9
    rw = X1 - rx
    text(rx, 2.6, rw, 0.9, [P(('핵심 기술: 출원 200건 이상에 부여된 접합 기호', 15, NAVY, True))], anchor='m')
    rows = [['CPC 기호', '내용', '출원 수(비중)'],
            ['H01L 24/80', '접합 패드를 이용한 연결 방법', '529 (57.0%)'],
            ['H01L 24/08', '연결 후 접합 패드의 구조', '513 (55.3%)'],
            ['H01L 24/05', '연결 전 접합 패드의 구조', '350 (37.7%)'],
            ['H01L 24/03', '접합 패드의 제조 방법', '222 (23.9%)']]
    table(rx, 3.6, [3.9, rw - 3.9 - 3.6, 3.6], rows, [0.95] + [1.05] * 4, size=12.5, head_size=12.5, aligns=[C, L, C], bold_cols=(0,))
    yy = 9.25
    for k, (h, t) in enumerate([('클러스터', '네 기호는 출원 200건 이상에 부여된 H01L24 기호 전부'),
                                ('네트워크', '연결 강도 상위 4개 기호 = 이 네 기호, 네 기호의 6개 쌍 = 가장 강한 6개 연결'),
                                ('CPC 체계', '모두 접합 패드 그룹: 범프 · 층 · 와이어 없이 패드끼리 연결')]):
        chip(rx, yy + k * 1.45, 2.8, 1.05, h, fill=NAVY, color=WHITE, size=12)
        text(rx + 3.1, yy + k * 1.45, rw - 3.1, 1.05, [P((t, 12.5, TEXT))], anchor='m')
    card(rx, 13.85, rw, 2.15, fill=HL)
    text(rx + 0.4, 13.85, rw - 0.8, 2.15, [P(('접합 출원 763건 중 620건(81.3%)이 핵심 기술 포함', 15, NAVY, True), align=C)], anchor='m')


def s_method(core=False):
    CS('3.4', '분석 방법', CH[3][0], take='모든 모형에 출원청을 통제하고, 계수는 인과효과가 아닌 조건부 연관으로 해석한다', script=(
        '분석 방법은 세 가지 회귀모형입니다. 첫 번째 주모형은 다항 로짓으로, 출원이 세 유형 가운데 어디에 속하는지를 기술범위와 출원연도로 설명합니다. '
        '기술범위는 두 클래스 안에서 출원이 가진 세부 분류 기호의 수입니다. 두 번째는 제조공정 분류를 가졌는지를 보는 이항 로짓으로, 보완분석입니다. '
        '세 번째는 세 번째 연구 질문을 위한 이항 로짓으로, 접합 출원 763건에서 핵심 기술 보유와 출원 시기의 상호작용을 봅니다. 출원 시기는 출원이 드물었던 '
        '2014년 이전과 2015년부터 2019년, 2020년부터 2024년의 세 시기로 나누었습니다. 여기에 제조공정 영역별 결합 강도인 Jaccard 계수와 기술 네트워크의 '
        '매개 중심성을 보조 지표로 사용했습니다.'))
    boxes = [('모형 1 · 다항 로짓 (주모형)', 'RQ1 · RQ2', [('종속변수', '공동분류 유형 (기준 = 접합 단독형)'), ('설명변수', '기술범위, 출원연도'), ('보고', '상대위험비(RRR) · 예측확률'), ('표본', '928건')]),
             ('모형 2 · 이항 로짓 (보완)', 'RQ1 · RQ2', [('종속변수', '제조공정 분류 보유 여부'), ('설명변수', '기술범위, 출원연도'), ('보고', '승산비(OR)'), ('표본', '928건')]),
             ('모형 3 · 이항 로짓', 'RQ3', [('종속변수', '통합형 여부'), ('설명변수', '핵심 기술 × 출원 시기'), ('보고', '시기별 승산비(OR)'), ('표본', '접합 출원 763건')])]
    w, g = 9.6, 0.75
    x0 = (SW - 3 * w - 2 * g) / 2
    for k, (head, rq, items) in enumerate(boxes):
        x = x0 + k * (w + g)
        paras = []
        for lab, val in items:
            paras += [P((lab, 11.5, MUTED, True), after=0), P((val, 14.5, TEXT, True), after=7)]
        tab_box(x, 3.15, w, 7.2, head, paras, tab_w=w - 1.0)
        chip(x + w - 3.3, 9.25, 2.9, 0.8, rq, fill=HL, color=NAVY, size=12)
    # variables & auxiliaries
    vy = 11.0
    items = [('기술범위', '두 클래스 안의 서로 다른 CPC 서브그룹 수 (평균 5.69)'),
             ('출원 시기', '–2014 · 2015–2019 · 2020–2024 (첫 시기: 출원 희소)'),
             ('통제변수', '출원청 8개 (US · CN · KR · TW · JP · EP · WO · 기타)'),
             ('보조 지표 (RQ3)', 'Jaccard 계수(영역별 결합 강도) · 매개 중심성(연결 기호)')]
    if core:
        items = [('핵심 기술 (RQ3)', '접합 패드 기호 24/03 · 05 · 08 · 80: 출원 200건 이상, 클러스터 · 네트워크로 식별'),
                 ('기술범위', '두 클래스 안의 서로 다른 CPC 서브그룹 수 (평균 5.69)'),
                 ('출원 시기', '–2014 · 2015–2019 · 2020–2024 (첫 시기: 출원 희소)'),
                 ('통제 · 보조 지표', '출원청 8개 통제 · Jaccard 계수 · 매개 중심성 (RQ3)')]
    hw = (CW - 0.6) / 2
    for k, (h, t) in enumerate(items):
        x = X0 + (k % 2) * (hw + 0.6); yy = vy + (k // 2) * 2.25
        card(x, yy, hw, 1.95, fill=PANEL)
        chip(x + 0.35, yy + 0.45, 4.1, 1.05, h, fill=NAVY, color=WHITE, size=12)
        text(x + 4.75, yy, hw - 5.05, 1.95, [P((t, 13, TEXT))], anchor='m')


# ================================================================ 04 연구 결과
def s_res_rq1a():
    CS('4.1', 'RQ1) 융합의 수준', CH[4][0], take='두 공정 단계를 함께 담은 통합형이 가장 많았다: 공정 단계 간 융합은 예외가 아니다', script=(
        '이제 연구 결과를 말씀드리겠습니다. 첫 번째 연구 질문의 결과입니다. 928건 가운데 가장 많은 유형은 두 영역의 기호를 모두 가진 통합형으로, 425건, '
        '45.8%였습니다. 접합 단독형이 36.4%, 공정 단독형이 17.8%였습니다. 제조공정 분류에 패키징 단계의 공정도 들어 있다는 점을 고려해 이를 제외한 '
        '좁은 정의를 적용해도 통합형은 37.2%였습니다. 하이브리드 본딩에서 공정 단계 간 융합은 예외적인 현상이 아니었습니다.'))
    cd = CategoryChartData(); cd.categories = ['통합형', '접합 단독형', '공정 단독형']
    cd.add_series('유형', (425, 338, 165))
    gf = K.s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, cm(X0 + 0.2), cm(2.55), cm(11.6), cm(13.0), cd)
    ch = gf.chart
    ch.has_title = False
    style_chart(ch, legend=False, size=12, axes=False)
    ser = ch.plots[0].series[0]
    for i, col in enumerate((INTEG, BOND, PROC)):
        pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(col)
        pt.format.line.color.rgb = rgb(WHITE); pt.format.line.width = Pt(2)
    dn = ch._chartSpace.find('.//' + qn('c:doughnutChart'))
    hs = dn.find(qn('c:holeSize'))
    if hs is None: hs = etree.SubElement(dn, qn('c:holeSize'))
    hs.set('val', '58')
    text(X0 + 0.2, 7.55, 11.6, 3.0, [P(('N = 928', 20, NAVY, True), align=C), P(('출원', 12, MUTED), align=C)], anchor='m')
    # legend with numbers
    lx = X0 + 12.4
    for k, (lab, n, pct, col) in enumerate([('통합형', 425, '45.8%', INTEG), ('접합 단독형', 338, '36.4%', BOND), ('공정 단독형', 165, '17.8%', PROC)]):
        yy = 4.2 + k * 3.2
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, lx, yy + 0.15, 0.55, 2.1, fill=col, radius=0.3)
        text(lx + 0.9, yy, 6.0, 0.9, [P((lab, 15, col, True))], anchor='m')
        text(lx + 0.9, yy + 0.9, 6.0, 1.4, [P((pct, 26 if k == 0 else 22, TEXT, True), (f'  {n}건', 13, MUTED))], anchor='m')
    rx = X0 + 20.4
    rq_panel(rx, 2.75, X1 - rx, 1, [
        P(('통합형 45.8% — 가장 큰 비중', 16, NAVY, True), after=8),
        P(('좁은 정의(패키징 단계 공정 제외)에서도 37.2%', 13.5, TEXT, True), bullet=('■', NAVY, 70), indent=0.5, after=5),
        P(('제조공정 기호 보유 출원 590건 (63.6%)', 13.5, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=5),
        P(('→ 공정 단계 간 융합은 예외가 아님', 14, INTEG, True))], h=9.6)


def s_res_rq1b():
    CS('4.1', 'RQ1) 융합의 형태: 기술범위와 통합형', CH[4][0], take='세부 기술을 더 다양하게 포함한 출원일수록 통합형일 가능성이 높았다', script=(
        '그렇다면 어떤 특허가 경계를 넘었을까요? 다항 로짓 결과, 세부 분류가 하나 늘 때마다 접합 단독형 대비 통합형일 상대위험이 40% 높아졌습니다. '
        '그래프로 보면, 기술범위가 3인 출원은 통합형일 확률이 0.26이지만 8인 출원은 0.68입니다. 즉, 경계를 넘는 특허는 더 다양한 세부 기술을 포함하고 '
        '있었습니다. 다만 통합형은 정의상 기호가 두 개 이상이라는 기계적인 측면이 있어, 유형의 최소 기호 수를 뺀 지표로도 확인했는데 상대위험비는 1.262로 '
        '방향이 유지되었습니다.'))
    pb = D['pred_breadth']
    bs = list(range(1, 16))
    cd = CategoryChartData(); cd.categories = [str(b) for b in bs]
    cd.add_series('통합형', [round(pb[str(b)][1], 3) for b in bs])
    cd.add_series('접합 단독형', [round(pb[str(b)][0], 3) for b in bs])
    cd.add_series('공정 단독형', [round(pb[str(b)][2], 3) for b in bs])
    gf = K.s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, cm(X0), cm(2.5), cm(16.6), cm(14.0), cd)
    ch = gf.chart
    style_chart(ch, legend=True, size=12)
    for ser, col in zip(ch.plots[0].series, (INTEG, BOND, PROC)):
        ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(3.5); ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.NONE
    integ = ch.plots[0].series[0]
    for idx in (2, 7):
        pt = integ.points[idx]
        pt.marker.style = XL_MARKER_STYLE.CIRCLE; pt.marker.size = 9
        pt.marker.format.fill.solid(); pt.marker.format.fill.fore_color.rgb = rgb(INTEG); pt.marker.format.line.color.rgb = rgb(INTEG)
        dl = pt.data_label; dl.has_text_frame = True
        dl.text_frame.text = '0.26' if idx == 2 else '0.68'
        dl.position = XL_LABEL_POSITION.ABOVE if idx == 7 else XL_LABEL_POSITION.BELOW
        for p in dl.text_frame.paragraphs:
            for r in p.runs: font_run(r, 13, INTEG, True)
    ch.value_axis.maximum_scale = 1.0; ch.value_axis.minimum_scale = 0; ch.value_axis.major_unit = 0.2
    ch.value_axis.tick_labels.number_format = '0.0'; ch.value_axis.tick_labels.number_format_is_linked = False
    axis_title(ch.value_axis, '예측확률'); axis_title(ch.category_axis, '기술범위 (두 클래스 내 CPC 서브그룹 수)')
    rx = X0 + 17.3
    rw = X1 - rx
    rows = [['모형', '관계', '효과비'],
            ['다항 로짓', '기술범위 → 통합형', 'RRR 1.403***'],
            ['', '기술범위 → 공정 단독형', 'RRR 0.905**'],
            ['이항 로짓', '기술범위 → 제조공정 분류', 'OR 1.267***'],
            ['예측확률', '통합형: 범위 3 → 8', '0.26 → 0.68']]
    table(rx, 2.75, [2.7, rw - 2.7 - 3.6, 3.6], rows, [0.9, 1.0, 1.0, 1.0, 1.0], size=12.5, head_size=12, aligns=[C, L, C],
          hl=[(1, 2), (4, 2)], merge=[((1, 0), (2, 0))])
    text(rx, 7.75, rw, 0.6, [P(('*** p < 0.01, ** p < 0.05 · 다항 로짓 기준 범주 = 접합 단독형 · 출원청 통제', 10, MUTED))], anchor='m')
    rq_panel(rx, 8.55, rw, 1, [
        P(('세부 분류 1개 증가 → 통합형 상대위험 +40%', 14, NAVY, True), after=6),
        P(('정의상 최소 기호 수 차이를 뺀 지표로도 RRR 1.262 (방향 유지)', 12.5, TEXT), bullet=('■', NAVY, 70), indent=0.5)], h=7.85)


def s_res_rq2a(reg=False):
    CS('4.2', 'RQ2) 융합의 시간적 변화: 세 국면', CH[4][0], take='융합은 단선적으로 늘지 않았다: 형성기 → 확산기 → 재편기의 세 국면', script=(
        '두 번째 연구 질문의 결과입니다. 출원을 세 시기로 나누어 보았습니다. 2014년 이전 47년간의 출원이 179건으로 이후 5년간의 209건보다 적을 만큼 '
        '드물었기 때문에 하나의 시기로 묶었습니다. 2014년 이전에는 공정 단독형이 43.6%로 가장 많았고, 2015년부터 2019년까지는 통합형이 62.2%로 정점에 '
        '이르렀으며, 2020년 이후에는 통합형이 44.3%로 낮아지고 접합 단독형이 45.4%까지 늘었습니다. 저는 이 세 국면을 형성기, 확산기, 재편기라고 이름 붙였습니다. '
        '융합은 계속 늘어난 것이 아니라 오르내렸습니다.'))
    cd = CategoryChartData()
    cd.categories = ['형성기\n–2014\n(N = 179)', '확산기\n2015–2019\n(N = 209)', '재편기\n2020–2024\n(N = 540)']
    cd.add_series('공정 단독형', (43.6, 14.8, 10.4)); cd.add_series('통합형', (31.3, 62.2, 44.3)); cd.add_series('접합 단독형', (25.1, 23.0, 45.4))
    gf = K.s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, cm(X0), cm(2.5), cm(17.4), cm(14.1), cd)
    ch = gf.chart
    style_chart(ch, legend=True, size=12.5)
    pl = ch.plots[0]; pl.gap_width = 45; pl.overlap = 100
    for ser, col in zip(pl.series, (PROC, INTEG, BOND)):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
        dl = ser.data_labels
        dl.show_value = True; dl.number_format = '0.0"%"'; dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.CENTER
        dl.font.size = Pt(14); dl.font.bold = True; dl.font.color.rgb = rgb(WHITE); dl.font.name = FONT
    ch.value_axis.maximum_scale = 100; ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
    rx = X0 + 18.2
    rw = X1 - rx
    ph = [('형성기', '–2014', '공정 단독형 43.6%로 최다', '직접 접합을 가능하게 하는 공정 발명이 단독으로', PROC, PROC_L),
          ('확산기', '2015–2019', '통합형 62.2%로 정점', '두 공정의 기술을 하나의 발명에 함께', INTEG, INTEG_L),
          ('재편기', '2020–2024', '접합 단독형 45.4% ≈ 통합형 44.3%', '통합형 감소 · 접합 단독형 증가', BOND, BOND_L)]
    for k, (nm, yr, a, b, col, colL) in enumerate(ph):
        yy = 2.75 + k * 3.55
        card(rx, yy, rw, 3.2, fill=colL)
        text(rx + 0.45, yy + 0.2, 4.4, 1.0, [P((nm, 17, col, True))], anchor='m')
        text(rx + 0.45, yy + 1.1, 4.4, 0.8, [P((yr, 12, MUTED))], anchor='m')
        text(rx + 4.6, yy + 0.25, rw - 5.0, 1.3, [P((a, 15, TEXT, True))], anchor='m')
        text(rx + 4.6, yy + 1.55, rw - 5.0, 1.3, [P((b, 12.5, MUTED))], anchor='m')
    if reg:
        card(rx, 13.45, rw, 3.1, fill=HL)
        text(rx + 0.45, 13.5, rw - 0.9, 3.0, [P(('출원연도 → 공정 단독형 RRR 0.913*** (연 8.7% 감소)', 13.5, NAVY, True), after=4),
                                             P(('예측확률 2010 → 2022: 공정 단독형 0.27 → 0.12, 접합 단독형 0.31 → 0.40', 12, TEXT), after=3),
                                             P(('첫 시기는 47년간 179건 < 이후 5년간 209건 → 하나로 묶음', 11, MUTED))], anchor='m')
    else:
        text(rx, 13.5, rw, 2.9, [P(('Research Question 2', 15, NAVY, True), after=3), P((RQ[2], 12, MUTED), after=6),
                                 P(('첫 시기는 47년간 179건으로 이후 5년간 209건보다 적어 하나로 묶음', 11, MUTED))])


def s_res_rq2b():
    CS('4.2', 'RQ2) 출원의 무게중심 이동', CH[4][0], take='출원의 무게중심은 제조공정에서 접합 쪽으로 이동하였다', script=(
        '회귀분석에서도 같은 흐름이 확인되었습니다. 출원연도가 1년 늦어질 때마다 접합 단독형 대비 공정 단독형의 상대위험은 약 8.7%씩 감소했습니다. '
        '2010년 출원과 2022년 출원을 비교하면 공정 단독형 예측확률은 0.27에서 0.12로 줄고, 접합 단독형은 0.31에서 0.40으로 늘었습니다. 출원의 무게중심이 '
        '제조공정에서 접합 쪽으로 이동한 것입니다. 한편 기술범위가 넓을수록 통합형이 되는 경향은 세 국면 모두에서 유지되었습니다.'))
    py_ = D['pred_year']
    ys = list(range(2005, 2025))
    cd = CategoryChartData(); cd.categories = [str(y) if y % 5 == 0 or y == 2024 else '' for y in ys]
    cd.add_series('통합형', [round(py_[str(y)][1], 3) for y in ys])
    cd.add_series('접합 단독형', [round(py_[str(y)][0], 3) for y in ys])
    cd.add_series('공정 단독형', [round(py_[str(y)][2], 3) for y in ys])
    gf = K.s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, cm(X0), cm(2.5), cm(16.6), cm(14.0), cd)
    ch = gf.chart
    style_chart(ch, legend=True, size=12)
    sers = list(ch.plots[0].series)
    for ser, col in zip(sers, (INTEG, BOND, PROC)):
        ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(3.5); ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.NONE
    for ser, col, vals in ((sers[1], BOND, ('0.31', '0.40')), (sers[2], PROC, ('0.27', '0.12'))):
        for idx, v in zip((ys.index(2010), ys.index(2022)), vals):
            pt = ser.points[idx]
            pt.marker.style = XL_MARKER_STYLE.CIRCLE; pt.marker.size = 9
            pt.marker.format.fill.solid(); pt.marker.format.fill.fore_color.rgb = rgb(col); pt.marker.format.line.color.rgb = rgb(col)
            dl = pt.data_label; dl.has_text_frame = True; dl.text_frame.text = v
            dl.position = XL_LABEL_POSITION.BELOW if (col == PROC or idx == ys.index(2022)) else XL_LABEL_POSITION.ABOVE
            for p in dl.text_frame.paragraphs:
                for r in p.runs: font_run(r, 13, col, True)
    ch.value_axis.maximum_scale = 1.0; ch.value_axis.minimum_scale = 0; ch.value_axis.major_unit = 0.2
    ch.value_axis.tick_labels.number_format = '0.0'; ch.value_axis.tick_labels.number_format_is_linked = False
    ca = ch.category_axis._element
    el = etree.SubElement(ca, qn('c:tickLblSkip')); el.set('val', '1')
    axis_title(ch.value_axis, '예측확률'); axis_title(ch.category_axis, '출원연도')
    rx = X0 + 17.3
    rw = X1 - rx
    rows = [['분석', '효과비'],
            ['출원연도 → 공정 단독형 (연간)', 'RRR 0.913***'],
            ['출원연도 → 통합형 (연간)', 'RRR 0.993'],
            ['공정 단독형 (기간 더미)', '0.273*** / 0.101***'],
            ['통합형 (기간 더미)', '1.707* / 0.762'],
            ['기술범위 → 통합형 (국면별)', '1.179 · 1.513 · 1.414***']]
    table(rx, 2.75, [rw - 6.4, 6.4], rows, [0.9] + [1.0] * 5, size=11.5, head_size=12, aligns=[L, C], hl=[(1, 1), (3, 1)])
    text(rx, 8.75, rw, 0.6, [P(('기간 더미의 기준 = 2014년 이전 · *** p < 0.01, * p < 0.10', 10, MUTED))], anchor='m')
    rq_panel(rx, 9.55, rw, 2, [
        P(('공정 단독형 상대위험 연 8.7% 감소', 14, NAVY, True), after=5),
        P(('2010 → 2022 예측확률: 공정 단독형 0.27 → 0.12, 접합 단독형 0.31 → 0.40', 12.5, TEXT), bullet=('■', NAVY, 70), indent=0.5)], h=6.85)


def or_plot(x, y, w, rows, rs=1.25, label_w=2.6, size=14):
    """odds ratios on a log2 axis (shapes). rows: (name, or, lo, hi, label, colour)"""
    ax0, ax1 = x + label_w + 0.4, x + w - 0.4
    lo_, hi_ = math.log2(0.25), math.log2(32)
    def xp(v): return ax0 + (math.log2(v) - lo_) / (hi_ - lo_) * (ax1 - ax0)
    n = len(rows)
    line(xp(1), y, xp(1), y + n * rs + 0.1, '9AA3B2', 1.5, dash='dash')
    for i, (nm, orv, l, h, lab, col) in enumerate(rows):
        yy = y + 0.55 + i * rs
        text(x, yy - 0.45, label_w, 0.9, [P((nm, size, col, True))], anchor='m')
        line(xp(l), yy, xp(h), yy, col, 3.5)
        shape(MSO_SHAPE.OVAL, xp(orv) - 0.27, yy - 0.27, 0.54, 0.54, fill=col)
        if i == 0:
            text(xp(orv) - 4.4, yy - 1.0, 3.9, 0.6, [P((lab, 12.5, col, True), align=R)], anchor='m')
        else:
            text(xp(h) + 0.25, yy - 0.32, 4.6, 0.64, [P((lab, 12.5, col, True))], anchor='m')
    axy = y + n * rs + 0.2
    line(ax0, axy, ax1, axy, '9AA3B2', 1.25)
    for v in (0.25, 0.5, 1, 2, 4, 8, 16, 32):
        line(xp(v), axy, xp(v), axy + 0.15, '9AA3B2', 1.25)
        text(xp(v) - 0.8, axy + 0.15, 1.6, 0.55, [P((('%g' % v), 11, MUTED), align=C)], anchor='m')
    text(xp(1) - 3.5, axy + 0.7, 7.0, 0.55, [P(('승산비 (로그 척도, 1 = 차이 없음)', 11, MUTED), align=C)], anchor='m')


def s_res_rq3a():
    CS('4.3', 'RQ3) 핵심 기술의 경계 넘기 우위', CH[4][0], take='핵심 기술의 "경계 넘기" 우위는 형성기에 뚜렷했으나 재편기에는 확인되지 않았다', script=(
        '세 번째 연구 질문의 결과입니다. 핵심 기술을 포함한 접합 출원은 형성기에 79.6%가 통합형이었지만, 핵심 기술이 없는 출원은 27.7%에 그쳤습니다. '
        '출원청을 통제한 승산비는 11.1이었습니다. 그런데 이 우위는 확산기에 2.5로 줄었고, 재편기에는 0.66으로 유의한 차이가 확인되지 않았습니다. '
        '그사이 접합 출원 가운데 핵심 기술을 포함한 비중은 53.5%에서 88.0%로 높아졌습니다. 발명이 하나의 핵심 설계로 수렴하면서, 그 핵심을 다루는 발명이 '
        '다른 접합 발명보다 제조공정 지식을 더 자주 담지는 않게 된 것으로 해석됩니다.'))
    rows = [['출원 시기', '통합형 비중\n핵심 보유', '통합형 비중\n핵심 없음', '핵심 기술의\n승산비', '95% CI'],
            ['형성기 (–2014)', '79.6%', '27.7%', '11.071***', '4.211–29.107'],
            ['확산기 (2015–19)', '77.1%', '57.9%', '2.452**', '1.158–5.194'],
            ['재편기 (2020–24)', '48.4%', '56.9%', '0.658', '0.374–1.160']]
    tw = 19.0
    table(X0, 2.75, [5.0, 3.4, 3.4, 3.2, 4.0], rows, [1.35, 1.1, 1.1, 1.1], size=12, head_size=12, aligns=[L, C, C, C, C],
          hl=[(1, 3), (3, 3)], bold_cols=(0,))
    text(X0, 7.5, tw, 0.6, [P(('접합 출원 N = 763 · 종속변수: 통합형 여부 · 출원청 통제 · *** p < 0.01, ** p < 0.05', 10, MUTED))], anchor='m')
    card(X0, 8.35, tw, 7.9, fill=PANEL)
    text(X0 + 0.5, 8.5, tw - 1.0, 0.8, [P(('시기별 핵심 기술의 승산비 (95% 신뢰구간)', 13.5, NAVY, True))], anchor='m')
    or_plot(X0 + 0.5, 9.45, tw - 1.0, [('형성기', 11.071, 4.211, 29.107, '11.07***', PROC), ('확산기', 2.452, 1.158, 5.194, '2.45**', INTEG),
                                     ('재편기', 0.658, 0.374, 1.160, '0.66 (비유의)', BOND)], rs=1.25)
    rx = X0 + tw + 0.8
    rw = X1 - rx
    rq_panel(rx, 2.75, rw, 3, [
        P(('형성기 11.1 → 확산기 2.5 → 재편기 0.66', 14, NAVY, True), after=6),
        P(('2014년 이전 대비 배수 (상호작용)', 12, MUTED, True), after=2),
        P(('2015–19: 0.221** · 2020–24: 0.059***', 13, TEXT, True), after=8),
        P(('접합 출원 중 핵심 기술 비중', 12, MUTED, True), after=2),
        P(('53.5% → 78.7% → 88.0%', 15, TEXT, True), after=8),
        P(('→ 발명이 하나의 핵심 설계로 수렴', 13, BOND, True))], h=13.5, qh=1.95)


def s_res_rq3b():
    CS('4.3', 'RQ3) 결합 중심의 이동', CH[4][0], take='핵심 기술은 중심에 남고, 접합과 제조공정을 잇는 연결은 공정 장치 · 패키징 단계 공정으로 이동하였다', script=(
        '그렇다면 융합은 어디로 옮겨 갔을까요? 접합 기호와 함께 부여되는 정도를 제조공정 영역별로 보면, 웨이퍼 제조공정은 확산기 0.598에서 재편기 0.159로 '
        '크게 낮아진 반면, 공정 장치와 패키징 단계 공정은 꾸준히 높아졌습니다. 기술 네트워크에서도 핵심 기호는 세 시기 내내 허브의 자리를 지켰지만, '
        '접합과 제조공정을 잇는 연결 기호는 웨이퍼 제조공정에서 공정 장치로 바뀌었습니다. 핵심 기술과 함께 부여된 공정도 반도체 몸체끼리의 접합에서 '
        '임시 기판을 이용한 다이 지지와 봉지 공정으로 바뀌었습니다.'))
    cd = CategoryChartData(); cd.categories = ['형성기', '확산기', '재편기']
    cd.add_series('웨이퍼 제조공정', (0.306, 0.598, 0.159)); cd.add_series('공정 장치', (0.046, 0.059, 0.176)); cd.add_series('패키징 단계 공정', (0.194, 0.234, 0.268))
    gf = K.s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, cm(X0), cm(2.5), cm(14.6), cm(9.6), cd)
    ch = gf.chart
    style_chart(ch, legend=True, size=12)
    for ser, col in zip(ch.plots[0].series, (PROC, TEAL, GOLD)):
        ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(3.5); ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = 9
        ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col); ser.marker.format.line.color.rgb = rgb(col)
        dl = ser.data_labels; dl.show_value = True; dl.number_format = '0.000'; dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.ABOVE if col != TEAL else XL_LABEL_POSITION.BELOW
        dl.font.size = Pt(11); dl.font.bold = True; dl.font.color.rgb = rgb(col); dl.font.name = FONT
    ch.value_axis.maximum_scale = 0.7; ch.value_axis.minimum_scale = 0; ch.value_axis.major_unit = 0.2
    for ser, col, pos in ((ch.plots[0].series[0], PROC, XL_LABEL_POSITION.BELOW), (ch.plots[0].series[1], TEAL, XL_LABEL_POSITION.ABOVE)):
        dl = ser.points[2].data_label; dl.has_text_frame = True
        dl.text_frame.text = '0.159' if col == PROC else '0.176'; dl.position = pos
        for p in dl.text_frame.paragraphs:
            for r in p.runs: font_run(r, 11, col, True)
    ch.value_axis.tick_labels.number_format = '0.0'; ch.value_axis.tick_labels.number_format_is_linked = False
    axis_title(ch.value_axis, '접합 기호와의 Jaccard 계수')
    rows = [['매개 중심성 비중', '형성기', '확산기', '재편기'],
            ['웨이퍼 제조공정', '53.6%', '42.5%', '9.0%'],
            ['공정 장치', '0.0%', '1.8%', '26.1%'],
            ['패키징 단계 공정', '0.3%', '0.6%', '7.3%'],
            ['접합 기호(H01L24)', '46.1%', '55.1%', '57.6%']]
    table(X0, 12.35, [5.2, 3.1, 3.1, 3.2], rows, [0.82] * 5, size=12, head_size=11.5, aligns=[L, C, C, C], hl=[(1, 3), (2, 3)], bold_cols=(0,))
    rx = X0 + 15.4
    rw = X1 - rx
    card(rx, 2.75, rw, 5.6, fill=PANEL)
    text(rx + 0.5, 2.95, rw - 1.0, 0.8, [P(('핵심 기술과 함께 부여된 제조공정 기호', 14, NAVY, True))], anchor='m')
    for k, (ph_, syms, col) in enumerate([('확산기', [('21/187', '반도체 몸체끼리의 접합'), ('21/76898', '기판 관통 배선')], PROC),
                                         ('재편기', [('21/6835', '임시 기판을 이용한 지지'), ('21/568', '봉지 공정용 임시 기판')], BOND)]):
        yy = 4.0 + k * 2.1
        chip(rx + 0.5, yy + 0.25, 2.4, 1.1, ph_, fill=col, color=WHITE, size=12.5)
        for j, (sy, lab) in enumerate(syms):
            text(rx + 3.2 + j * ((rw - 3.7) / 2), yy, (rw - 3.7) / 2, 0.8, [P((sy, 13.5, col, True))], anchor='m')
            text(rx + 3.2 + j * ((rw - 3.7) / 2), yy + 0.75, (rw - 3.7) / 2, 0.8, [P((lab, 11.5, TEXT))], anchor='m')
    card(rx, 8.65, rw, 3.25, fill=PANEL)
    text(rx + 0.5, 8.8, rw - 1.0, 0.8, [P(('기술 네트워크의 허브와 연결 기호', 14, NAVY, True))], anchor='m')
    text(rx + 0.5, 9.6, rw - 1.0, 2.3, [P(('허브(연결 강도 상위 3개): 세 시기 모두 핵심 기호 24/80 · 24/08 · 24/05', 12, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=3),
                                        P(('연결 기호(매개 중심성 최대 제조공정 기호): 21/187 → 21/67121 · 21/67766 (공정 장치)', 12, TEXT), bullet=('■', NAVY, 70), indent=0.5)])
    rq_panel(rx, 12.2, rw, 3, [P(('융합의 중심: 웨이퍼 제조공정 → 공정 장치 · 패키징 단계 공정', 13, NAVY, True))], h=4.1)


def s_robust():
    CS('4.4', '강건성 검증', CH[4][0], take='RQ1 · RQ3는 모든 재추정에서, RQ2는 hybrid bonding 표본을 제외한 모든 재추정에서 유지되었다', script=(
        '결과가 표본이나 정의에 따라 달라지는지도 확인했습니다. 같은 제목의 출원을 군집으로 본 표준오차, 같은 제목의 중복을 제거한 표본, 패키징 단계 공정을 '
        '제외한 좁은 정의, 검색어가 다른 뜻으로 쓰인 68건을 제외한 표본, 제목에 hybrid bonding이 들어간 467건 표본 등으로 다시 추정했습니다. 첫 번째와 '
        '세 번째 연구 질문의 핵심 결과는 모든 재추정에서 유지되었고, 두 번째 질문의 시간 추세는 2014년 이전 출원이 30건뿐인 hybrid bonding 표본을 제외하고 '
        '모두 유지되었습니다.'))
    rows = [['표본 · 변수', 'N', 'RQ1\n기술범위 → 통합형', 'RQ2\n출원연도 → 공정 단독형', 'RQ3\n2020–24 상호작용'],
            ['기준', '928 / 763', '1.403***', '0.913***', '0.059***'],
            ['동일 제목 군집 표준오차', '928 / 763', '1.403***', '0.913***', '0.059***'],
            ['동일 제목 중복 제거 (제목당 1건)', '473 / 361', '1.295***', '0.935***', '0.131***'],
            ['좁은 제조공정 정의 (패키징 단계 공정 제외)', '908 / 763', '1.352***', '0.898***', '0.050***'],
            ['검색어의 무관 용법 68건 제외', '860 / 712', '1.384***', '0.820***', '0.185**'],
            ['hybrid bonding 표본 (제목 기준)', '467', '1.485***', '1.029', '추정하지 않음'],
            ['핵심 기호 외 접합 기호 수 통제', '763', '—', '—', '0.078***'],
            ['핵심 기술을 24/08 · 24/80으로 한정', '763', '—', '—', '0.051***']]
    table(X0, 2.65, [10.4, 3.6, 5.4, 6.1, 5.37], rows, [1.35] + [1.06] * 8, size=12.5, head_size=12, aligns=[L, C, C, C, C],
          warn=[(6, 3)], colors={(r, c): MUTED for r in (7, 8) for c in (2, 3)})
    text(X0, 13.7, CW, 2.7, [
        P(('RQ1 · RQ2: 다항 로짓 상대위험비 (N 왼쪽) · RQ3: 핵심 기술 모형의 2014년 이전 대비 배수 (N 오른쪽) · *** p < 0.01, ** p < 0.05', 11, MUTED), after=4),
        P(('hybrid bonding 표본은 2014년 이전 출원이 30건뿐이어서 시간 추세를 식별할 수 없고, 2019년 이전 핵심 기술 없는 접합 출원이 12건이라 RQ3 모형은 추정하지 않음', 12, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=3),
        P(('확산기의 승산비는 군집 표준오차와 무관 용법 제외 표본에서 유의성을 잃으므로, 우위가 사라진 시점(확산기/재편기)은 확정하지 않음', 12, TEXT), bullet=('■', NAVY, 70), indent=0.5)])


def s_res_summary(num='', title='Chapter 4  Summary', robust=False):
    CS(num, title, CH[4][0], take='공정 단계 간 융합은 국면에 따라 내용과 중심이 재편되는 동태적 과정이었다', script=(
        '결과를 하나의 그림으로 정리하면 이렇습니다. 형성기에는 제조공정 발명이 단독으로 많이 출원되며 접합에 필요한 지식이 쌓였고, 확산기에는 두 영역의 '
        '지식이 하나의 발명 안에서 결합되며 통합형이 정점에 이르렀으며, 재편기에는 핵심 설계가 고착화되면서 결합의 중심이 다이 지지와 봉지 같은 주변 공정으로 '
        '옮겨 갔습니다. 이론에서 말씀드린 지식 융합, 기술 융합, 응용 융합의 흐름과 맞닿아 있습니다.'))
    cols = [('형성기', '–2014', PROC, PROC_L, ['공정 단독형 43.6% (최다)', '11.07***', '53.5%', '웨이퍼 제조공정', '지식 융합']),
            ('확산기', '2015–2019', INTEG, INTEG_L, ['통합형 62.2% (정점)', '2.45**', '78.7%', '웨이퍼 제조공정\n(21/187 · 21/76898)', '기술 융합']),
            ('재편기', '2020–2024', BOND, BOND_L, ['접합 단독형 45.4%\n통합형 44.3%', '0.66 (비유의)', '88.0%', '공정 장치 · 패키징 단계 공정\n(21/6835 · 21/568)', '응용 융합으로 이행'])]
    labels = ['가장 두드러진 유형', '핵심 기술의 승산비', '접합 출원 중 핵심 기술 비중', '접합과 결합한 제조공정의 중심', '단계 모형과의 대응']
    lw_ = 6.2
    w, g = (CW - lw_ - 2 * 0.5) / 3, 0.5
    rh = [1.75, 1.45, 1.45, 2.05, 1.45]
    y0 = 4.6
    for i, lab in enumerate(labels):
        yy = y0 + sum(rh[:i]) + i * 0.18
        text(X0, yy, lw_ - 0.3, rh[i], [P((lab, 13, MUTED, True), align=R)], anchor='m')
    for k, (nm, yr, col, colL, vals) in enumerate(cols):
        x = X0 + lw_ + k * (w + g)
        sh = shape(MSO_SHAPE.PENTAGON if k < 2 else MSO_SHAPE.RECTANGLE, x, 2.7, w + (0.45 if k < 2 else 0), 1.55, fill=col)
        shape_text(sh, [P((nm, 18, WHITE, True), (f'  {yr}', 12.5, WHITE), align=C)])
        card(x, y0 - 0.1, w, sum(rh) + 5 * 0.18 + 0.05, fill=colL)
        for i, v in enumerate(vals):
            yy = y0 + sum(rh[:i]) + i * 0.18
            bold = i in (0, 1, 4)
            text(x + 0.3, yy, w - 0.6, rh[i], [P((ln_, 14.5 if bold else 13.5, col if i == 4 else TEXT, bold), align=C) for ln_ in v.split('\n')], anchor='m')
            if i < 4: line(x + 0.6, yy + rh[i] + 0.09, x + w - 0.6, yy + rh[i] + 0.09, 'FFFFFF', 1.5)
    if robust:
        text(X0, 14.25, CW, 1.5, [P(('강건성: RQ1 · RQ3 결과는 여섯 가지 재추정에서 모두, RQ2의 시간 추세는 2014년 이전 출원이 30건뿐인 hybrid bonding 표본을 제외하고 모두 유지', 13, NAVY, True), align=C)], anchor='m')


# ================================================================ 05 결론
def s_conclusion():
    CS('5.1', '연구 결과 요약 및 이론적 함의', CH[5][0], script=(
        '연구 결과를 종합해 보면, 첫째, 기술융합은 서로 다른 산업 사이뿐 아니라 한 산업의 공정 단계 사이에서도 개별 발명으로 나타났습니다. 둘째, 그 융합은 '
        '한 방향으로 쌓이지 않고 형성, 확산, 재편의 국면을 거쳤습니다. 셋째, 핵심 설계가 고착화되면서 융합의 중심이 핵심 기술에서 주변 공정으로 이동했습니다. '
        '특히 재편기에 연결 역할을 한 기호가 공정 장치로 바뀐 것은, 공정 지식의 수렴이 공작기계라는 장비를 매개로 확산되었다는 Rosenberg의 관찰과도 맥을 '
        '같이합니다. 이것이 이 연구에서 가장 핵심이 되는 발견이라고 생각합니다.'))
    cells = [(1, '통합형 45.8% (좁은 정의 37.2%)\n세부 기술이 다양할수록 통합형', '기술융합 연구의 범위를 산업 간에서\n공정 단계 간으로 확장'),
             (2, '형성기 → 확산기 → 재편기\n출원의 무게중심: 제조공정 → 접합', '융합은 정태적 결합이 아닌\n유형별 구성이 바뀌는 동태적 과정'),
             (3, '핵심 기술의 우위: 11.1 → 0.66 (비유의)\n결합 중심: 웨이퍼 제조공정 → 장치 · 패키징', '핵심–주변 구조와 기술 순환으로\n융합의 전개를 설명')]
    hw, hh = (CW - 0.4) / 2, 6.95
    pos = [(X0, 2.6), (X0 + hw + 0.4, 2.6), (X0, 2.6 + hh + 0.4)]
    for (n, res, imp), (x, y) in zip(cells, pos):
        card(x, y, hw, hh, fill=PANEL)
        text(x + 0.6, y + 0.35, hw - 1.2, 0.9, [P((f'Research Question {n}', 18, NAVY, True), (f'   {RQ_SHORT[n]}', 12.5, MUTED))], anchor='m')
        text(x + 0.6, y + 1.35, hw - 1.2, 2.3, [P((ln_, 14.5, TEXT, True)) for ln_ in res.split('\n')], anchor='m')
        line(x + 0.6, y + 3.95, x + hw - 0.6, y + 3.95, RULE, 1.0)
        chip(x + 0.6, y + 4.3, 3.2, 0.85, '이론적 함의', fill=HL, color=NAVY, size=12)
        text(x + 0.6, y + 5.25, hw - 1.2, 1.55, [P((ln_, 14, NAVY, True)) for ln_ in imp.split('\n')], anchor='m')
    x, y = pos[2][0] + hw + 0.4, pos[2][1]
    card(x, y, hw, hh, fill=NAVY)
    text(x + 0.6, y + 0.35, hw - 1.2, 0.9, [P(('결론', 18, 'D4A72C', True))], anchor='m')
    text(x + 0.6, y + 1.35, hw - 1.2, 5.3, [
        P(('공정 단계 간 기술융합은 한 방향으로 축적된 과정이라기보다, 핵심 기술 설계의 고착화와 함께 융합의 중심이 핵심에서 주변으로 이동한 동태적 과정이었다', 15.5, WHITE, True), line=1.15, after=8),
        P(('공정 장치가 연결의 중심에 선 것은 장비를 매개로 한 공정 지식의 확산(Rosenberg, 1963)과도 맥을 같이함', 12, DIM), line=1.1)], anchor='t')


def s_practice():
    CS('5.2', '실무적 · 정책적 시사점', CH[5][0], take='패키징 경쟁력은 접합 원리 하나에 좌우되지 않을 수 있다', script=(
        '실무적으로는 두 가지 시사점이 있습니다. 첫째, 상당수의 출원이 제조공정과 접합 기술을 함께 담고 있다는 것은, 하이브리드 본딩에서 접합 기술 역량만으로는 '
        '충분하지 않고 평탄화, 오염 제어, 표면 화학, 정렬을 아우르는 인접 공정 지식이 중요해질 수 있음을 뜻합니다. 둘째, 재편기에는 경쟁의 초점이 핵심 원리의 '
        '발명에서 그 원리를 다이 단위 공정과 결합하는 응용의 발명으로 옮겨 가고 있을 가능성이 있습니다. 차세대 HBM에 하이브리드 본딩을 도입하려는 한국 메모리 '
        '산업에는, 제조공정 역량에 더해 얇은 다이를 다루고 적층하는 공정 역량까지 필요할 수 있다는 의미입니다.'))
    cw2 = (CW - 0.8) / 2
    items = [('FaBuilding', '기업: 인접 공정 지식의 확보', ['상당수 출원이 제조공정과 접합 기술을 함께 포함', '접합 기술 역량만으로는 충분하지 않을 수 있음',
              '평탄화 · 오염 제어 · 표면 화학 · 정렬을 아우르는 인접 공정 지식', '직접 생산하는 범위보다 넓은 지식 (Kapoor & Adner, 2012)']),
             ('FaMemory', '한국 메모리 산업: 차세대 HBM', ['재편기: 핵심 기술의 우위가 확인되지 않고, 결합 공정은 다이 지지 · 봉지로',
              '경쟁의 초점이 핵심 원리 → 다이 단위 공정과의 결합(응용)으로 이동 가능성', '평탄화 · 표면 화학 · 정렬 계측', '+ 얇은 다이를 다루고 적층하는 공정 역량'])]
    for k, (ic, head, pts) in enumerate(items):
        x = X0 + k * (cw2 + 0.8)
        card(x, 2.7, cw2, 13.8, fill=PANEL)
        icon_circle(ic, x + 0.6, 3.2, 2.2, NAVY)
        text(x + 3.2, 3.2, cw2 - 3.6, 2.2, [P((head, 19, NAVY, True))], anchor='m')
        paras = []
        for j, t in enumerate(pts):
            last = j == len(pts) - 1
            paras.append(P((t, 16, NAVY if last else TEXT, last), bullet=('■', BOND if last else NAVY, 70), indent=0.6, after=20, line=1.15))
        text(x + 0.7, 6.1, cw2 - 1.4, 10.0, paras)


def s_limits():
    CS('5.3', '한계점과 향후 연구 방향', CH[5][0], take='다른 공정기술 사례와 비교하여 형성–확산–재편의 일반성을 검증하는 것이 다음 과제이다', script=(
        '이 연구에는 한계도 있습니다. 첫째, 공동분류는 융합의 대리측정치이고, CPC가 과거 출원까지 소급해 재분류될 수 있어 시간적 변화의 일부는 분류 체계의 '
        '변화를 반영했을 수 있습니다. 둘째, 제목 키워드로 검색했기 때문에 제목에 관련어가 없는 출원은 빠졌습니다. 셋째, 특허 패밀리 정보가 없어 같은 제목으로 '
        '중복을 근사했습니다. 넷째, 출원인이나 응용처 같은 변수가 없고, 결과는 인과관계가 아니라 조건부 연관입니다. 향후에는 출원인과 패밀리 정보를 결합하고, '
        '디스플레이나 배터리처럼 하류 단계가 상류 역량에 의존하게 된 다른 공정기술과 비교하여 이 세 국면이 일반적인 현상인지 검증해 보고자 합니다.'))
    text(X0, 2.5, 8.0, 0.9, [P(('한계', 17, NAVY, True))], anchor='m')
    lim = [('FaTags', 'Measurement', '공동분류는 대리측정치\nCPC 소급 재분류 가능성'), ('FaFilter', 'Sampling', '제목 키워드 검색\n관련어 없는 출원 누락'),
           ('FaCopy', 'Patent Family', '패밀리 식별자 없음\n동일 제목으로 근사'), ('FaScaleBalanced', 'Causality', '조건부 연관 · 누락변수\n시기 = 지배적 설계의 근사')]
    w, g = 7.2, 0.69
    for k, (ic, en, ko) in enumerate(lim):
        x = X0 + k * (w + g)
        card(x, 3.45, w, 5.6, fill=PANEL)
        icon(ic, x + w / 2 - 0.85, 3.85, 1.7, NAVY)
        text(x, 5.7, w, 0.8, [P((en, 14.5, NAVY, True), align=C)], anchor='m')
        text(x + 0.3, 6.55, w - 0.6, 2.3, [P((ln_, 13, TEXT), align=C) for ln_ in ko.split('\n')], anchor='m')
    text(X0, 9.45, 8.0, 0.9, [P(('향후 연구', 17, NAVY, True))], anchor='m')
    fut = [('FaUsers', 'Applicant & Family', '출원인 · 패밀리 정보를 결합한 재분석'),
           ('FaShuffle', 'Other Industries', '디스플레이 리소그래피 · 배터리 셀-투-팩 · 항공우주 적층제조와 비교'),
           ('FaCircleNodes', 'Generality', '고차 공출현 · 일반성 지수로 범용기술적 성격 검정 (Petralia, 2020)')]
    w, g = (CW - 2 * 0.69) / 3, 0.69
    for k, (ic, en, ko) in enumerate(fut):
        x = X0 + k * (w + g)
        card(x, 10.4, w, 6.05, fill=PANEL2)
        icon(ic, x + 0.55, 10.85, 1.5, NAVY)
        text(x + 2.4, 10.8, w - 2.8, 1.6, [P((en, 15, NAVY, True))], anchor='m')
        text(x + 0.55, 12.8, w - 1.1, 3.4, [P((ko, 14, TEXT), line=1.12)], anchor='t')


def s_close():
    image_slide('wafer.png', '마무리', 'Q&A', script=SCRIPT or (
        '이상으로 발표를 마치겠습니다. 경청해 주셔서 감사합니다. 질문 주시면 성실히 답변드리겠습니다.'))
    sh = K.s.shapes.build_freeform(cm(12.0), cm(0))
    sh.add_line_segments([(cm(SW), cm(0)), (cm(SW), cm(SH)), (cm(6.0), cm(SH)), (cm(12.0), cm(0))])
    ff = sh.convert_to_shape(); K._nostyle(ff)
    ff.fill.solid(); ff.fill.fore_color.rgb = rgb('05070B'); ff.line.fill.background()
    shape(MSO_SHAPE.RECTANGLE, 9.0, 4.2, 20.0, 10.6, line=WHITE, lw=2.5)
    text(9.0, 5.2, 20.0, 3.2, [P(('Q&A', 54, WHITE, True), align=C)], anchor='m')
    text(9.0, 8.4, 20.0, 1.6, [P(('Thanks for Listening', 28, WHITE, True), align=C)], anchor='m')
    text(9.0, 10.6, 20.0, 1.2, [P(('경청해 주셔서 감사합니다', 16, DIM), align=C)], anchor='m')
    text(9.0, 12.2, 20.0, 1.0, [P(('이형규 · 한양대학교 기술경영전문대학원 기술경영학과', 13, DIM), align=C)], anchor='m')


# ================================================================ 부록 (질의응답용)
def appendix_head(n, title, script):
    return CS(f'부록 {n}', title, '부록', script=script, title_size=20)


def s_app_data():
    appendix_head(1, '검색 전략과 자료의 특성', '검색 전략과 표본의 특성을 묻는 질문에 쓰는 슬라이드입니다. 1차로 CPC 두 클래스, 2차로 제목 키워드를 적용했고, '
                  'hybrid bonding 표본 467건, 무관 용법 68건, 동일 제목 473개를 강건성 검증에 사용했습니다.')
    rows = [['단계', '조건', '내용'],
            ['1차: CPC', 'H01L21', '반도체 제조공정. 웨이퍼 직접 접합(21/18, 21/20), 소자 분리·배선·TSV(21/76), 기판·절연층 처리(21/02), CMP(21/30–21/32);\n패키징 단계 제조공정(21/48, 21/50–21/58, 21/60, 21/78)과 공정 장치(21/67–21/68) 포함'],
            ['1차: CPC', 'H01L24', '반도체 · 고체소자의 연결을 위한 커넥터 구조, 범프 · 층 · 와이어 연결, 관련 연결 방법 및 장치'],
            ['2차: 키워드', '제목', 'hybrid bonding · direct bonding · Cu–Cu · copper to copper · metal-oxide bond 가운데 하나 이상']]
    table(X0, 2.65, [3.4, 2.6, CW - 6.0], rows, [0.9, 2.0, 1.15, 1.15], size=12, head_size=12, aligns=[C, C, L], bold_cols=(0, 1))
    facts = [('표본', '출원 928건 (1968–2024) · 출원–CPC 쌍 5,277건 · 서브그룹 기호 262개 (H01L21 202개, H01L24 60개)'),
             ('hybrid bonding 표본', '제목에 hybrid bonding이 포함된 467건 → 재추정 (부록 6)'),
             ('무관 용법', '68건 (direct bonding 63건 · hybrid bonding 5건): DBC 기판, 구리 합금, 전력 모듈 등'),
             ('동일 제목', '고유 제목 473개 · 이 중 149개는 둘 이상의 출원청에 출원 → 제목당 1건 표본, 군집 표준오차'),
             ('H01L21 기호 1,636건', '웨이퍼 제조공정 905 (55.3%) · 공정 장치 328 (20.0%) · 패키징 단계 공정 403 (24.6%)'),
             ('출원청', 'US 382 · CN 175 · WO(PCT) 113 · EP 83 · TW 75 · KR 58 · JP 26 · 기타 16')]
    for k, (h, t) in enumerate(facts):
        yy = 8.25 + k * 1.38
        chip(X0, yy, 5.2, 1.08, h, fill=NAVY, color=WHITE, size=12)
        text(X0 + 5.5, yy, CW - 5.5, 1.08, [P((t, 13, TEXT))], anchor='m')


def s_app_vars():
    appendix_head(2, '변수 정의와 기술통계', '변수 정의를 묻는 질문에 쓰는 슬라이드입니다. 기술범위는 평균 5.69개이며, 통합형 7.89개, 접합 단독형 4.16개, '
                  '공정 단독형 3.15개입니다. 유형 최소치 조정 범위는 탐색적 강건성 지표로만 사용했습니다.')
    rows = [['변수', '정의', '유형', '해석상 주의'],
            ['breadth', 'H01L21 · H01L24 안의 서로 다른 서브그룹 수', '카운트, 1–21', '두 클래스 내부의 기술적 다양성 (법적 청구범위 아님)'],
            ['breadth_adj', '보유 클래스마다 최소 1개 기호를 뺀 추가 서브그룹 수', '카운트, 0–19', '탐색적 지표 (유형 정보 이용)'],
            ['tech_cat', '접합 단독형 / 통합형 / 공정 단독형', '범주, 3', '분류의 조합'],
            ['has_process', 'H01L21 기호를 하나라도 가지면 1', '이항', '전공정 자체가 아닌 제조공정 분류'],
            ['yc', '출원연도 (1968년 = 0), 평균 49.7 (2017.7년)', '정수, 0–56', '기술 성숙의 직접 측정치 아님'],
            ['period', '–2014(기준) · 2015–2019 · 2020–2024', '범주, 3', '첫 시기가 이후 두 시기보다 김'],
            ['core', '핵심 접합 기호(24/03 · 05 · 08 · 80) 보유 시 1', '이항', 'H01L24 보유 출원(N = 763)에서만 정의'],
            ['ccode', 'US(기준) · CN · KR · TW · JP · EP · WO · 기타', '범주, 8', '통제변수 (해석하지 않음)']]
    table(X0, 2.65, [3.9, 11.6, 3.4, CW - 18.9], rows, [0.9] + [1.08] * 8, size=11.5, head_size=12, aligns=[C, L, C, L], bold_cols=(0,))
    text(X0, 12.65, CW, 2.6, [
        P(('기술범위: 평균 5.69 (표준편차 4.07, 범위 1–21) · 통합형 7.89 · 접합 단독형 4.16 · 공정 단독형 3.15', 13, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=4),
        P(('유형: 통합형 425 (45.8%) · 접합 단독형 338 (36.4%) · 공정 단독형 165 (17.8%) · 제조공정 기호 보유 590 (63.6%)', 13, TEXT), bullet=('■', NAVY, 70), indent=0.5)])


def s_app_cluster():
    appendix_head(3, '기술 클러스터와 기술 네트워크', '핵심 기술의 근거를 묻는 질문에 쓰는 슬라이드입니다. 클러스터 상위 기호와 네트워크의 가장 강한 기호 쌍이 모두 '
                  '접합 패드와 직접 접합 기호로 이루어져 있습니다. 인덱싱 코드는 회귀 자료에 없으므로 핵심 기술은 회귀 자료의 H01L24 기호 네 개로 정의했습니다.')
    rows = [['순위', 'CPC 기호', '공식 정의(요약)', '출원 수', '비중'],
            ['1', '24/80', '접합 패드를 이용한 연결 방법', '529', '57.0%'],
            ['2', '24/08', '연결 후 개별 접합 패드의 구조', '513', '55.3%'],
            ['3', '2224/80895', '도전성 표면끼리의 직접 접합 (구리–구리)', '504', '54.3%'],
            ['4', '2224/80896', '절연성 표면끼리의 직접 접합 (산화막)', '488', '52.6%'],
            ['5', '2224/08145', '범프 없는 적층 칩 사이의 접합', '436', '47.0%'],
            ['6', '25/50', '소자 조립체의 다단계 제조공정', '363', '39.1%'],
            ['7', '24/05', '연결 전 개별 접합 패드의 구조', '350', '37.7%'],
            ['8', '25/0657', '소자의 적층 배치', '274', '29.5%'],
            ['9', '2224/05647', '구리를 주성분으로 하는 접합 패드', '239', '25.8%'],
            ['10', '24/03', '접합 패드의 제조 방법', '222', '23.9%'],
            ['11', '2224/80357', '표면과 같은 높이의 접합 계면', '205', '22.1%']]
    hl = [(i, 1) for i in (1, 2, 7, 10)]
    table(X0, 2.65, [1.4, 3.3, 9.0, 2.0, 2.0], rows, [0.85] + [0.98] * 11, size=11.5, head_size=11.5, aligns=[C, C, L, C, C], hl=hl)
    rx = X0 + 18.4
    rw = X1 - rx
    rows2 = [['기호 쌍 (상위 5)', '출원 수'],
             ['2224/80895 – 2224/80896', '460'],
             ['24/08 – 24/80', '439'],
             ['24/80 – 2224/80895', '434'],
             ['24/08 – 2224/80895', '420'],
             ['24/80 – 2224/80896', '418']]
    table(rx, 2.65, [rw - 2.6, 2.6], rows2, [0.85] + [0.98] * 5, size=11.5, head_size=11.5, aligns=[L, C])
    text(rx, 8.9, rw, 5.6, [
        P(('강조한 네 기호 = 회귀 자료의 핵심 기술', 12.5, NAVY, True), after=5),
        P(('회귀 자료 네트워크의 연결 강도: 24/80 3,235 · 24/08 3,140 · 24/05 2,518 · 24/03 1,778 (상위 4개)', 12, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=4),
        P(('네 기호의 6개 쌍(186–439건)이 가장 강한 6개 연결', 12, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=4),
        P(('다섯째 24/83(층 형태 커넥터)은 164건으로 기준 미달', 12, TEXT), bullet=('■', NAVY, 70), indent=0.5)])
    text(X0, 15.0, 18.0, 0.6, [P(('N = 928 · 기호 앞 H01L 생략 · 2224, 25 계열은 같은 출원의 전체 CPC를 별도 집계', 10, MUTED))], anchor='m')


def s_app_logit():
    appendix_head(4, '이항 · 다항 로짓 추정 결과', '회귀 결과의 세부를 묻는 질문에 쓰는 슬라이드입니다. 이항 로짓의 AUC는 0.707로 수용 가능한 수준의 하단이며, '
                  '다항 로짓의 McFadden 의사 R²는 0.192입니다.')
    hw = (CW - 1.0) / 2
    text(X0, 2.6, hw, 0.8, [P(('제조공정 분류 여부의 이항 로짓 (종속변수 has_process)', 14, NAVY, True))], anchor='m')
    rows = [['', '승산비', '로그 계수 SE', 'p'], ['기술범위', '1.267***', '(0.027)', '<0.001'], ['출원연도(yc)', '0.946***', '(0.011)', '<0.001']]
    table(X0, 3.5, [4.3, 3.6, 3.6, hw - 11.5], rows, [0.95, 1.05, 1.05], size=13, aligns=[L, C, C, C], bold_cols=(0,), hl=[(1, 1)])
    text(X0, 6.75, hw, 2.2, [P(('N = 928 · 출원청 더미 7개 포함 · 로그우도 −547.6 · 우도비 χ²(9) = 122.0 (p < 0.001) · McFadden 의사 R² = 0.100 · AIC 1,115.1 · AUC 0.707', 12, MUTED), line=1.1)])
    x2 = X0 + hw + 1.0
    text(x2, 2.6, hw, 0.8, [P(('공동분류 유형의 다항 로짓 (기준 = 접합 단독형)', 14, NAVY, True))], anchor='m')
    rows = [['', '통합형\nRRR', 'SE', '공정 단독형\nRRR', 'SE'],
            ['기술범위', '1.403***', '(0.031)', '0.905**', '(0.048)'],
            ['출원연도(yc)', '0.993', '(0.015)', '0.913***', '(0.015)']]
    table(x2, 3.5, [3.6, 2.9, 2.4, 3.3, hw - 12.2], rows, [1.3, 1.05, 1.05], size=13, head_size=11.5, aligns=[L, C, C, C, C], bold_cols=(0,), hl=[(1, 1), (2, 3)])
    text(x2, 6.75, hw, 2.2, [P(('N = 928 · 출원청 더미 7개 포함 · 로그우도 −774.5 · 우도비 χ²(18) = 367.4 (p < 0.001) · McFadden 의사 R² = 0.192 · AIC 1,589.0 · p: 0.654(통합형 연도), 0.039(공정 단독형 범위)', 12, MUTED), line=1.1)])
    card(X0, 9.4, CW, 6.3, fill=PANEL)
    text(X0 + 0.6, 9.6, CW - 1.2, 0.8, [P(('해석 원칙', 14, NAVY, True))], anchor='m')
    text(X0 + 0.6, 10.5, CW - 1.2, 5.0, [
        P(('OR은 승산의 배율, RRR은 기준 범주 대비 상대위험의 배율: 확률이 같은 비율로 변한다는 뜻이 아님 → 평균 예측확률을 함께 제시', 13, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=5),
        P(('세 범주는 크기 순서가 아닌 서로 다른 CPC 구성 → 순서형이 아닌 다항 로짓 (McFadden, 1974)', 13, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=5),
        P(('5% 유의수준을 쓰되 단일 p값이 아니라 대안 표본 · 정의 · 군집 표준오차에서의 방향 · 크기 · 유의성의 일관성으로 판단', 13, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=5),
        P(('모든 계수는 인과효과가 아닌, 관찰된 변수를 통제한 조건부 연관', 13, TEXT), bullet=('■', NAVY, 70), indent=0.5)])


def s_app_core():
    appendix_head(5, '핵심 기술 모형의 세부 결과', '핵심 기술 모형의 세부입니다. 두 상호작용의 결합 Wald 검정은 χ²(2) 26.43으로 유의합니다. 변화는 핵심 기술을 '
                  '포함한 출원에서 비롯되었고, 핵심 기술이 없는 접합 출원의 통합형 비중은 오히려 높아졌습니다.')
    rows = [['출원 시기', '핵심 기술 보유\n통합형/N', '핵심 기술 없음\n통합형/N', '승산비', '95% CI', 'p', '2014년 이전\n대비 배수'],
            ['–2014', '43/54 (79.6%)', '13/47 (27.7%)', '11.071***', '4.211–29.107', '<0.001', '—'],
            ['2015–2019', '108/140 (77.1%)', '22/38 (57.9%)', '2.452**', '1.158–5.194', '0.019', '0.221 (p = 0.016)'],
            ['2020–2024', '206/426 (48.4%)', '33/58 (56.9%)', '0.658', '0.374–1.160', '0.148', '0.059 (p < 0.001)']]
    table(X0, 2.65, [3.4, 4.6, 4.6, 3.3, 4.6, 2.6, CW - 23.1], rows, [1.35, 1.1, 1.1, 1.1], size=12.5, head_size=12, aligns=[C] * 7, bold_cols=(0,), hl=[(1, 3), (3, 6)])
    text(X0, 7.5, CW, 0.7, [P(('N = 763 (H01L24 보유) · 출원청 통제 · 이분산 강건 표준오차 · 결합 Wald χ²(2) = 26.43 (p < 0.001) · McFadden 의사 R² = 0.069', 11.5, MUTED))], anchor='m')
    cw2 = (CW - 0.8) / 2
    for k, (h, pts) in enumerate([('변화는 핵심 기술 보유 출원에서', ['핵심 기술 보유: 통합형 79.6% → 77.1% → 48.4%', '핵심 기술 없음: 통합형 27.7% → 57.9% → 56.9% (오히려 상승)', '접합 출원 중 핵심 기술 비중 53.5% → 78.7% → 88.0%']),
                                  ('추정식', ['log[P/(1 − P)] = β₀ + β₁·core + Σδₜ·periodₜ + Σθₜ·(core × periodₜ) + Σγₖ·ccodeₖ', '시기별 승산비 = exp(β₁ + θₜ), 2014년 이전 대비 배수 = exp(θₜ)', '형성기의 해당 출원 43건은 소수의 동일 제목에 집중 → 대표 기호 제시하지 않음'])]):
        x = X0 + k * (cw2 + 0.8)
        card(x, 8.6, cw2, 7.4, fill=PANEL)
        text(x + 0.5, 8.8, cw2 - 1.0, 0.8, [P((h, 14, NAVY, True))], anchor='m')
        text(x + 0.5, 9.8, cw2 - 1.0, 6.0, [P((t, 13, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=8, line=1.1) for t in pts])


def s_app_robust1():
    appendix_head(6, '강건성: 다항 로짓 재추정', '다항 로짓 재추정의 세부입니다. 군집 표준오차에서 유의성을 잃은 계수는 기술범위가 공정 단독형에 미치는 효과 하나뿐이며, '
                  '유형 최소치 조정 범위를 쓰면 통합형의 상대위험비는 1.262로 줄지만 방향은 유지됩니다.')
    rows = [['표본 · 변수', 'N', '범위 → 통합형', '연도 → 공정 단독형'],
            ['기준', '928', '1.403***', '0.913***'],
            ['기준 + 동일 제목 군집 표준오차', '928', '1.403***', '0.913***'],
            ['동일 제목 중복 제거 (제목당 1건)', '473', '1.295***', '0.935***'],
            ['좁은 제조공정 정의 (패키징 단계 공정 제외)', '908', '1.352***', '0.898***'],
            ['무관 용법 68건 제외', '860', '1.384***', '0.820***'],
            ['hybrid bonding 표본 (제목 기준)', '467', '1.485***', '1.029']]
    table(X0, 2.65, [11.0, 2.4, 4.4, 5.4], rows, [0.95] + [1.02] * 6, size=12.5, head_size=12, aligns=[L, C, C, C], warn=[(6, 3)])
    rx = X0 + 24.0
    rw = X1 - rx
    text(rx, 2.65, rw, 9.0, [
        P(('군집 표준오차: 범위 → 공정 단독형만 유의성 상실 (RRR 0.905, p = 0.29)', 12.5, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=6, line=1.1),
        P(('breadth_adj: 통합형 RRR 1.403 → 1.262, 방향 유지', 12.5, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=6, line=1.1),
        P(('2010년 이후(N = 830): 통합형 연도 RRR 0.934 (p < 0.05)', 12.5, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=6, line=1.1),
        P(('좁은 정의 유형 분포: 접합 단독형 46.8% · 통합형 37.2% · 공정 단독형 16.0% → 최다 유형은 정의에 의존', 12.5, TEXT), bullet=('■', NAVY, 70), indent=0.5, line=1.1)])
    card(X0, 10.4, CW, 5.6, fill=PANEL)
    text(X0 + 0.6, 10.6, CW - 1.2, 0.8, [P(('표본별 유형 분포', 14, NAVY, True))], anchor='m')
    rows = [['표본', '통합형', '접합 단독형', '공정 단독형'],
            ['전체 (928)', '45.8%', '36.4%', '17.8%'],
            ['무관 용법 제외 (860)', '47.4%', '35.3%', '17.2%'],
            ['hybrid bonding (467)', '49.5%', '38.3%', '12.2%'],
            ['좁은 제조공정 정의 (908)', '37.2%', '46.8%', '16.0%']]
    table(X0 + 0.6, 11.5, [8.0, 4.0, 4.0, 4.0], rows, [0.8] * 5, size=12, head_size=11.5, aligns=[L, C, C, C])
    text(X0 + 21.4, 11.5, CW - 22.0, 4.2, [P(('hybrid bonding 표본은 2014년 이전 출원이 30건뿐: 공정 단독형의 감소는 direct bonding 시기의 웨이퍼 접합 공정 출원에서 hybrid bonding 시기의 구조 출원으로 이행한 변화를 상당 부분 반영', 12, MUTED), line=1.1)])


def s_app_robust2():
    appendix_head(7, '강건성: 핵심 기술 모형 재추정', '핵심 기술 모형 재추정의 세부입니다. 2014년 이전 승산비가 1보다 크고, 2020–2024년 승산비가 1과 유의하게 '
                  '다르지 않으며, 그 차이가 유의하게 1보다 작다는 결과는 일곱 가지 추정에서 모두 유지되었습니다. 가장 약한 경우는 무관 용법 제외 표본입니다.')
    rows = [['표본 · 변수', 'N', '–2014', '2015–2019', '2020–2024', '2020–24 상호작용'],
            ['기준', '763', '11.071***', '2.452**', '0.658', '0.059***'],
            ['핵심 기호 외 접합 기호 수 통제', '763', '8.768***', '2.377**', '0.684', '0.078***'],
            ['기준 + 동일 제목 군집 표준오차', '763', '11.071***', '2.452', '0.658', '0.059***'],
            ['동일 제목 중복 제거 (제목당 1건)', '361', '6.191***', '4.433***', '0.814', '0.131***'],
            ['좁은 제조공정 정의 (패키징 단계 공정 제외)', '763', '15.334***', '3.185***', '0.771', '0.050***'],
            ['무관 용법 제외', '712', '3.981**', '1.817', '0.738', '0.185**'],
            ['핵심 기술을 24/08 · 24/80으로 한정', '763', '14.474***', '3.445***', '0.732', '0.051***']]
    table(X0, 2.65, [11.0, 2.4, 4.0, 4.0, 4.0, CW - 25.4], rows, [0.95] + [1.02] * 7, size=12.5, head_size=12, aligns=[L, C, C, C, C, C],
          warn=[(3, 3), (6, 3)], hl=[(i, 5) for i in range(1, 8)])
    text(X0, 11.4, CW, 4.6, [
        P(('핵심 기호 외 접합 기호 수 자체의 승산비 1.100 (p < 0.01) → 통제해도 시기별 승산비는 거의 변하지 않음', 13, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=5),
        P(('무관 용법 제외 표본: 무관 출원 51건 중 28건이 2014년 이전 핵심 기술 없는 출원(그중 25건 접합 단독형) → 2014년 이전 승산비 3.981로 작아짐', 13, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=5),
        P(('좁은 정의: 제조공정 기호가 패키징 단계 공정뿐인 통합형을 접합 단독형으로 재분류 (통합형 338건)', 13, TEXT), bullet=('■', NAVY, 70), indent=0.5, after=5),
        P(('hybrid bonding 표본: 2019년 이전 핵심 기술 없는 접합 출원 12건 → 추정하지 않음', 13, TEXT), bullet=('■', NAVY, 70), indent=0.5)])


QA = [
    ('H01L21–H01L24 공동분류를 공정 단계 간 융합으로 볼 수 있는가?',
     '대리측정치로 사용했습니다. H01L21에는 패키징 단계 공정(24.6%)과 공정 장치(20.0%)도 포함되어, 패키징 단계 공정을 뺀 좁은 정의로 재추정했고 '
     '회귀 결론은 유지되었습니다. 다만 통합형이 최다 유형이라는 기술통계는 정의에 의존합니다(좁은 정의에서는 접합 단독형 46.8%가 최다).'),
    ('기술범위와 통합형의 관계는 정의상 당연한 것 아닌가?',
     '통합형은 정의상 기호가 2개 이상이라는 기계적 측면이 있습니다. 보유 클래스마다 1개씩 뺀 유형 최소치 조정 범위로 재추정하면 RRR이 1.403에서 1.262로 '
     '줄지만 방향은 유지됩니다. 이 관계는 세 국면 모두에서 유지되었고, 인과가 아닌 구성적 연관으로 해석합니다.'),
    ('CPC가 소급 재분류되면 시간적 변화가 분류 체계의 산물일 수 있지 않은가?',
     '그 가능성은 한계로 밝혔습니다. 결과는 분류 체계와 심사 관행의 변화까지 포함한 관찰 결과로 해석하며, 대안 표본과 정의에서 방향이 일관되게 '
     '유지되었다는 점을 함께 판단 근거로 삼았습니다.'),
    ('1968–2014년을 하나의 시기로 묶은 근거는?',
     '47년간의 출원(179건)이 이후 5년간(209건)보다 적을 만큼 드물어 출원 집중도를 고려했습니다. 시기는 기술 수명주기 단계를 근사하는 것이며(Lee et al., 2017), '
     '연도를 선형 항으로 넣은 모형에서도 같은 결론(공정 단독형 RRR 0.913)이었습니다.'),
    ('같은 발명이 여러 특허청에 출원되어 관측치가 독립적이지 않은 것 아닌가?',
     '패밀리 식별자가 없어 동일 제목으로 근사했습니다. 928건은 고유 제목 473개이며, 제목당 1건만 남긴 표본과 동일 제목 군집 표준오차 모두에서 '
     '핵심 결과는 유지되었습니다.'),
    ('핵심 기술 우위가 사라진 것이 지배적 설계 때문이라고 할 수 있는가?',
     '인과관계를 주장하지 않습니다. 시기 구분은 지배적 설계의 형성 시점을 직접 측정한 것이 아니라 출원 시기로 근사한 것이고, 확산기 승산비는 일부 '
     '재추정에서 유의성을 잃으므로 우위가 사라진 시점이 확산기인지 재편기인지는 확정하지 않았습니다.'),
    ('hybrid bonding 표본에서 연도 계수가 유의하지 않은 이유는?',
     '이 표본은 2014년 이전 출원이 30건뿐입니다. 공정 단독형의 감소는 direct bonding 시기의 웨이퍼 접합 공정 출원에서 hybrid bonding 시기의 구조 출원으로 '
     '이행한 변화를 상당 부분 반영하는데, 이 표본으로는 그 이전 시기를 충분히 관찰할 수 없습니다. 기술범위의 RRR(1.485)은 유의했습니다.'),
    ('다항 로짓의 IIA 가정은 점검했는가?',
     '정식 검정 대신 대안 제외 민감도 분석으로 점검했습니다. 공정 단독형을 제외한 표본에서 통합형 대 접합 단독형의 이항 로짓을 추정하여 '
     '실질적 해석이 유지되는지 확인했습니다(논문 부록 A).'),
    ('텍스트 마이닝이 아닌 CPC 공동분류를 쓴 이유는?',
     'CPC 공동분류는 재현 가능하지만 발명보다 늦게 부여됩니다. 텍스트 기반 방법은 더 이른 탐지가 가능한 대신 학습된 모형에 의존하므로, H01L21과 H01L24의 '
     '경계가 제도적으로 명확한 본 연구에서는 재현성을 우선했습니다(Preschitschek et al., 2013; Lei et al., 2019; Zhu & Motohashi, 2022).')]


def s_app_qa(k):
    items = QA[3 * k:3 * k + 3]
    appendix_head(8 + k, f'예상 질문 ({k + 1}/3)', '예상 질문과 답변 요지입니다. 질문을 받으면 해당 답변을 바탕으로 짧게 답하고, 필요하면 앞의 부록 표로 넘어가 수치를 보여 드립니다.\n\n'
                  + '\n\n'.join(f'Q. {q}\nA. {a}' for q, a in items))
    for j, (q, a) in enumerate(items):
        yy = 2.6 + j * 4.75
        card(X0, yy, CW, 4.45, fill=PANEL)
        chip(X0 + 0.5, yy + 0.4, 1.3, 0.95, 'Q', fill=NAVY, color=WHITE, size=14)
        text(X0 + 2.1, yy + 0.35, CW - 2.6, 1.05, [P((q, 15, NAVY, True))], anchor='m')
        chip(X0 + 0.5, yy + 1.65, 1.3, 0.95, 'A', fill=HL, color=NAVY, size=14)
        text(X0 + 2.1, yy + 1.55, CW - 2.6, 2.75, [P((a, 12.5, TEXT), line=1.1)], anchor='t')


# ================================================================ merged slides (both versions) and 15-minute slides
def s_th_conv_stage():
    CS('2.1', '기술융합과 융합의 단계 모형', CH[2][0], take='본 연구가 특허 공동분류로 측정하는 것은 "기술 융합" 단계이다', script=(
        '먼저 기술융합의 개념입니다. Rosenberg는 19세기 총기와 재봉틀을 위해 개발된 가공 기술이 공통의 생산공정을 통해 다른 산업으로 퍼지고, 그 공정 '
        '지식을 축적한 공작기계 산업이 등장한 과정을 기술적 수렴이라고 불렀습니다. Kodama는 분리된 기술 분야가 결합하여 어느 한쪽만으로는 갖지 못한 '
        '역량을 만드는 과정을 기술융합으로 정의했습니다. Hacklin 등은 융합이 지식, 기술, 응용, 산업의 네 단계로 진행된다고 보았습니다. 이를 하이브리드 '
        '본딩에 대입하면, 평탄화 같은 제조공정 지식이 칩 연결 문제에 쓰이기 시작하는 것이 지식 융합, 그 지식이 접합 구조와 하나의 발명 안에서 결합되는 것이 '
        '기술 융합, HBM 같은 제품으로 구현되는 것이 응용 융합, 파운드리가 첨단 패키징을 내재화하는 것이 산업 융합입니다. 이 가운데 특허의 공동분류로 '
        '직접 관찰할 수 있는 것은 기술 융합 단계입니다.'))
    cols = [('기술적 수렴', 'Rosenberg (1963)', '공통의 생산공정을 통한 기술 확산 → 공작기계 산업의 등장'),
            ('기술융합', 'Kodama (1992)', '분리된 기술의 결합으로 어느 한쪽도 갖지 못한 역량 창출'),
            ('단계 모형', 'Hacklin et al. (2009, 2010)', '지식 → 기술 → 응용 → 산업 융합, 공진화적 순환')]
    w, g = 9.6, 1.0
    x0 = (SW - 3 * w - 2 * g) / 2
    for k, (ko, who, desc) in enumerate(cols):
        x = x0 + k * (w + g)
        obox(x, 2.6, w, 1.9, ko, who, size=17, sub_size=11, color=NAVY, fill=WHITE, lw=1.75)
        text(x, 4.6, w, 1.05, [P((desc, 12.5, TEXT), align=C)], anchor='m')
        if k < 2: arrow(x + w + 0.15, 3.55, x + w + g - 0.15, 3.55, NAVY, 2)
    line(X0, 6.05, X1, 6.05, RULE, 1.0)
    text(X0, 6.2, CW, 0.75, [P(('단계 모형을 하이브리드 본딩에 적용하면', 13.5, NAVY, True), align=C)], anchor='m')
    st = [('지식 융합', 'Knowledge', '제조공정 지식이 칩 연결 문제에 적용', '평탄화 · 표면 활성화 · 웨이퍼 직접 접합', '형성의 출발점', None),
          ('기술 융합', 'Technological', '그 지식이 접합 구조와 하나의 발명 안에서 결합', 'H01L21 + H01L24 공동분류', '특허 공동분류로 직접 관찰', 'hl'),
          ('응용 융합', 'Applicational', '결합된 기술이 제품으로 구현', '이미지센서 · 3D NAND · HBM · 칩렛', 'RQ3에서 간접적으로', None),
          ('산업 융합', 'Industrial', '공정 단계 간 분업 구조의 재편', '파운드리의 첨단 패키징 내재화', '분석 범위 밖', 'out')]
    w, g = 7.35, 0.42
    x0 = (SW - 4 * w - 3 * g) / 2
    for k, (ko, en, d1, d2, status, kind) in enumerate(st):
        x = x0 + k * (w + g)
        hlk = kind == 'hl'
        sh = shape(MSO_SHAPE.CHEVRON if k else MSO_SHAPE.PENTAGON, x, 7.15, w + 0.35, 1.85, fill=NAVY if hlk else PANEL2)
        try: sh.adjustments[0] = 0.28
        except IndexError: pass
        shape_text(sh, [P((ko, 17, WHITE if hlk else NAVY, True), align=C), P((en, 10.5, 'C8D3E6' if hlk else MUTED), align=C)])
        card(x + 0.1, 9.25, w - 0.2, 4.35, fill=PANEL if not hlk else 'EAF0FA', line_=NAVY if hlk else None)
        text(x + 0.4, 9.4, w - 0.8, 2.0, [P((d1, 14, TEXT, True), align=C)], anchor='m')
        line(x + 1.2, 11.5, x + w - 1.2, 11.5, RULE, 1.0)
        text(x + 0.4, 11.6, w - 0.8, 1.85, [P((d2, 12.5, MUTED), align=C)], anchor='m')
        fill_, col_ = {'hl': (HL, NAVY), 'out': ('EEF0F3', MUTED), None: (WHITE, NAVY)}[kind]
        chip(x + 0.4, 13.85, w - 0.8, 1.15, status, fill=fill_, color=col_, size=13, line_=None if kind else 'C9D2E0')
    source('Curran & Leker (2011); Lau (2022); KnowMade (2024)', y=15.55, align=R)


def s_res_rq1():
    CS('4.1', 'RQ1) 융합의 수준과 형태', CH[4][0], take='통합형이 가장 많았고, 세부 기술을 더 다양하게 포함한 출원일수록 통합형일 가능성이 높았다', script=(
        '이제 연구 결과를 말씀드리겠습니다. 첫 번째 연구 질문의 결과입니다. 928건 가운데 가장 많은 유형은 두 영역의 기호를 모두 가진 통합형으로, '
        '45.8%였습니다. 제조공정 분류에 패키징 단계 공정도 들어 있다는 점을 고려해 이를 뺀 좁은 정의를 적용해도 37.2%였습니다. 공정 단계 간 융합은 '
        '예외적인 현상이 아니었습니다. 그렇다면 어떤 특허가 경계를 넘었을까요? 세부 분류가 하나 늘 때마다 접합 단독형 대비 통합형일 상대위험이 40% '
        '높아졌습니다. 가운데 그래프처럼 기술범위가 3인 출원은 통합형일 확률이 0.26이지만, 8인 출원은 0.68입니다. 통합형은 정의상 기호가 두 개 이상이라는 '
        '기계적인 측면이 있어 이를 조정한 지표로도 확인했는데, 방향은 그대로 유지되었습니다.'))
    cd = CategoryChartData(); cd.categories = ['통합형', '접합 단독형', '공정 단독형']
    cd.add_series('유형', (425, 338, 165))
    gf = K.s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, cm(X0), cm(2.6), cm(9.6), cm(9.0), cd)
    ch = gf.chart; ch.has_title = False
    style_chart(ch, legend=False, size=12, axes=False)
    ser = ch.plots[0].series[0]
    for i, col in enumerate((INTEG, BOND, PROC)):
        pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(col)
        pt.format.line.color.rgb = rgb(WHITE); pt.format.line.width = Pt(2)
    dn = ch._chartSpace.find('.//' + qn('c:doughnutChart'))
    hs = dn.find(qn('c:holeSize'))
    if hs is None: hs = etree.SubElement(dn, qn('c:holeSize'))
    hs.set('val', '56')
    text(X0, 6.0, 9.6, 2.2, [P(('N = 928', 17, NAVY, True), align=C), P(('출원', 11, MUTED), align=C)], anchor='m')
    for k, (lab, n, pct, col) in enumerate([('통합형', 425, '45.8%', INTEG), ('접합 단독형', 338, '36.4%', BOND), ('공정 단독형', 165, '17.8%', PROC)]):
        yy = 11.8 + k * 1.3
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, X0 + 0.6, yy + 0.25, 0.5, 0.6, fill=col, radius=0.3)
        text(X0 + 1.35, yy, 3.6, 1.1, [P((lab, 13.5, col, True))], anchor='m')
        text(X0 + 4.9, yy, 4.7, 1.1, [P((pct, 17 if k == 0 else 15, TEXT, True), (f'  {n}건', 11, MUTED))], anchor='m')
    chip(X0 + 0.5, 15.65, 8.6, 0.95, '좁은 정의에서도 통합형 37.2%', fill=HL, color=NAVY, size=12.5)
    pb = D['pred_breadth']
    bs = list(range(1, 16))
    cd = CategoryChartData(); cd.categories = [str(b) for b in bs]
    cd.add_series('통합형', [round(pb[str(b)][1], 3) for b in bs])
    cd.add_series('접합 단독형', [round(pb[str(b)][0], 3) for b in bs])
    cd.add_series('공정 단독형', [round(pb[str(b)][2], 3) for b in bs])
    cx = X0 + 10.0
    gf = K.s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, cm(cx), cm(2.5), cm(12.3), cm(14.2), cd)
    ch = gf.chart
    style_chart(ch, legend=True, size=11.5)
    for ser, col in zip(ch.plots[0].series, (INTEG, BOND, PROC)):
        ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(3.25); ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.NONE
    integ = ch.plots[0].series[0]
    for idx in (2, 7):
        pt = integ.points[idx]
        pt.marker.style = XL_MARKER_STYLE.CIRCLE; pt.marker.size = 9
        pt.marker.format.fill.solid(); pt.marker.format.fill.fore_color.rgb = rgb(INTEG); pt.marker.format.line.color.rgb = rgb(INTEG)
        dl = pt.data_label; dl.has_text_frame = True
        dl.text_frame.text = '0.26' if idx == 2 else '0.68'
        dl.position = XL_LABEL_POSITION.ABOVE if idx == 7 else XL_LABEL_POSITION.BELOW
        for p in dl.text_frame.paragraphs:
            for r in p.runs: font_run(r, 13, INTEG, True)
    ch.value_axis.maximum_scale = 1.0; ch.value_axis.minimum_scale = 0; ch.value_axis.major_unit = 0.2
    ch.value_axis.tick_labels.number_format = '0.0'; ch.value_axis.tick_labels.number_format_is_linked = False
    axis_title(ch.value_axis, '예측확률', 11); axis_title(ch.category_axis, '기술범위 (CPC 서브그룹 수)', 11)
    rx = cx + 12.7
    rw = X1 - rx
    rq_panel(rx, 2.75, rw, 1, [
        P(('통합형 45.8%: 가장 큰 비중', 13.5, NAVY, True), after=7),
        P(('기술범위 → 통합형', 11.5, MUTED, True)), P(('RRR 1.403*** (서브그룹 1개당 +40%)', 13, TEXT, True), after=6),
        P(('예측확률 (통합형)', 11.5, MUTED, True)), P(('범위 3: 0.26 → 범위 8: 0.68', 13, TEXT, True), after=6),
        P(('최소 기호 수를 조정해도', 11.5, MUTED, True)), P(('RRR 1.262 (방향 유지)', 13, TEXT, True))], h=13.85, qh=2.9)


def s_rqs_15():
    CS('3.1', '연구 질문과 이론적 기대', CH[3][0], take='이론적 기대는 통계적 가설이 아니라, 결과를 해석하는 기준으로 사용한다', script=(
        '이 이론들로부터 세 가지 연구 질문을 세웠습니다. 첫째, 융합이 어느 정도로, 어떤 형태로 나타나는가. 둘째, 시간에 따라 어떻게 변화해 왔는가. '
        '셋째, 융합의 중심이 핵심 기술과 주변 기술 사이에서 어떻게 재편되는가입니다. 오른쪽은 각 질문에 대한 이론적 기대인데, 통계적 가설로 검정하는 것이 '
        '아니라 결과를 해석하는 기준으로 사용했습니다.'))
    rows = [(1, '인터페이스가 불안정해진 공정 단계 사이에서 두 영역을 함께 포함하는 발명이 예외적이지 않은 비중으로 나타나고, 더 다양한 세부 기술을 포함한다', '지식 경계 · 산업 아키텍처 · 재조합'),
            (2, '융합은 단선적으로 확대되기보다 발전 단계에 따라 수준이 달라지고, 출원의 구성은 제조공정 중심에서 접합 중심으로 이동한다', '단계 모형 · 기술 순환 모형'),
            (3, '핵심 설계가 고착화되면 핵심 기술의 제조공정 지식 포함 우위가 줄고, 융합의 중심은 응용에 필요한 주변 공정으로 이동한다', '핵심–주변 구조 · 아키텍처 지식')]
    text(X0 + 15.4, 2.45, 15.0, 0.6, [P(('이론적 기대', 12.5, MUTED, True))], anchor='m')
    for k, (n, exp, basis) in enumerate(rows):
        y = 3.1 + k * 4.55
        card(X0, y, CW, 4.2, fill=PANEL)
        chip(X0 + 0.45, y + 0.4, 2.0, 0.95, f'RQ{n}', fill=NAVY, color=WHITE, size=14)
        text(X0 + 2.75, y + 0.4, 11.8, 0.95, [P((RQ_SHORT[n], 14, NAVY, True))], anchor='m')
        text(X0 + 0.45, y + 1.5, 14.1, 2.5, [P((RQ[n], 14, TEXT, True), line=1.1)], anchor='m')
        shape(MSO_SHAPE.RIGHT_ARROW, X0 + 14.75, y + 1.6, 0.8, 1.0, fill='C9D2E0')
        text(X0 + 15.8, y + 0.3, CW - 16.2, 2.9, [P((exp, 13, NAVY), line=1.1)], anchor='m')
        text(X0 + 15.8, y + 3.2, CW - 16.2, 0.8, [P((f'근거: {basis}', 11, MUTED))], anchor='m')


def s_impl_limits_15():
    CS('5.2', '시사점과 한계', CH[5][0], take='다른 공정기술과 비교하여 형성–확산–재편의 일반성을 검증하는 것이 다음 과제이다', script=(
        '실무적으로는 하이브리드 본딩의 경쟁력이 접합 원리 하나에 좌우되지 않을 수 있습니다. 차세대 HBM에 하이브리드 본딩을 도입하려는 한국 메모리 '
        '산업에는 평탄화, 표면 화학, 정렬 계측 같은 제조공정 역량과 함께 얇은 다이를 다루고 적층하는 공정 역량이 필요할 수 있습니다. 한계도 있습니다. '
        '공동분류는 대리측정치이고, CPC 재분류의 영향을 배제할 수 없으며, 제목 검색과 패밀리 정보에도 한계가 있고, 결과는 인과가 아닌 조건부 연관입니다. '
        '향후에는 출원인과 패밀리 정보를 결합하고, 다른 공정기술과 비교하여 이 세 국면의 일반성을 검증하고자 합니다.'))
    cw2 = (CW - 0.8) / 2
    blocks = [('FaBuilding', '실무적 · 정책적 시사점', [('접합 역량만으로는 부족: 평탄화 · 오염 제어 · 표면 화학 · 정렬을 아우르는 인접 공정 지식', False),
                                                       ('재편기: 경쟁의 초점이 핵심 원리 → 다이 단위 공정과의 결합(응용)으로 이동 가능성', False),
                                                       ('차세대 HBM: 제조공정 역량 + 얇은 다이를 다루고 적층하는 공정 역량', True)]),
              ('FaScaleBalanced', '한계와 향후 연구', [('공동분류는 대리측정치 · CPC 소급 재분류 가능성', False),
                                                     ('제목 키워드 검색 · 패밀리 식별자 없음 (동일 제목으로 근사)', False),
                                                     ('조건부 연관 · 출원인 등 누락변수 · 시기는 지배적 설계의 근사', False),
                                                     ('향후: 출원인 · 패밀리 결합, 디스플레이 · 배터리 등 다른 공정기술과 비교', True)])]
    for k, (ic, head, pts) in enumerate(blocks):
        x = X0 + k * (cw2 + 0.8)
        card(x, 2.7, cw2, 13.85, fill=PANEL)
        icon_circle(ic, x + 0.6, 3.15, 2.0, NAVY)
        text(x + 3.0, 3.15, cw2 - 3.4, 2.0, [P((head, 18, NAVY, True))], anchor='m')
        text(x + 0.7, 5.9, cw2 - 1.4, 10.4, [P((t_, 15.5, NAVY if last else TEXT, last), bullet=('■', BOND if last else NAVY, 70), indent=0.6, after=20, line=1.15)
                                             for t_, last in pts])


# ================================================================ build
S20 = {
    'division': '반도체 제조는 웨이퍼 위에 트랜지스터를 만드는 전공정과, 웨이퍼를 잘라 조립하고 패키징하는 후공정으로 나뉩니다. 전공정은 종합반도체기업과 '
                '파운드리가, 후공정은 조립·테스트 전문업체가 맡는 분업이 반세기 넘게 이어졌고, 이는 두 단계 사이의 인터페이스가 안정적이었기에 가능했습니다. '
                '그런데 미세화가 한계에 부딪히자 산업은 다이를 나누고 쌓아 연결하는 첨단 패키징에서 성능 향상의 기회를 찾았고, 한 단계의 성능이 다른 단계의 '
                '공정 역량에 더 직접적으로 의존하게 되었습니다.',
    'hb': '하이브리드 본딩은 솔더 범프 없이 구리 패드와 주변 산화막을 직접 맞붙이는 기술입니다. 범프 접합이 두 판을 납이라는 접착제로 붙이는 방식이라면, '
          '하이브리드 본딩은 두 표면을 거울처럼 평평하게 갈아 그대로 맞붙이는 방식입니다. 덕분에 10마이크로미터 이하의 연결 간격을 구현합니다. 그런데 이 기술의 '
          '성패를 좌우하는 평탄화, 표면 활성화, 파티클 제어, 정밀 정렬은 후공정이 쌓아 온 역량이 아니라 전공정의 제조공정 역량입니다.',
    'gap': '이 현상을 기술융합 연구에서 찾아보았지만 아쉬운 점이 세 가지 있었습니다. 첫째, 대부분 서로 다른 산업 사이의 융합을 다루었고, 한 산업 안에서 '
           '공정 단계를 가로지르는 융합은 드물었습니다. 둘째, 융합을 분야 수준의 지수로 측정해 어떤 발명이 융합을 담는지는 설명하기 어려웠습니다. 셋째, 이론은 '
           '융합을 단계적 과정으로 보지만 실증 연구는 융합이 늘었는지를 확인하는 데 머물렀습니다. 그래서 저는 공정 단계 사이의 융합을 개별 특허 수준에서, '
           '시간의 흐름에 따라 분석했습니다.',
    'conv': '먼저 기술융합의 개념입니다. Rosenberg는 총기와 재봉틀을 위해 개발된 가공 기술이 공통의 생산공정을 통해 다른 산업으로 퍼지고 공작기계 산업이 '
            '등장한 과정을 기술적 수렴이라 불렀고, Kodama는 분리된 기술 분야가 결합해 새로운 역량을 만드는 과정을 기술융합으로 정의했습니다. Hacklin 등은 '
            '융합이 지식, 기술, 응용, 산업의 네 단계로 진행된다고 보았습니다. 하이브리드 본딩에 대입하면, 제조공정 지식이 칩 연결 문제에 쓰이기 시작하는 것이 '
            '지식 융합, 그 지식이 접합 구조와 하나의 발명 안에서 결합되는 것이 기술 융합, HBM 같은 제품으로 구현되는 것이 응용 융합, 분업 구조가 바뀌는 것이 '
            '산업 융합입니다. 이 가운데 특허 공동분류로 직접 관찰할 수 있는 것은 기술 융합 단계입니다.',
    'boundary': '두 번째 이론은 수직 전문화와 지식 경계입니다. 반도체 산업은 설계, 제조, 조립이 분리되며 팹리스, 파운드리, 조립·테스트 전문업체가 등장한 '
                '대표적인 산업입니다. 그런데 Kapoor와 Adner는 DRAM 산업의 12개 제품 세대를 분석해, 외주화한 부품의 지식까지 보유한 기업이 기술 전환에서 성과 '
                '우위를 얻었음을 보였습니다. 즉, 생산의 경계와 지식의 경계는 다릅니다. 인터페이스가 불안정해지면, 그 징후는 조직이나 산업 구조보다 먼저 두 '
                '단계의 지식을 함께 담은 발명으로 나타날 것입니다.',
    'cycle': '세 번째는 기술 순환 모형과 핵심–주변 구조입니다. 새로운 기술은 여러 설계가 경쟁하는 격동기를 거쳐 지배적 설계로 수렴하고, 그 뒤에는 점진적 '
             '변화가 이어집니다. Murmann과 Frenken은 지배적 설계를 다른 요소와 안정된 인터페이스 역할을 하는 핵심 요소의 설계로 보았고, 핵심이 안정되면 이후의 '
             '변화는 주변 요소에서 일어납니다. 도로에 비유하면, 규격이 정해지기 전에는 설계와 시공이 한 몸처럼 움직이지만, 규격이 정해진 뒤에는 각자 그 규격을 '
             '전제로 일하는 것과 같습니다.',
    'rq2why': '두 번째는 시간적 변화입니다. 단계 모형과 기술 순환 모형을 결합하면, 기술이 형성될 때는 직접 접합을 가능하게 하는 공정 발명이 주를 이루고, 확산될 '
              '때는 공정 지식과 접합 발명의 결합이 늘며, 지배적 설계가 자리 잡은 뒤에는 결합이 응용에 필요한 주변 공정으로 옮겨 갈 것으로 기대됩니다. 실제로 '
              '하이브리드 본딩은 웨이퍼 단위 접합에서 HBM과 로직의 다이 단위 접합으로 확대되었고, 발명의 중심도 공정에서 접합 구조와 적층으로 옮겨 갔습니다. '
              '그래서 두 번째 연구 질문은 "공정 단계 간 기술융합은 시간에 따라 어떻게 변화해 왔는가"입니다.',
    'rq3why': '세 번째는 융합의 중심 변화입니다. 하이브리드 본딩의 핵심은 범프 없이 접합 패드끼리 직접 연결하는 방식입니다. 이 인터페이스가 확정되기 전에는 '
              '핵심 기술을 발명하려면 평탄화나 웨이퍼 접합 같은 제조공정 지식을 함께 다루어야 합니다. 그러나 설계가 고착화되면 그 지식은 주어진 조건이 되고, '
              '새로운 결합은 다이를 다루고 패키지로 완성하는 주변 공정에서 일어날 것입니다. 그래서 세 번째 연구 질문은 "공정 단계 간 기술융합의 중심이 핵심 '
              '기술과 주변 기술 사이에서 시기에 따라 어떻게 재편되는가"입니다.',
    'object': '분석 대상인 하이브리드 본딩은 기반 공정, 접합 형태, 소자 구조, 적용 제품의 네 층위로 정리할 수 있습니다. 기반 공정은 CPC의 제조공정 클래스인 '
              'H01L21에, 접합 형태는 접합·인터커넥트 클래스인 H01L24에 주로 대응합니다. 오른쪽의 실제 공정 흐름을 보면, 다이를 자르기 전까지 증착, 식각, 도금, '
              '연마가 이어지는데 모두 전공정의 공정과 장비입니다. 주요 출원인에도 TSMC, Intel, Samsung처럼 파운드리와 메모리 기업이 함께 들어 있습니다.',
    'method': '분석 방법은 세 가지 회귀모형입니다. 주모형은 다항 로짓으로, 출원이 세 유형 중 어디에 속하는지를 기술범위와 출원연도로 설명합니다. 기술범위는 '
              '두 클래스 안에서 출원이 가진 세부 분류 기호의 수입니다. 보완분석으로 제조공정 분류 보유 여부의 이항 로짓을, 세 번째 질문을 위해 접합 출원 '
              '763건에서 핵심 기술과 출원 시기의 상호작용을 보는 이항 로짓을 사용했습니다. 여기에 Jaccard 계수와 매개 중심성을 보조 지표로 썼습니다.',
    'rq1': '이제 연구 결과입니다. 첫 번째 연구 질문의 결과, 928건 가운데 가장 많은 유형은 통합형으로 45.8%였고, 패키징 단계 공정을 뺀 좁은 정의에서도 '
           '37.2%였습니다. 공정 단계 간 융합은 예외가 아니었습니다. 그렇다면 어떤 특허가 경계를 넘었을까요? 세부 분류가 하나 늘 때마다 통합형일 상대위험이 40% '
           '높아져, 기술범위 3에서 0.26이던 예측확률이 8에서는 0.68이 되었습니다. 통합형은 정의상 기호가 두 개 이상이라는 기계적 측면이 있어 이를 조정한 '
           '지표로도 확인했는데, 방향은 유지되었습니다.',
    'practice': '실무적으로는 두 가지 시사점이 있습니다. 첫째, 상당수의 출원이 제조공정과 접합 기술을 함께 담고 있다는 것은 접합 기술 역량만으로는 충분하지 '
                '않고 평탄화, 오염 제어, 표면 화학, 정렬을 아우르는 인접 공정 지식이 중요해질 수 있음을 뜻합니다. 둘째, 재편기에는 경쟁의 초점이 핵심 원리에서 '
                '다이 단위 공정과의 결합으로 옮겨 가고 있을 가능성이 있습니다. 차세대 HBM에는 얇은 다이를 다루고 적층하는 공정 역량까지 필요할 수 있습니다.',
    'limits': '한계도 있습니다. 첫째, 공동분류는 대리측정치이고 CPC가 소급 재분류될 수 있어 시간적 변화의 일부는 분류 체계의 변화를 반영했을 수 있습니다. '
              '둘째, 제목 키워드 검색이라 관련어가 없는 출원은 빠졌습니다. 셋째, 패밀리 정보가 없어 같은 제목으로 중복을 근사했습니다. 넷째, 결과는 인과가 아닌 '
              '조건부 연관입니다. 향후에는 출원인과 패밀리 정보를 결합하고, 다른 공정기술과 비교하여 세 국면의 일반성을 검증하고자 합니다.'}

S15 = {
    'cover': '안녕하십니까. 기술경영학과 이형규입니다. "반도체 하이브리드 본딩의 공정 단계 간 기술융합: 특허 기반 종단 분석"이라는 주제로 발표를 시작하겠습니다.',
    'contents': '발표는 다섯 부분으로 구성되어 있습니다. 왜 이 연구를 시작했는지, 어떤 이론으로 연구 질문을 세웠고 어떻게 분석했는지, 그리고 어떤 결과가 나왔으며 무엇을 의미하는지 말씀드리겠습니다.',
    'hook': '생성형 AI, AI 가속기, 그리고 HBM. 이 세 단어를 잇는 공통의 과제는 무엇일까요? 생성형 AI는 AI 가속기를, AI 가속기는 D램을 여러 층 쌓은 HBM을 '
            '필요로 합니다. 메모리를 더 높이, 더 촘촘하게 쌓아 연결하는 것이 관건이 되면서, 그 연결을 위한 차세대 접합 기술로 주목받는 것이 하이브리드 본딩입니다.',
    'hb': ('반도체 제조는 웨이퍼 위에 소자를 만드는 전공정과, 이를 잘라 조립하는 후공정으로 반세기 넘게 나뉘어 왔고, 이 분업은 두 단계 사이의 인터페이스가 '
          '안정적이었기에 가능했습니다. 하이브리드 본딩은 솔더 범프 없이 구리 패드와 산화막을 직접 맞붙여 10마이크로미터 이하의 연결 간격을 구현하는 후공정 '
          '기술입니다. 범프 접합이 접착제로 붙이는 방식이라면, 하이브리드 본딩은 거울처럼 평평한 두 면을 그대로 맞붙이는 방식입니다. 그런데 그 성패를 좌우하는 '
          '평탄화, 표면 활성화, 파티클 제어, 정밀 정렬은 전공정의 제조공정 역량입니다. 후공정 기술의 성패를 전공정 역량이 좌우하게 된 것입니다.'),
    'gap': '기존 기술융합 연구는 주로 서로 다른 산업 사이의 융합을 분야 수준의 지수로 측정했고, 융합이 늘었는지를 확인하는 데 머물렀습니다. 그래서 저는 '
           '한 산업의 공정 단계 사이에서 일어나는 융합을 개별 특허 수준에서, 시간의 흐름에 따라 분석하였습니다.',
    'conv': 'Rosenberg는 공통의 생산공정을 통해 기술이 산업의 경계를 넘어 퍼지는 현상을 기술적 수렴이라 했고, Kodama는 분리된 기술이 결합하여 새로운 '
            '역량을 만드는 과정을 기술융합으로 정의했습니다. Hacklin 등은 융합이 지식, 기술, 응용, 산업의 단계로 진행된다고 보았습니다. 하이브리드 본딩에서 '
            '제조공정 지식이 접합 구조와 하나의 발명 안에서 결합되는 것이 기술 융합 단계이며, 특허 공동분류로 직접 관찰할 수 있는 것도 이 단계입니다.',
    'frame': ('여기에 두 가지 이론을 더했습니다. 첫째, 수직 전문화와 지식 경계입니다. Kapoor와 Adner는 DRAM 산업에서 외주화한 부품의 지식까지 보유한 기업이 '
             '기술 전환에서 성과 우위를 얻었음을 보였습니다. 생산의 경계와 지식의 경계는 다르기 때문에, 인터페이스가 불안정해지면 그 징후는 조직보다 먼저 경계를 '
             '넘는 발명으로 나타날 것입니다. 둘째, 기술 순환 모형과 핵심–주변 구조입니다. 지배적 설계가 고착화되면 새로운 결합은 핵심이 아닌 주변에서 일어날 '
             '것입니다. 도로의 규격이 정해지면 새로운 일이 그 규격 위에서 이루어지는 것과 같습니다.'),
    'data': 'PATSTAT에서 제조공정 분류 H01L21이나 접합 분류 H01L24를 가지면서 제목에 하이브리드 본딩 관련 키워드가 들어간 출원 928건을 추출했습니다. '
            '출원은 2010년대 후반에 급증하여 2022년에 정점에 이르렀고, 2023년 이후의 감소는 공개 시차 때문일 가능성이 큽니다.',
    'measure': '융합은 CPC 공동분류로 측정했습니다. 실제 출원 세 건을 보시면, 제조공정 기호만 가지면 공정 단독형, 접합 기호만 가지면 접합 단독형, 둘을 '
               '함께 가지면 공정 단계의 경계를 넘는 통합형입니다. 다만 이것은 대리측정치이며 법적 청구범위를 뜻하지는 않습니다.',
    'method': ('분석은 세 유형을 설명하는 다항 로짓을 주모형으로 했습니다. 설명변수는 두 클래스 안에서 출원이 가진 세부 분류 기호의 수인 기술범위와 출원연도입니다. '
              '제조공정 분류 보유 여부의 이항 로짓은 보완분석이고, 세 번째 질문은 접합 출원 763건에서 핵심 기술과 출원 시기의 상호작용을 이항 로짓으로 '
              '보았습니다. 출원 시기는 2014년 이전, 2015년부터 2019년, 2020년부터 2024년으로 나누었습니다. 핵심 기술은 기술 클러스터와 네트워크로 식별한 '
              '접합 패드 기호 네 개이며, 접합 출원의 81.3%가 이를 포함합니다.'),
    'rq1': '첫 번째 결과입니다. 통합형이 45.8%로 가장 많았고, 패키징 단계 공정을 뺀 좁은 정의에서도 37.2%였습니다. 또 세부 분류가 하나 늘 때마다 '
           '통합형일 상대위험이 40% 높아져, 기술범위 3에서 0.26이던 예측확률이 8에서는 0.68이 되었습니다. 공정 단계 간 융합은 예외가 아니었고, 경계를 넘는 '
           '특허는 더 다양한 세부 기술을 담고 있었습니다.',
    'rq2': ('두 번째 결과입니다. 2014년 이전 47년간의 출원은 이후 5년간보다 적을 만큼 드물어 하나의 시기로 묶었습니다. 2014년 이전에는 공정 단독형이 43.6%로 '
           '가장 많았고, 2015년부터 2019년까지는 통합형이 62.2%로 정점에 이르렀으며, 2020년 이후에는 통합형이 44.3%로 낮아지고 접합 단독형이 45.4%로 '
           '늘었습니다. 저는 이를 형성기, 확산기, 재편기라 이름 붙였습니다. 회귀분석에서도 공정 단독형의 상대위험은 매년 8.7%씩 줄어, 출원의 무게중심이 '
           '제조공정에서 접합 쪽으로 이동했습니다.'),
    'rq3a': '세 번째 결과입니다. 핵심 기술을 포함한 접합 출원이 통합형일 승산은 형성기에 11.1배였지만, 확산기에는 2.5배로 줄었고, 재편기에는 유의한 차이가 '
            '확인되지 않았습니다. 그사이 핵심 기술을 포함한 출원의 비중은 53.5%에서 88.0%로 높아져, 발명이 하나의 핵심 설계로 수렴한 것으로 보입니다.',
    'rq3b': '그렇다면 융합은 어디로 옮겨 갔을까요? 웨이퍼 제조공정과 접합의 결합 강도는 확산기 이후 크게 낮아졌고, 공정 장치와 패키징 단계 공정은 높아졌습니다. '
            '핵심 기호는 네트워크의 허브로 남았지만, 접합과 제조공정을 잇는 연결 기호는 공정 장치로 바뀌었고, 핵심 기술과 함께 부여된 공정도 다이 지지와 봉지 '
            '공정으로 바뀌었습니다.',
    'summary': ('결과를 하나로 정리하면, 형성기에는 제조공정 발명이 단독으로 많이 출원되며 접합에 필요한 지식이 쌓였고, 확산기에는 두 영역의 지식이 하나의 발명 '
               '안에서 결합되며 통합형이 정점에 이르렀으며, 재편기에는 핵심 설계가 고착화되면서 결합의 중심이 다이 지지와 봉지 같은 주변 공정으로 옮겨 '
               '갔습니다. 이 결과는 여러 재추정에서도 유지되었고, 시간 추세만 2014년 이전 출원이 30건뿐인 hybrid bonding 표본에서 식별되지 않았습니다.'),
    'concl': '결론적으로, 기술융합은 한 산업의 공정 단계 사이에서도 개별 발명으로 나타나며, 한 방향으로 쌓이는 것이 아니라 핵심 설계의 고착화와 함께 중심이 '
             '핵심에서 주변으로 이동하는 동태적 과정이었습니다. 이 연구는 기술융합 연구의 범위를 공정 단계 간 융합으로 넓히고, 그 전개를 핵심–주변 구조와 기술 '
             '순환의 관점에서 설명할 실증적 근거를 제시했다는 데 의의가 있습니다.',
    'rqs': ('이 이론들로부터 세 가지 연구 질문을 세웠습니다. 첫째, 융합이 어느 정도로, 어떤 형태로 나타나는가입니다. 지식 경계의 논리에 따르면 두 영역을 함께 '
            '담은 발명은 예외가 아닐 것입니다. 둘째, 시간에 따라 어떻게 변화해 왔는가입니다. 융합은 꾸준히 늘기보다 발전 단계에 따라 달라질 것입니다. 셋째, '
            '융합의 중심이 핵심 기술과 주변 기술 사이에서 어떻게 재편되는가입니다. 핵심 설계가 고착화되면 중심은 주변 공정으로 이동할 것입니다. 이 기대들은 '
            '통계적 가설이 아니라 결과를 해석하는 기준으로 사용했습니다.'),
    'close': '이상으로 발표를 마치겠습니다. 경청해 주셔서 감사합니다. 질문 주시면 성실히 답변드리겠습니다.'}

if VER == '20':
    seq = [(s_cover, {}), (s_contents, {}),
           (s_intro_hook, {}), (s_intro_division, {'script': S20['division']}), (s_intro_hb, {'script': S20['hb']}), (s_intro_gap, {'script': S20['gap']}),
           (s_th_conv_stage, {'script': S20['conv']}), (s_th_boundary, {'script': S20['boundary']}), (s_th_cycle, {'script': S20['cycle']}), (s_th_summary, {}),
           (s_rq1_why, {}), (s_rq2_why, {'script': S20['rq2why']}), (s_rq3_why, {'script': S20['rq3why']}),
           (s_object, {'script': S20['object']}), (s_data, {}), (s_measure, {}), (s_core, {}), (s_method, {'script': S20['method']}),
           (s_res_rq1, {'script': S20['rq1']}), (s_res_rq2a, {}), (s_res_rq2b, {}), (s_res_rq3a, {}), (s_res_rq3b, {}), (s_robust, {}), (s_res_summary, {}),
           (s_conclusion, {}), (s_practice, {'script': S20['practice']}), (s_limits, {'script': S20['limits']}), (s_close, {})]
    backup = []
else:
    seq = [(s_cover, {}), (s_contents, {'script': S15['contents']}),
           (s_intro_hook, {}), (s_intro_hb, {'script': S15['hb']}), (s_intro_gap, {'script': S20['gap']}),
           (s_th_conv_stage, {'script': S20['conv']}), (s_th_summary, {'script': S15['frame'], 'num': '2.2', 'title': '지식 경계와 기술 순환: 이론적 틀'}),
           (s_rqs_15, {'script': S15['rqs']}),
           (s_data, {}), (s_measure, {}), (s_method, {'script': S15['method'], 'core': True}),
           (s_res_rq1, {'script': S20['rq1']}), (s_res_rq2a, {'script': S15['rq2'], 'reg': True}),
           (s_res_rq3a, {}), (s_res_rq3b, {}),
           (s_res_summary, {'script': S15['summary'], 'num': '4.4', 'title': '결과 종합', 'robust': True}),
           (s_conclusion, {}), (s_impl_limits_15, {}), (s_close, {})]
    # slides of the 20-minute talk that the 15-minute talk leaves for questions
    backup = [s_intro_division, s_th_boundary, s_th_cycle, s_object, s_core, s_res_rq2b, s_robust, s_practice]
for fn, kw in seq:
    sc = kw.pop('script', None)
    run(fn, script=sc, **kw)
N_MAIN = len(K.prs.slides)
n = 0
for fn in backup:
    n += 1
    run(fn, appx=n)          # keeps the 20-minute notes, useful when answering
APP = [s_app_data, s_app_vars, s_app_cluster, s_app_logit, s_app_core, s_app_robust1, s_app_robust2]
for fn in APP:
    n += 1
    run(fn, appx=n)
save(OUT)
json.dump({'version': VER, 'n_main': N_MAIN, 'notes': K.NOTES}, open(OUT.replace('.pptx', '_notes.json'), 'w'), ensure_ascii=False, indent=1)

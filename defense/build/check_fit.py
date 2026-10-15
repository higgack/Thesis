# -*- coding: utf-8 -*-
"""Overflow check for every text frame of a deck as PowerPoint lays it out with 맑은 고딕 (adapted from poster/build).
Korean wraps at spaces only; widths = Noto Sans CJK x WSCALE (1.2 calibrated on the user's PowerPoint screenshots).
Usage: python3 check_fit.py deck.pptx [WSCALE=1.2]   -> one line per overflowing frame; exit 1 if any."""
import sys
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn
from PIL import ImageFont

FONT = {False: ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', 1000, index=1),
        True: ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc', 1000, index=1)}
PT_CM = 0.03528


def runs_of(p):
    return [(r.text, r.font.size.pt if r.font.size else 18.0, bool(r.font.bold)) for r in p.runs]


def para_lines(p, avail, scale):
    toks = []
    for t, sz, b in runs_of(p):
        cur = ''
        for ch in t:
            if ch == ' ':
                if cur: toks.append((FONT[b].getlength(cur) / 1000 * sz * PT_CM * scale, False)); cur = ''
                toks.append((FONT[b].getlength(' ') / 1000 * sz * PT_CM * scale, True))
            else:
                cur += ch
        if cur: toks.append((FONT[b].getlength(cur) / 1000 * sz * PT_CM * scale, False))
    words, cur = [], 0.0
    for w, sp in toks:
        if sp:
            if cur: words.append(cur); cur = 0.0
            words.append(-w)
        else:
            cur += w
    if cur: words.append(cur)
    lines, x, widest, pend = 1, 0.0, 0.0, 0.0
    for w in words:
        if w < 0: pend += -w; continue
        widest = max(widest, w)
        if x > 0 and x + pend + w > avail: lines += 1; x = w
        else: x += (pend if x > 0 else 0) + w
        pend = 0.0
    return lines, widest


def pitch(p):
    pPr = p._p.find(qn('a:pPr')); sizes = [s for _, s, _ in runs_of(p)] or [18.0]
    if pPr is not None:
        ls = pPr.find(qn('a:lnSpc'))
        if ls is not None:
            pts, pct = ls.find(qn('a:spcPts')), ls.find(qn('a:spcPct'))
            if pts is not None: return int(pts.get('val')) / 100 * PT_CM
            if pct is not None: return max(sizes) * 1.2 * int(pct.get('val')) / 100000 * PT_CM
    return max(sizes) * 1.2 * PT_CM


def spacing(p):
    pPr = p._p.find(qn('a:pPr')); tot = 0.0
    if pPr is not None:
        for tag in ('a:spcBef', 'a:spcAft'):
            el = pPr.find(qn(tag))
            if el is not None and el.find(qn('a:spcPts')) is not None:
                tot += int(el.find(qn('a:spcPts')).get('val')) / 100 * PT_CM
    return tot


def indent(p):
    pPr = p._p.find(qn('a:pPr'))
    return Emu(int(pPr.get('marL'))).cm if pPr is not None and pPr.get('marL') else 0.0


def check(tf, w, h, scale):
    bp = tf._txBody.find(qn('a:bodyPr'))
    ins = {k: (Emu(int(bp.get(k))).cm if bp is not None and bp.get(k) else d) for k, d in (('lIns', 0.254), ('rIns', 0.254), ('tIns', 0.127), ('bIns', 0.127))}
    avail = w - ins['lIns'] - ins['rIns']
    need, wide = 0.0, []
    ps = list(tf.paragraphs)
    for i, p in enumerate(ps):
        if not ''.join(r.text for r in p.runs).strip():
            need += pitch(p) if i < len(ps) - 1 else 0; continue
        a = avail - indent(p)
        n, wd = para_lines(p, a, scale)
        need += n * pitch(p) + (spacing(p) if i < len(ps) - 1 else 0)
        if wd > a + 0.01: wide.append((''.join(r.text for r in p.runs)[:24], round(wd, 2), round(a, 2)))
    return need, h - ins['tIns'] - ins['bIns'], wide


def main(path, scale):
    prs = Presentation(path); bad = 0
    for si, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            frames = []
            if sh.has_text_frame and sh.text_frame.text.strip():
                frames.append((sh.text_frame, Emu(sh.width).cm, Emu(sh.height).cm, sh.text_frame.text.strip().replace('\n', ' / ')[:36]))
            if getattr(sh, 'has_table', False) and sh.has_table:
                tbl = sh.table
                for row in tbl.rows:
                    for ci, cell in enumerate(row.cells):
                        if cell.text.strip() and not cell.is_spanned:
                            span = cell.span_width if cell.is_merge_origin else 1
                            w = sum(Emu(tbl.columns[ci + k].width).cm for k in range(span))
                            frames.append((cell.text_frame, w, Emu(row.height).cm, 'cell: ' + cell.text.strip()[:28]))
            for tf, w, h, lab in frames:
                need, room, wide = check(tf, w, h, scale)
                if need > room + 0.02 or wide:
                    bad += 1
                    print(f'slide {si:2d}: need {need:4.2f} > room {room:4.2f} (w {w:4.1f})  {lab}' + (f'  WIDE {wide}' if wide else ''))
    print('frames over:', bad, '(scale', scale, ')')
    return bad


if __name__ == '__main__':
    sys.exit(1 if main(sys.argv[1], float(sys.argv[2]) if len(sys.argv) > 2 else 1.2) else 0)

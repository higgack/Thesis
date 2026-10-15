# -*- coding: utf-8 -*-
"""Helpers for the thesis-defense deck (16:9). Structure follows the reference talk the user shared on 10-15
(numbered "| 3.1 title" header, one take-away line per slide, chapter summaries, research-question cards);
the look is our own (user, 10-15: the reference design and colours are only an example): white content slides,
navy as the dominant colour, the poster's type colours (process-only blue, integrated purple, bonding-only
orange) as the motif, and dark image slides for the cover, section dividers, RQ cards and the close.
Text is 맑은 고딕 everywhere (latin + ea + cs), sized for PowerPoint's 맑은 고딕 widths (check_fit.py, x1.2).
Layout "본문(머리글)" carries the header bar, rule, emblem, footer and page number;
layout "전체 이미지" (blank) is used for the full-image slides."""
import copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from lxml import etree

FONT = '맑은 고딕'
SW, SH = 33.867, 19.05
# palette (light content slides; the same navy and type colours as the Research Day poster)
NAVY, NAVY2, WHITE = '13294B', '24406E', 'FFFFFF'
TEXT, MUTED, RULE = '1E2532', '5B6475', 'D5DBE3'
PANEL, PANEL2, BOXL = 'F2F5FA', 'E8EEF7', '9AA6B8'
PROC, INTEG, BOND = '2563C9', '7C3AED', 'EA7A28'          # process-only / integrated / bonding-only
PROC_L, INTEG_L, BOND_L = 'E3ECFB', 'EEE6FD', 'FDEBDD'
TEAL, GOLD, GOLDB, CORE = '10A39A', 'E0A100', 'C9A227', 'B5441B'
HL, WARN = 'FFF0C2', 'FBE3E0'                            # highlighted table cells (key result / caveat)
DIM = 'B9C6DD'                                           # secondary text on dark slides
L, C, R = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
X0, X1 = 1.5, SW - 1.5
CW = X1 - X0
TAKE_Y = 16.95

prs = s = None
LAY_TITLE = LAY_BLANK = None
ASSET = ICON = None
NOTES = []          # (slide number, section, header, script) for the script document


def cm(v): return Emu(int(round(v * 360000)))
def rgb(h): return RGBColor.from_string(h)


# ---------------------------------------------------------------- text primitives (same conventions as the poster kit)
def font_run(r, size, color=TEXT, bold=False, italic=False):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', FONT)
    rPr.set('lang', 'ko-KR'); rPr.set('altLang', 'en-US')


def set_para(p, align=L, space_after=0, line=None, bullet=None, indent=0.0, space_before=0):
    p.alignment = align
    p.space_after = Pt(space_after); p.space_before = Pt(space_before)
    if line: p.line_spacing = line
    pPr = p._p.get_or_add_pPr()
    for t in ('a:buNone', 'a:buChar', 'a:buFont', 'a:buClr', 'a:buSzPct'):
        for el in pPr.findall(qn(t)): pPr.remove(el)
    if bullet:
        ch, col = bullet[:2]
        pPr.set('marL', str(int(cm(indent)))); pPr.set('indent', str(-int(cm(indent))))
        bc = etree.SubElement(pPr, qn('a:buClr')); sc = etree.SubElement(bc, qn('a:srgbClr')); sc.set('val', col)
        bs = etree.SubElement(pPr, qn('a:buSzPct')); bs.set('val', str(int((bullet[2] if len(bullet) > 2 else 100) * 1000)))
        bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', 'Arial')
        b = etree.SubElement(pPr, qn('a:buChar')); b.set('char', ch)
    else:
        pPr.set('marL', '0'); pPr.set('indent', '0')
        etree.SubElement(pPr, qn('a:buNone'))


def norm(rr):
    rr = tuple(rr)
    return (rr[0], rr[1] if len(rr) > 1 else 16, rr[2] if len(rr) > 2 else TEXT, rr[3] if len(rr) > 3 else False,
            rr[4] if len(rr) > 4 else False)


def P(*runs, **kw):
    d = {'runs': [r if isinstance(r, tuple) else (r, kw.get('size', 16), kw.get('color', TEXT), kw.get('bold', False)) for r in runs]}
    d.update({k: v for k, v in kw.items() if k not in ('size', 'color', 'bold')}); return d


LINE_K = 1.2


def _pitch(pd):
    mx = max(norm(r)[1] for r in pd['runs'])
    return Pt(round(mx * LINE_K * (pd.get('line') or 1.0), 1))


def _fill_tf(tf, paras, default_align=L):
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_para(p, pd.get('align', default_align), pd.get('after', 0), _pitch(pd), pd.get('bullet'), pd.get('indent', 0), pd.get('before', 0))
        for (t, size, col, bold, ital) in map(norm, pd['runs']):
            r = p.add_run(); r.text = t; font_run(r, size, col, bold, ital)


def _anchor(a): return {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[a]


def _noautofit(tf):
    bp = tf._txBody.find(qn('a:bodyPr'))
    for el in list(bp): bp.remove(el)
    etree.SubElement(bp, qn('a:noAutofit'))


def text(x, y, w, h, paras, anchor='t', margin=0.0, name=None, align=None):
    if isinstance(paras, dict): paras = [paras]
    tb = s.shapes.add_textbox(cm(x), cm(y), cm(w), cm(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = cm(margin)
    tf.vertical_anchor = _anchor(anchor)
    _noautofit(tf)
    if align is not None:
        paras = [dict(p, align=p.get('align', align)) for p in paras]
    _fill_tf(tf, paras)
    return tb


def _nostyle(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None: sh._element.remove(st)


def _alpha(clr_parent, alpha):
    clr = clr_parent.find(qn('a:srgbClr'))
    a = etree.SubElement(clr, qn('a:alpha')); a.set('val', str(int(alpha * 100000)))


def shape(kind, x, y, w, h, fill=None, line=None, lw=0.0, radius=None, alpha=None, name=None, dash=None, shadow=False):
    sh = s.shapes.add_shape(kind, cm(x), cm(y), cm(w), cm(h)); _nostyle(sh)
    if name: sh.name = name
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        if alpha is not None: _alpha(sh.fill._xPr.find(qn('a:solidFill')), alpha)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
        if dash:
            ln = sh.line._get_or_add_ln(); d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', dash)
    else:
        sh.line.fill.background()
    if radius is not None:
        try: sh.adjustments[0] = radius
        except IndexError: pass
    if shadow:
        spPr = sh._element.spPr
        eff = etree.SubElement(spPr, qn('a:effectLst'))
        sd = etree.SubElement(eff, qn('a:outerShdw'))
        for k, v in (('blurRad', '127000'), ('dist', '25400'), ('dir', '5400000'), ('algn', 't'), ('rotWithShape', '0')): sd.set(k, v)
        c = etree.SubElement(sd, qn('a:srgbClr')); c.set('val', '1E2532')
        a = etree.SubElement(c, qn('a:alpha')); a.set('val', '16000')
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = cm(0.1)
    return sh


def shape_text(sh, paras, anchor='m', margin=0.15, align=C):
    if isinstance(paras, dict): paras = [paras]
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = cm(margin); tf.margin_top = tf.margin_bottom = cm(0.05)
    tf.vertical_anchor = _anchor(anchor)
    _noautofit(tf)
    _fill_tf(tf, paras, default_align=align)
    return sh


def line(x1, y1, x2, y2, color=BOXL, w=1.5, dash=None, head=None, tail=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cm(x1), cm(y1), cm(x2), cm(y2)); _nostyle(c)
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', dash)
    for tag, kind in (('a:headEnd', head), ('a:tailEnd', tail)):
        if kind:
            e = etree.SubElement(ln, qn(tag)); e.set('type', kind); e.set('w', 'med'); e.set('len', 'med')
    return c


def arrow(x1, y1, x2, y2, color=BOXL, w=2.0):
    return line(x1, y1, x2, y2, color, w, tail='triangle')


def picture(path, x, y, w=None, h=None):
    return s.shapes.add_picture(path, cm(x), cm(y), cm(w) if w else None, cm(h) if h else None)


def icon(name, x, y, size, color=NAVY):
    return picture(f'{ICON}/{name}_{color}.png', x, y, size, size)


def icon_circle(name, x, y, d, fill, color='FFFFFF', pad=0.25):
    shape(MSO_SHAPE.OVAL, x, y, d, d, fill=fill)
    return icon(name, x + d * pad, y + d * pad, d * (1 - 2 * pad), color)


def to_back(sh):
    sp = s.shapes._spTree; sp.remove(sh._element); sp.insert(2, sh._element)


# ---------------------------------------------------------------- composite pieces used across slides
def obox(x, y, w, h, title, sub=None, size=16, sub_size=11, color=NAVY, fill=WHITE, tcolor=None, lw=1.75, radius=0.18,
         sub_color=MUTED, shadow=False, bold=True):
    """rounded box with a bold title and an optional small explanatory line"""
    sh = shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=color, lw=lw, radius=radius, shadow=shadow)
    paras = [P((title, size, tcolor or color, bold), align=C)]
    if sub:
        for ln_ in sub.split('\n'):
            paras.append(P((ln_, sub_size, sub_color), align=C))
    shape_text(sh, paras)
    return sh


def chip(x, y, w, h, t, fill=PANEL2, color=NAVY, size=12, bold=True, line_=None, radius=0.5):
    return shape_text(shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line_, lw=1.0 if line_ else 0, radius=radius),
                      P((t, size, color, bold), align=C), margin=0.08)


def card(x, y, w, h, fill=PANEL, line_=None, radius=0.06, shadow=False):
    return shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line_, lw=1.0 if line_ else 0, radius=radius, shadow=shadow)


def tab_box(x, y, w, h, tab, paras, tab_w=None, color=NAVY, tab_fill=NAVY):
    """box with a label tab sitting on its top edge (reference: Dependent / Independent / Control Variable)"""
    shape(MSO_SHAPE.RECTANGLE, x, y, w, h, fill=WHITE, line=color, lw=1.25)
    tw = tab_w or (w - 1.6)
    shape_text(shape(MSO_SHAPE.RECTANGLE, x + (w - tw) / 2, y - 0.45, tw, 0.9, fill=tab_fill), P((tab, 13, WHITE, True), align=C), margin=0.05)
    text(x + 0.4, y + 0.7, w - 0.8, h - 0.9, paras, anchor='t')


def takeaway(t, size=16):
    sh = shape(MSO_SHAPE.ROUNDED_RECTANGLE, X0, TAKE_Y, CW, 1.05, fill=PANEL2, radius=0.5, name='Takeaway')
    shape_text(sh, P((t, size, NAVY, True), align=C), margin=0.4)
    return sh


def source(t, y=16.35, x=X0, w=CW, align=L):
    text(x, y, w, 0.5, [P((t, 10, MUTED), align=align)], anchor='m', name='Source')


# ---------------------------------------------------------------- tables
NO_GRID = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'


def _cell_border(cell, side, color, w):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {'L': 'a:lnL', 'R': 'a:lnR', 'T': 'a:lnT', 'B': 'a:lnB'}[side]
    for el in tcPr.findall(qn(tag)): tcPr.remove(el)
    ln = etree.Element(qn(tag)); ln.set('w', str(int(Pt(w)))); ln.set('cap', 'flat'); ln.set('cmpd', 'sng'); ln.set('algn', 'ctr')
    if color:
        sf = etree.SubElement(ln, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', color)
    else:
        etree.SubElement(ln, qn('a:noFill'))
    etree.SubElement(ln, qn('a:prstDash')).set('val', 'solid')
    order = ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB']
    pos = sum(1 for t in order[:order.index(tag)] if tcPr.find(qn(t)) is not None)
    tcPr.insert(pos, ln)


def table(x, y, colw, rows, rowh, size=13, head_size=None, aligns=None, hl=None, warn=None, bold_cols=(), head_fill=NAVY,
          body=(WHITE, 'F4F7FB'), rule='C9D2E0', merge=None, head_rows=1, colors=None):
    """rows: list of lists (first `head_rows` rows = header). hl: cells (i, j) highlighted (key result);
    warn: cells shown as caveats; colors: {(i, j): text colour}; merge: list of ((r0, c0), (r1, c1))."""
    nR, nC = len(rows), len(rows[0])
    rh = rowh if isinstance(rowh, list) else [rowh] * nR
    gf = s.shapes.add_table(nR, nC, cm(x), cm(y), cm(sum(colw)), cm(sum(rh)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None: sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
    sid.text = NO_GRID
    tbl.first_row = True; tbl.horz_banding = False
    for j, w_ in enumerate(colw): tbl.columns[j].width = cm(w_)
    for i in range(nR): tbl.rows[i].height = cm(rh[i])
    hl = set(hl or []); warn = set(warn or []); colors = colors or {}
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            head = i < head_rows
            fill = head_fill if head else body[(i - head_rows) % 2]
            col = WHITE if head else TEXT
            bold = head or (j in bold_cols)
            if (i, j) in hl: fill, col, bold = HL, NAVY, True
            if (i, j) in warn: fill, col, bold = WARN, '9B2C1F', True
            if (i, j) in colors: col = colors[(i, j)]
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fill)
            cell.margin_left = cell.margin_right = cm(0.12); cell.margin_top = cell.margin_bottom = cm(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            al = C if head else (aligns[j] if aligns else C)
            sz = (head_size or size) if head else size
            _fill_tf(tf, [P((ln_, sz, col, bold), align=al) for ln_ in str(val).split('\n')])
            for side in 'LR': _cell_border(cell, side, None, 0)
            _cell_border(cell, 'T', None, 0)
            _cell_border(cell, 'B', rule if i < nR - 1 else NAVY, 0.75 if i < nR - 1 else 1.5)
    for (a, b) in (merge or []):
        tbl.cell(*a).merge(tbl.cell(*b))
    return gf


# ---------------------------------------------------------------- charts
def _chart_bg_none(ch):
    cs = ch._chartSpace
    sp = etree.Element(qn('c:spPr')); etree.SubElement(sp, qn('a:noFill'))
    ln = etree.SubElement(sp, qn('a:ln')); etree.SubElement(ln, qn('a:noFill'))
    cs.find(qn('c:chart')).addnext(sp)
    pa = cs.find(qn('c:chart')).find(qn('c:plotArea'))
    sp = etree.Element(qn('c:spPr')); etree.SubElement(sp, qn('a:noFill'))
    ln = etree.SubElement(sp, qn('a:ln')); etree.SubElement(ln, qn('a:noFill'))
    ext = pa.find(qn('c:extLst'))
    if ext is not None: ext.addprevious(sp)
    else: pa.append(sp)


def style_chart(ch, legend=True, size=13, axes=True):
    fg, axis, grid = TEXT, '9AA3B2', 'DDE2EA'
    ch.font.size = Pt(size); ch.font.name = FONT; ch.font.color.rgb = rgb(fg)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size); ch.legend.font.name = FONT; ch.legend.font.color.rgb = rgb(fg)
    if axes:
        for ax in (ch.category_axis, ch.value_axis):
            ax.tick_labels.font.size = Pt(size); ax.tick_labels.font.name = FONT; ax.tick_labels.font.color.rgb = rgb(fg)
            ax.format.line.color.rgb = rgb(axis)
        ch.value_axis.has_major_gridlines = True
        ch.value_axis.major_gridlines.format.line.color.rgb = rgb(grid)
        ch.value_axis.format.line.fill.background()
    for d in ch._chartSpace.iter(qn('a:defRPr')):
        if d.find(qn('a:ea')) is None:
            ea = etree.SubElement(d, qn('a:ea')); ea.set('typeface', FONT)
    _chart_bg_none(ch)


def axis_title(ax, t, size=12):
    ax.has_title = True
    tf = ax.axis_title.text_frame
    tf.text = t
    for p in tf.paragraphs:
        for r in p.runs: font_run(r, size, MUTED)


# ---------------------------------------------------------------- deck scaffolding
def _theme(prs_):
    tp = prs_.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(tp.blob)
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    root.set('name', 'Defense')
    cs = root.find('.//a:clrScheme', ns); cs.set('name', 'Defense')
    cols = {'dk1': TEXT, 'lt1': 'FFFFFF', 'dk2': NAVY, 'lt2': PANEL, 'accent1': NAVY2, 'accent2': PROC,
            'accent3': INTEG, 'accent4': BOND, 'accent5': TEAL, 'accent6': GOLD, 'hlink': PROC, 'folHlink': INTEG}
    for k, v in cols.items():
        el = cs.find(f'a:{k}', ns)
        for c in list(el): el.remove(c)
        etree.SubElement(el, qn('a:srgbClr')).set('val', v)
    fs = root.find('.//a:fontScheme', ns); fs.set('name', 'Malgun Gothic')
    for tag in ('a:majorFont', 'a:minorFont'):
        f = fs.find(tag, ns)
        f.find('a:latin', ns).set('typeface', FONT)
        f.find('a:ea', ns).set('typeface', FONT)
        for el in f.findall('a:font', ns):
            if el.get('script') == 'Hang': el.set('typeface', FONT)
    tp._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def _move(ph, x, y, w, h):
    ph.left, ph.top, ph.width, ph.height = cm(x), cm(y), cm(w), cm(h)


FOOTER = '이형규 · 석사학위논문 심사'


def setup(asset, icondir):
    """new 16:9 deck with the theme, title style and the header/footer decoration on the content layout"""
    global prs, s, LAY_TITLE, LAY_BLANK, ASSET, ICON
    ASSET, ICON = asset, icondir
    prs = Presentation()
    prs.slide_width, prs.slide_height = cm(SW), cm(SH)
    prs.core_properties.title = '반도체 하이브리드 본딩의 공정 단계 간 기술융합: 특허 기반 종단 분석 (석사학위논문 심사 발표)'
    prs.core_properties.author = '이형규'
    _theme(prs)
    m = prs.slide_master
    m.background.fill.solid(); m.background.fill.fore_color.rgb = rgb(WHITE)
    ts = m._element.find(qn('p:txStyles')).find(qn('p:titleStyle')).find(qn('a:lvl1pPr'))
    ts.set('algn', 'l')
    d = ts.find(qn('a:defRPr')); d.set('sz', '2200'); d.set('b', '1')
    for c in list(d):
        if c.tag == qn('a:solidFill'): d.remove(c)
    sf = etree.Element(qn('a:solidFill')); etree.SubElement(sf, qn('a:schemeClr')).set('val', 'tx2'); d.insert(0, sf)
    for ph in m.placeholders:
        t = str(ph.placeholder_format.type)
        if 'TITLE' in t: _move(ph, 1.6, 0.4, 25.6, 1.25)
        elif 'BODY' in t: _move(ph, X0, 2.6, CW, 14.0)
        elif 'DATE' in t: _move(ph, X0, 18.25, 6, 0.6)
        elif 'FOOTER' in t: _move(ph, 11, 18.25, 12, 0.6)
        elif 'SLIDE_NUMBER' in t: _move(ph, 30.6, 18.25, 2.6, 0.6)
    lays = list(prs.slide_layouts)
    LAY_TITLE, LAY_BLANK = lays[5], lays[6]
    for lay in lays:
        if lay not in (LAY_TITLE, LAY_BLANK): prs.slide_layouts.remove(lay)
    LAY_TITLE.name = '본문(머리글)'; LAY_BLANK.name = '전체 이미지'
    for ph in LAY_TITLE.placeholders:
        if ph.placeholder_format.idx == 0: _move(ph, 1.6, 0.4, 25.6, 1.25)
    # header/footer decoration, drawn on a scratch slide and moved onto the layout
    tmp = prs.slides.add_slide(LAY_BLANK)
    s = tmp
    els = [shape(MSO_SHAPE.RECTANGLE, 1.2, 0.62, 0.13, 0.82, fill=NAVY, name='Header bar')._element,
           line(0, 1.92, SW, 1.92, RULE, 0.75)._element,
           text(X0, 18.25, 12.0, 0.6, [P((FOOTER, 10, MUTED))], anchor='m', name='Footer')._element]
    pn = text(30.6, 18.25, 2.6, 0.6, [P(('#', 10, MUTED), align=R)], anchor='m', name='Page number')
    r = pn.text_frame.paragraphs[0].runs[0]._r
    fld = etree.Element(qn('a:fld')); fld.set('id', '{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}'); fld.set('type', 'slidenum')
    fld.append(copy.deepcopy(r.find(qn('a:rPr')))); etree.SubElement(fld, qn('a:t')).text = '‹#›'
    r.addprevious(fld); r.getparent().remove(r)
    els.append(pn._element)
    em = picture(f'{asset}/emblem.png', SW - 1.2 - 1.25, 0.32, 1.25, 1.25); em.name = 'Emblem'
    _, rid = LAY_TITLE.part.get_or_add_image_part(f'{asset}/emblem.png')
    eel = copy.deepcopy(em._element)
    eel.find('.//' + qn('a:blip')).set(qn('r:embed'), rid)
    els.append(eel)
    spTree = LAY_TITLE.shapes._spTree
    for k, el in enumerate(els):
        el = copy.deepcopy(el)
        el.find('.//' + qn('p:cNvPr')).set('id', str(200 + k))
        spTree.append(el)
    sldIdLst = prs.slides._sldIdLst
    sid = list(sldIdLst)[0]
    prs.part.drop_rel(sid.rId); sldIdLst.remove(sid)
    s = None
    return prs


def _notes(script, section, header):
    if script:
        s.notes_slide.notes_text_frame.text = script
    NOTES.append((len(prs.slides), section, header, script or ''))


def content_slide(num, title, section, script=None, take=None, title_size=22):
    """slide on the header layout; num like '3.1' (may be '')"""
    global s
    s = prs.slides.add_slide(LAY_TITLE)
    tf = s.shapes.title.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    runs = ([(num + '  ', 16, MUTED, False)] if num else []) + [(title, title_size, NAVY, True)]
    _fill_tf(tf, [P(*runs)])
    if take: takeaway(take)
    _notes(script, section, (num + ' ' if num else '') + title)
    return s


def image_slide(img, section, header, script=None):
    global s
    s = prs.slides.add_slide(LAY_BLANK)
    pic = picture(f'{ASSET}/{img}', 0, 0, SW, SH); pic.name = 'Background image'
    _notes(script, section, header)
    return s


def save(out):
    prs.save(out); print('saved', out, len(prs.slides), 'slides')

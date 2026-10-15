"""Re-apply the yellow highlights of an earlier Word file to a later one whose text carries the same passages.

Usage: python3 tools/reapply_marks.py manuscript_v76_ko.docx manuscript_v79_ko.docx manuscript_v79_ko_yellow.docx

v79 (the user's final thesis in the Hanyang form) has the text of v76 without its highlights; this copies every
highlighted stretch of v76 onto v79 and changes nothing else (the user's file is kept as it is, runs are only split
where a highlight begins or ends). Each highlighted v76 paragraph is matched to the v79 paragraph with the same text
after the thesis-form conversions of v77/v78 (제2.5절 -> 제2장 제5절, heading "2.5 " -> "제5절 ", 제N절 -> 제N장,
particles after four table numbers); each highlighted stretch is then found inside that paragraph.
Audit (exit 1 on failure): every highlighted stretch of v76 is found, and the highlighted characters of the output
equal those of v76 after the same conversions.
"""
import copy, re, sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
RPR_ORDER = ('rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid '
             'vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs '
             'em lang eastAsianLayout specVanish oMath').split()


def conv(t):
    """The thesis-form conversions applied to the v76 text in v77/v78."""
    t = t.replace('본 절에서는 제5절의', '본 장에서는 제5절의')
    t = re.sub(r'제(\d)절', r'제\1장', t)
    t = re.sub(r'제(\d)\.(\d)–\1\.(\d)절', r'제\1장 제\2–\3절', t)
    t = re.sub(r'제(\d)\.(\d)절', r'제\1장 제\2절', t)
    for a, b in (('다음으로 표 9은', '다음으로 표 9는'), ('표 13는 다항 로짓', '표 13은 다항 로짓'),
                 ('표 14은 핵심 기술', '표 14는 핵심 기술'), ('표 16는 본 연구의', '표 16은 본 연구의')):
        t = t.replace(a, b)
    return t


def conv_par(t):
    t = conv(t)
    t = re.sub(r'^(\d)\.(\d) ', r'제\2절 ', t)
    t = re.sub(r'^(\d)\. ', lambda m: f'제{m.group(1)}장 ', t)
    return t.replace('제1장 서론', '제1장 서 론').replace('제7장 결론', '제7장 결 론')


def runs_of(p):
    return [r for r in p.iter(qn('w:r')) if r.find(qn('w:t')) is not None]


def text_of(p):
    return ''.join(t.text or '' for r in runs_of(p) for t in r.findall(qn('w:t')))


def highlighted(p):
    """[(start, end)] of highlighted characters (adjacent runs merged)."""
    out, pos = [], 0
    for r in runs_of(p):
        n = sum(len(t.text or '') for t in r.findall(qn('w:t')))
        rPr = r.find(qn('w:rPr'))
        if rPr is not None and rPr.find(qn('w:highlight')) is not None and n:
            if out and out[-1][1] == pos:
                out[-1] = (out[-1][0], pos + n)
            else:
                out.append((pos, pos + n))
        pos += n
    return out


def set_highlight(r):
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr'); r.insert(0, rPr)
    if rPr.find(qn('w:highlight')) is not None:
        return
    h = OxmlElement('w:highlight'); h.set(qn('w:val'), 'yellow')
    rank = RPR_ORDER.index('highlight')
    for k in rPr:
        name = k.tag[len(W):] if isinstance(k.tag, str) and k.tag.startswith(W) else ''
        if name in RPR_ORDER and RPR_ORDER.index(name) > rank:
            k.addprevious(h); return
    rPr.append(h)


def split_run(r, at):
    """Split run r (one w:t) at character offset `at`; returns the second run."""
    ts = r.findall(qn('w:t'))
    assert len(ts) == 1, 'run with several text elements'
    t = ts[0]; s = t.text or ''
    r2 = copy.deepcopy(r)
    t.text = s[:at]; r2.find(qn('w:t')).text = s[at:]
    for x in (t, r2.find(qn('w:t'))):
        x.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.addnext(r2)
    return r2


def apply(p, start, end):
    pos = 0
    for r in list(runs_of(p)):
        n = sum(len(t.text or '') for t in r.findall(qn('w:t')))
        a, b = pos, pos + n
        pos = b
        if b <= start or a >= end or not n:
            continue
        if a < start:
            r = split_run(r, start - a); a = start
        if b > end:
            split_run(r, end - a); b = end
        set_highlight(r)


def main(src, dst, out):
    S = docx.Document(src); T = docx.Document(dst)
    tps = list(T.element.body.iter(qn('w:p')))
    ttexts = [text_of(p) for p in tps]
    k = 0; found = 0; want = []
    fails = 0
    for p in S.element.body.iter(qn('w:p')):
        spans = highlighted(p)
        if not spans:
            continue
        st = text_of(p)
        target = conv_par(st)
        j = next((i for i in range(k, len(tps)) if ttexts[i] == target), None)
        if j is None:
            j = next((i for i in range(len(tps)) if ttexts[i] == target), None)
        if j is None:
            print('PARAGRAPH NOT FOUND:', target[:100]); fails += 1; continue
        k = j + 1
        cursor = 0
        for a, b in spans:
            seg = conv(st[a:b])
            pos = ttexts[j].find(seg, cursor)
            if pos < 0:
                print('STRETCH NOT FOUND:', seg[:80], '| in', ttexts[j][:80]); fails += 1; continue
            apply(tps[j], pos, pos + len(seg))
            cursor = pos + len(seg); found += 1; want.append(seg)
    T.save(out)
    # audit: highlighted text of the output equals the converted highlighted text of the source
    O = docx.Document(out)
    got = []
    for p in O.element.body.iter(qn('w:p')):
        t = text_of(p)
        got += [t[a:b] for a, b in highlighted(p)]
    ok = ''.join(got) == ''.join(want)
    print(f'stretches {found}, characters {sum(map(len, want))}, output highlighted characters {sum(map(len, got))},',
          'audit OK' if ok and not fails else 'AUDIT FAILED')
    sys.exit(0 if ok and not fails else 1)


if __name__ == '__main__':
    main(*sys.argv[1:4])

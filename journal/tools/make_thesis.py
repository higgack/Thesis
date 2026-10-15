"""Hanyang University master's thesis layout (기술경영전문대학원) of the Korean manuscript (thesis track, v77+).

Usage: python3 tools/make_thesis.py 77            (add --no-render to skip the LibreOffice page-number passes)
  in : manuscript_v77_ko.md (tools/make_easy_ko.py), sources/ko_v77/08_abstract_en.md,
       forms/hanyang_mot_thesis_form.docx (the school's Word form, 3-2 학위논문 양식)
  out: thesis_v77_ko.md (the text as laid out), manuscript_v77_ko.docx (the thesis)

Built on the school's Word form (its styles, its 표제지/제출서/인준서 tables and its page set-up) and the university
guide (4-1 학위논문 작성 안내):
  order       표제지 - 제출서 - 인준서 - 차례(목차, 표 목차, 그림 목차) - 국문요지 - 본문 - 참고문헌 - 부록 - ABSTRACT
              - 국문 연구윤리서약서 - 영문 연구윤리서약서 (no acknowledgements)
  page no.    none on 표제지, 제출서, 인준서 and the two pledges; - i -, - ii - ... from 차례 to 국문요지;
              - 1 -, - 2 - ... from 제1장 to ABSTRACT (bottom centre)
  front pages the form's tables with its fonts (HY신명조 bold 16 pt, titles 20 pt); row heights fixed so that each page
              keeps the form's positions with a two-line Korean and a three-line English title; committee of three
  body pages  the form's section: A4, top/bottom 40 mm, left/right 35 mm, header/footer 15 mm; HY신명조 11 pt
              (the form: 10-11 pt), justified, first line indented two characters, line pitch 160 % (the HWP form);
              제N장 16 pt bold centred on a new page, 제M절 13 pt bold, as in the form
  ABSTRACT    English title, then "surname, given name / Dept. / Graduate School / University" (no "The")
  pledges     the university's Korean and English texts; degree, department, supervisor and name filled in; date and
              signature left for the student (the guide: signed in the submission period)
The three lists (목차, 표 목차, 그림 목차) are Word TOC fields (Word offers to update them when the file is opened);
their cached entries carry the page numbers of a LibreOffice rendering (two passes) so they read right before updating.
Table and figure entries come from TC fields holding the caption title without its parenthetical notes.
"""
import copy, json, os, pathlib, re, shutil, subprocess, sys, tempfile
sys.path.insert(0, str(pathlib.Path(__file__).parent))
import pypandoc
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, Emu, RGBColor
import docx_polish

VER = sys.argv[1] if len(sys.argv) > 1 else '77'
RENDER = '--no-render' not in sys.argv
SRC = pathlib.Path(f'sources/ko_v{VER}')
MD_IN = pathlib.Path(f'manuscript_v{VER}_ko.md')
MD_OUT = pathlib.Path(f'thesis_v{VER}_ko.md')
DOCX_OUT = pathlib.Path(f'manuscript_v{VER}_ko.docx')

# ---- thesis data (the user's instructions of 2026-10-15) ----
DEGREE_LINE = '석사학위논문'
SUBMIT = '이 논문을 공학 석사학위논문으로 제출합니다.'
APPROVE = '이 논문을 이형규의 석사학위 논문으로 인준함'
DATE = '2027년  2월'
SCHOOL = '한양대학교 기술경영전문대학원'
DEPT = '기술경영학과'
NAME_SPACED = '이 형 규'
ADVISOR_LINE = '지도교수 강 민 구'
COMMITTEE = ['심사위원장 :', '심사위원 :', '심사위원 :']   # master's committee of three
KO_PLEDGE = dict(degree='석사', dept='기술경영학과', advisor='강민구', name='이형규', to='한양대학교 기술경영전문대학원장 귀하')
EN_PLEDGE = dict(degree='Master', dept='DEPARTMENT OF TECHNOLOGY MANAGEMENT', advisor='Mingu Kang', name='LEE HYUNGKYU')

EA, LATIN, HANSI = 'HY신명조', 'HY신명조', '바탕체'   # the form's fonts
FORM = pathlib.Path('forms/hanyang_mot_thesis_form.docx')
BODY_PT, LINE_PT = 11, 17.6            # 11 pt, line pitch 160 %
TEXT_W_MM = 140                        # 210 - 2 x 35 (the form's body section)
TEXT_W_TW = int(TEXT_W_MM / 25.4 * 1440)


def fail(msg):
    print('FAIL:', msg); sys.exit(1)


# =============================================================== 1. thesis markdown
def compose_md():
    t = MD_IN.read_text(encoding='utf-8')
    titles = re.findall(r'^\*\*(.+?)\*\*$', t, re.M)[:2]
    ko_title, en_title = titles
    a = t.index('\n## 국문요지'); b = t.index('\n## 제1장')
    abstract = t[a:b].replace('\n---\n', '\n')
    r = t.index('\n## 참고문헌')
    body = t[b:r].rstrip().removesuffix('---').rstrip() + '\n'
    rest = t[r:]
    en = (SRC / '08_abstract_en.md').read_text(encoding='utf-8')
    md = abstract.strip() + '\n\n' + body + '\n' + rest.strip() + '\n\n' + en.strip() + '\n'
    md = re.sub(r'^---$', '', md, flags=re.M)
    md = re.sub(r'^## ', '# ', md, flags=re.M)
    md = re.sub(r'^### ', '## ', md, flags=re.M)
    md = re.sub(r'\n{3,}', '\n\n', md)
    # audits: nothing highlighted or red is left, no e-mail address, every heading in the thesis form
    if '{.mark}' in md or '{.red}' in md: fail('highlight left')
    if re.search(r'[\w.]+@[\w.]+', md): fail('e-mail address in the thesis')
    for h in re.findall(r'^# (.+)$', md, re.M):
        if not re.match(r'(국문요지|제\d장 .+|참고문헌|부록|ABSTRACT)$', h): fail('chapter heading ' + h)
    for h in re.findall(r'^## (.+)$', md, re.M):
        if not re.match(r'(제\d절 .+|부록 [A-Z]\. .+)$', h): fail('section heading ' + h)
    if re.search(r'제\d\.\d+절|제\d+절에서는 (이론|하이브리드)', md): fail('decimal section reference left')
    MD_OUT.write_text(md, encoding='utf-8')
    return ko_title, en_title, md


# =============================================================== 2. small XML helpers
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_fonts(rPr, latin=LATIN, ea=EA):
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = el('w:rFonts'); rPr.insert(0, rf)
    for k in list(rf.attrib):
        if k.endswith('Theme'):
            del rf.attrib[k]
    rf.set(qn('w:ascii'), latin); rf.set(qn('w:hAnsi'), HANSI if latin == LATIN else latin); rf.set(qn('w:cs'), latin)
    rf.set(qn('w:eastAsia'), ea)


def fmt_run(run, size=None, bold=None, spacing_pt=None, latin=LATIN, ea=EA, color=None):
    rPr = run._r.get_or_add_rPr()
    set_fonts(rPr, latin, ea)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if spacing_pt:
        sp = rPr.find(qn('w:spacing'))
        if sp is None:
            sp = el('w:spacing'); rPr.append(sp)
        sp.set(qn('w:val'), str(int(round(spacing_pt * 20))))
    if color: run.font.color.rgb = RGBColor.from_string(color)


def pfmt(p, align=None, line_pt=None, before=0, after=0, first_chars=None, left_tw=None, hanging_tw=None, keep=False):
    pf = p.paragraph_format
    if align is not None: pf.alignment = align
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if line_pt:
        pf.line_spacing = Pt(line_pt)
        pPr = p._p.get_or_add_pPr(); pPr.find(qn('w:spacing')).set(qn('w:lineRule'), 'exact')
    if keep: pf.keep_with_next = True
    pPr = p._p.get_or_add_pPr()
    if first_chars is not None or left_tw is not None or hanging_tw is not None:
        ind = pPr.find(qn('w:ind'))
        if ind is None:
            ind = el('w:ind'); pPr.append(ind)
        for k in list(ind.attrib): del ind.attrib[k]
        if left_tw is not None: ind.set(qn('w:left'), str(left_tw))
        if hanging_tw is not None: ind.set(qn('w:hanging'), str(hanging_tw))
        if first_chars is not None:
            ind.set(qn('w:firstLineChars'), str(first_chars * 100))
            ind.set(qn('w:firstLine'), str(int(first_chars * BODY_PT * 20)))


def style_ppr(st, align=None, line_pt=None, before=0, after=0, first_chars=None, hanging_tw=None, left_tw=None,
              page_break=False, keep=False, outline=None):
    pf = st.paragraph_format
    if align is not None: pf.alignment = align
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pPr = st.element.get_or_add_pPr()
    if line_pt:
        pf.line_spacing = Pt(line_pt); pPr.find(qn('w:spacing')).set(qn('w:lineRule'), 'exact')
    else:
        pf.line_spacing = 1.0
    pf.page_break_before = page_break
    pf.keep_with_next = keep
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = el('w:ind'); pPr.append(ind)
    for k in list(ind.attrib): del ind.attrib[k]
    ind.set(qn('w:left'), str(left_tw or 0)); ind.set(qn('w:right'), '0')
    if hanging_tw:
        ind.set(qn('w:hanging'), str(hanging_tw))
    else:
        ind.set(qn('w:firstLineChars'), str((first_chars or 0) * 100))
        ind.set(qn('w:firstLine'), str(int((first_chars or 0) * BODY_PT * 20)))
    if outline is not None:
        ol = pPr.find(qn('w:outlineLvl'))
        if ol is None:
            ol = el('w:outlineLvl'); pPr.append(ol)
        ol.set(qn('w:val'), str(outline))


def style_rpr(st, size, bold=False, color='000000'):
    rPr = st.element.get_or_add_rPr()
    set_fonts(rPr)
    for tag in ('w:color',):
        e = rPr.find(qn(tag))
        if e is not None: rPr.remove(e)
    st.font.size = Pt(size); st.font.bold = bold; st.font.italic = False
    st.font.color.rgb = RGBColor.from_string(color)


def field_runs(p, instr, cached=None):
    """Complex field inside paragraph p: begin, instruction, (separate, cached text), end."""
    def r(child):
        rr = el('w:r'); rr.append(child); p._p.append(rr); return rr
    r(el('w:fldChar', **{'w:fldCharType': 'begin'}))
    it = el('w:instrText'); it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); it.text = f' {instr} '
    r(it)
    if cached is not None:
        r(el('w:fldChar', **{'w:fldCharType': 'separate'}))
        rr = p.add_run(cached)
        r(el('w:fldChar', **{'w:fldCharType': 'end'}))
        return rr
    r(el('w:fldChar', **{'w:fldCharType': 'end'}))


def no_borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = tblPr.find(qn('w:tblBorders'))
    if b is not None: tblPr.remove(b)
    b = el('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b.append(el(f'w:{edge}', **{'w:val': 'nil'}))
    tblPr.append(b)
    mar = el('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        mar.append(el(f'w:{side}', **{'w:w': 0, 'w:type': 'dxa'}))
    tblPr.append(mar)


def fixed_table(tbl, widths_mm):
    tblPr = tbl._tbl.tblPr
    tblPr.append(el('w:tblLayout', **{'w:type': 'fixed'}))
    tw = [int(w / 25.4 * 1440) for w in widths_mm]
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = el('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(tw))); tblW.set(qn('w:type'), 'dxa')
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), tw):
        gc.set(qn('w:w'), str(w))
    for row in tbl.rows:
        for c, w in zip(row.cells, tw):
            c.width = Emu(int(w * 635))


def cell_border(cell, edge, sz=6):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = el('w:tcBorders'); tcPr.append(b)
    b.append(el(f'w:{edge}', **{'w:val': 'single', 'w:sz': sz, 'w:space': 0, 'w:color': '000000'}))


def valign(cell, v):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(el('w:vAlign', **{'w:val': v}))


# =============================================================== 3. front pages (the form's 표제지, 제출서, 인준서)
FRONT_TOP = 20.0   # mm, the form's top margin on these pages
LINE16, LINE20 = 24, 30   # exact line pitch (pt) of the 16 pt and 20 pt lines


def set_lines(p, lines, line_pt):
    """Replace the text of a form paragraph, keeping the format of its first run."""
    runs = p.runs
    keep = runs[0]
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for child in list(keep._r):
        if child.tag in (qn('w:t'), qn('w:br'), qn('w:lastRenderedPageBreak'), qn('w:tab')):
            keep._r.remove(child)
    for i, line in enumerate(lines):
        if i:
            keep._r.append(el('w:br'))
        t = el('w:t'); t.text = line; t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); keep._r.append(t)
    pfmt(p, line_pt=line_pt)


def spacer(cell):
    for q in cell.paragraphs[1:]:
        q._p.getparent().remove(q._p)
    q = cell.paragraphs[0]
    pfmt(q, line_pt=1)
    rPr = q._p.get_or_add_pPr().find(qn('w:rPr'))
    if rPr is not None:
        for tag in ('w:sz', 'w:szCs'):
            e = rPr.find(qn(tag))
            if e is None:
                e = el(tag); rPr.append(e)
            e.set(qn('w:val'), '2')


def exact(row, h_mm):
    row.height = Mm(h_mm); row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY


def layout(items):
    """items: (centre_mm, height_mm) of the text rows in order; returns the spacer heights before each text row and
    after the last one, so that each text row is centred at its page position."""
    gaps, y = [], FRONT_TOP
    for centre, h in items:
        top = centre - h / 2
        if top < y - 0.01: fail('front page rows overlap')
        gaps.append(top - y); y = top + h
    if y > 297 - 20 - 8: fail('front page overflows')
    gaps.append(2.0)
    return gaps


def fill_column_table(tbl, texts, centres):
    """tbl: the form's one-column table (spacer, text, spacer, text, ..., spacer); texts: (lines, size) per text row."""
    rows = tbl.rows
    heights = [len(lines) * (LINE20 if size == 20 else LINE16) / 72 * 25.4 + 1.5 for lines, size in texts]
    gaps = layout(list(zip(centres, heights)))
    assert len(rows) == 2 * len(texts) + 1, (len(rows), len(texts))
    for k, row in enumerate(rows):
        cell = row.cells[0]
        if k % 2 == 0:
            spacer(cell); exact(row, max(gaps[k // 2], 0.5))
        else:
            lines, size = texts[k // 2]
            p = next(q for q in cell.paragraphs if q.text.strip())
            for q in cell.paragraphs:
                if q._p is not p._p: q._p.getparent().remove(q._p)
            set_lines(p, lines, LINE20 if size == 20 else LINE16)
            exact(row, heights[k // 2])


def page_break_par(after_el):
    p = el('w:p'); pPr = el('w:pPr'); sp = el('w:spacing', **{'w:after': 0, 'w:line': 20, 'w:lineRule': 'exact'})
    pPr.append(sp); p.append(pPr)
    r = el('w:r'); rPr = el('w:rPr'); rPr.append(el('w:sz', **{'w:val': 2})); r.append(rPr)
    r.append(el('w:br', **{'w:type': 'page'})); p.append(r)
    after_el.addnext(p)
    return p


def en_title_lines(t):
    # title case for the cover; broken after "in" and after the colon so that each line fits at 20 pt
    small = {'a', 'an', 'and', 'in', 'of', 'the', 'for', 'on', 'to'}
    words = []
    for i, w in enumerate(t.split()):
        parts = w.split('-')
        parts = [p if (p.lower() in small and i and not words[-1].endswith(':')) else p[:1].upper() + p[1:] for p in parts]
        words.append('-'.join(parts))
    s = ' '.join(words)
    head, tail = s.split(': ', 1)
    hw = head.split(' in ', 1)
    return [hw[0] + ' in', hw[1] + ':', tail]


def ko_title_lines(t):
    head, tail = t.split(': ', 1)
    return [head + ':', tail]


def front_pages(ko_title, en_title):
    """The form's first section (three tables and the section-break paragraph), filled in; returns its elements."""
    form = Document(str(FORM))
    els = []
    for e in form.element.body:
        els.append(e)
        if e.tag == qn('w:p') and e.find('.//' + qn('w:sectPr')) is not None:
            break
    tbls = [e for e in els if e.tag == qn('w:tbl')]
    from docx.table import Table
    cover, submit, approve = (Table(t, form) for t in tbls)
    ko, en = ko_title_lines(ko_title), en_title_lines(en_title)
    hk = len(ko) * LINE20 / 72 * 25.4 + 1.5
    he = len(en) * LINE20 / 72 * 25.4 + 1.5
    fill_column_table(cover, [([DEGREE_LINE], 16), (ko, 20), (en, 20), ([NAME_SPACED], 16), ([SCHOOL], 16), ([DATE], 16)],
                      [49.3, 81.0 + hk / 2, 81.0 + hk + 4 + he / 2, 174.9, 216.9, 248.9])
    fill_column_table(submit, [([DEGREE_LINE], 16), (ko, 20), (en, 20), ([ADVISOR_LINE], 16), ([SUBMIT], 16), ([DATE], 16),
                               ([SCHOOL], 16), ([DEPT], 16), ([NAME_SPACED], 16)],
                      [45.0, 64.0 + hk / 2, 64.0 + hk + 4 + he / 2, 147.0, 165.0, 198.0, 221.0, 240.0, 259.0])
    # 인준서: rows sentence, date, committee (one row per member), school
    rows = approve.rows
    set_lines(next(q for q in rows[1].cells[0].paragraphs if q.text.strip()), [APPROVE], LINE16)
    set_lines(next(q for q in rows[3].cells[0].paragraphs if q.text.strip()), [DATE], LINE16)
    set_lines(next(q for q in rows[7].cells[0].paragraphs if q.text.strip()), [SCHOOL], LINE16)
    com = rows[5]._tr
    member_h = 12.7
    new_rows = []
    for i, label in enumerate(COMMITTEE):
        tr = copy.deepcopy(com)
        cells = tr.findall(qn('w:tc'))
        for ci, tc in enumerate(cells):
            ps = tc.findall(qn('w:p'))
            keep = ps[min(i, len(ps) - 1)] if ci in (1, 2, 3) else ps[0]
            for q in ps:
                if q is not keep: tc.remove(q)
        new_rows.append(tr)
    for tr in new_rows:
        com.addprevious(tr)
    com.getparent().remove(com)
    from docx.table import _Row
    rows = approve.rows
    h16 = LINE16 / 72 * 25.4 + 1.5
    centres = [52.0, 92.0]
    gaps = layout([(c, h16) for c in centres])
    exact(rows[0], gaps[0]); exact(rows[1], h16); exact(rows[2], gaps[1]); exact(rows[3], h16)
    exact(rows[4], 120.0 - (92.0 + h16 / 2))
    for i in range(len(COMMITTEE)):
        row = rows[5 + i]; exact(row, member_h)
        c = row.cells
        p = c[1].paragraphs[0]; set_lines(p, [COMMITTEE[i]], 36)
        q = c[2].paragraphs[0]
        for r in q.runs: r._r.getparent().remove(r._r)
        pfmt(q, line_pt=36); cell_border(c[2], 'bottom', 8)
        p = c[3].paragraphs[0]; pfmt(p, line_pt=36)
        for k in (1, 2, 3):
            tcPr = c[k]._tc.get_or_add_tcPr(); v = tcPr.find(qn('w:vAlign'))
            if v is not None: v.set(qn('w:val'), 'bottom')
        for k in (0, 4): spacer(c[k])
    n = 5 + len(COMMITTEE)
    y = 120.0 + member_h * len(COMMITTEE)
    exact(rows[n], 232.0 - h16 / 2 - y); exact(rows[n + 1], h16); exact(rows[n + 2], 2.0)
    for k in (0, 2, 4, n, n + 2):
        spacer(rows[k].cells[0])
    # pages: explicit breaks between the tables (the form relies on the tables filling their pages)
    out = []
    for e in els:
        if e.tag == qn('w:p') and e.find('.//' + qn('w:sectPr')) is None and not ''.join(e.itertext()).strip():
            continue                                       # the form's empty paragraphs between the tables
        out.append(e)
    holder = el('w:body')
    for e in out: holder.append(e)
    t_els = holder.findall(qn('w:tbl'))
    page_break_par(t_els[0]); page_break_par(t_els[1])
    sect_p = holder.findall(qn('w:p'))[-1]
    for r in sect_p.findall(qn('w:r')): sect_p.remove(r)
    ppr = sect_p.find(qn('w:pPr'))
    ppr.insert(0, el('w:spacing', **{'w:after': 0, 'w:line': 20, 'w:lineRule': 'exact'}))
    return list(holder)


# =============================================================== 4. pledges
def tiny_par(doc, page_break=False):
    p = doc.add_paragraph()
    pfmt(p, line_pt=1)
    r = p.add_run(); fmt_run(r, 1)
    if page_break: r.add_break(WD_BREAK.PAGE)
    return p


def pledge_box(doc, title, paras, date_line, info, sign, foot, latin=LATIN, size=12.5, title_pt=20):
    tbl = doc.add_table(rows=1, cols=1)
    no_borders(tbl); fixed_table(tbl, [170])
    cell = tbl.rows[0].cells[0]
    tbl.rows[0]._tr.get_or_add_trPr().append(el('w:cantSplit'))
    for edge in ('top', 'left', 'bottom', 'right'):
        cell_border(cell, edge, 8)
    tcPr = cell._tc.get_or_add_tcPr()
    mar = el('w:tcMar')
    for side in ('top', 'bottom'):
        mar.append(el(f'w:{side}', **{'w:w': 500, 'w:type': 'dxa'}))
    for side in ('left', 'right'):
        mar.append(el(f'w:{side}', **{'w:w': 560, 'w:type': 'dxa'}))
    tcPr.append(mar)
    p = cell.paragraphs[0]
    r = p.add_run(title); fmt_run(r, title_pt, True, latin=latin)
    pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_pt=30, before=12, after=36)
    for t in paras:
        p = cell.add_paragraph(); r = p.add_run(t); fmt_run(r, size, latin=latin)
        pfmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_pt=size * 1.7, after=18)
    p = cell.add_paragraph(); r = p.add_run(date_line); fmt_run(r, size, latin=latin, color='808080')
    pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_pt=size * 1.6, before=48, after=36)
    info_tbl = cell.add_table(rows=len(info), cols=3)
    no_borders(info_tbl); fixed_table(info_tbl, [40, 90, 20])
    info_tbl.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for i, ((k, v), row) in enumerate(zip(info, info_tbl.rows)):
        row.height = Mm(9); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for c, txt, al in ((row.cells[0], k, WD_ALIGN_PARAGRAPH.LEFT), (row.cells[1], v, WD_ALIGN_PARAGRAPH.LEFT)):
            r = c.paragraphs[0].add_run(txt); fmt_run(r, size - 1.5, latin=latin)
            pfmt(c.paragraphs[0], align=al, line_pt=(size - 1.5) * 1.6); valign(c, 'bottom')
        pfmt(row.cells[2].paragraphs[0], line_pt=(size - 1.5) * 1.6)
        if i == len(info) - 1:
            r = row.cells[2].paragraphs[0].add_run(sign); fmt_run(r, 9, latin=latin, color='808080')
            pfmt(row.cells[2].paragraphs[0], align=WD_ALIGN_PARAGRAPH.CENTER, line_pt=(size - 1.5) * 1.6)
            valign(row.cells[2], 'bottom')
    p = cell.add_paragraph()
    if foot:
        r = p.add_run(' '.join(foot.replace(' ', ''))); fmt_run(r, 15, True, latin=latin)   # evenly spaced, as in the guide
        pfmt(p, align=WD_ALIGN_PARAGRAPH.DISTRIBUTE, line_pt=26, before=56, after=6)
    else:
        pfmt(p, line_pt=6)
    return tbl


def pledges(doc):
    k = KO_PLEDGE
    pledge_box(doc, '연구 윤리 서약서', [
        '본인은 한양대학교 대학원생으로서 이 학위논문 작성 과정에서 다음과 같이 연구 윤리의 기본 원칙을 준수하였음을 서약합니다.',
        '첫째, 지도교수의 지도를 받아 정직하고 엄정한 연구를 수행하여 학위논문을 작성한다.',
        '둘째, 논문 작성시 위조, 변조, 표절 등 학문적 진실성을 훼손하는 어떤 연구 부정행위도 하지 않는다.',
        '셋째, 논문 작성시 논문유사도 검증시스템 "카피킬러"등을 거쳐야 한다.'],
        '년        월        일',   # dated and signed by the student in the submission period
        [('학위명 :', k['degree']), ('학과 :', k['dept']), ('지도교수 :', k['advisor']), ('성명 :', k['name'])],
        '(서명)', k['to'])
    tiny_par(doc, page_break=True)
    e = EN_PLEDGE
    pledge_box(doc, 'Declaration of Ethical Conduct in Research', [
        'I, as a graduate student of Hanyang University, hereby declare that I have abided by the following Code of '
        'Research Ethics while writing this dissertation thesis, during my degree program.',
        '"First, I have strived to be honest in my conduct, to produce valid and reliable research conforming with the '
        'guidance of my thesis supervisor, and I affirm that my thesis contains honest, fair and reasonable conclusions '
        'based on my own careful research under the guidance of my thesis supervisor.',
        'Second, I have not committed any acts that may discredit or damage the credibility of my research. These '
        'include, but are not limited to : falsification, distortion of research findings or plagiarism.',
        'Third, I need to go through with Copykiller Program(Internet-based Plagiarism-prevention service) before '
        'submitting a thesis."'],
        'Month        Day,        Year',
        [('Degree :', e['degree']), ('Department :', e['dept']), ('Thesis Supervisor :', e['advisor']), ('Name :', e['name'])],
        '(Signature)', None, size=12, title_pt=17)


# =============================================================== 5. lists (목차, 표 목차, 그림 목차)
def caption_title(text):
    """'표 7. 제조공정 분류 여부의 이항 로짓(종속변수 ...; ...).' -> '표 7. 제조공정 분류 여부의 이항 로짓'"""
    m = re.match(r'^((?:표|그림) \d+\.)\s*(.*)$', text, re.S)
    label, rest = m.group(1), m.group(2)
    depth, cut = 0, len(rest)
    for i, ch in enumerate(rest):          # first sentence end outside parentheses
        if ch == '(': depth += 1
        elif ch == ')': depth -= 1
        elif ch == '.' and depth == 0 and (i + 1 == len(rest) or rest[i + 1] == ' '):
            cut = i; break
    s = rest[:cut].strip()
    if s.endswith(')'):                    # drop a trailing parenthetical note
        depth = 0
        for i in range(len(s) - 1, -1, -1):
            if s[i] == ')': depth += 1
            elif s[i] == '(':
                depth -= 1
                if depth == 0:
                    s = s[:i].rstrip(); break
    return f'{label} {s}'


def toc_styles(doc):
    styles = doc.styles
    out = {}
    for name, left, bold, before in (('toc 1', 0, True, 6), ('toc 2', 2, False, 0), ('toc 3', 0, False, 0)):
        try:
            st = styles[name]
        except KeyError:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles['Normal']
        style_rpr(st, BODY_PT, bold)
        style_ppr(st, WD_ALIGN_PARAGRAPH.LEFT, LINE_PT, before, 0, left_tw=int(left * BODY_PT * 20), hanging_tw=None)
        pPr = st.element.get_or_add_pPr()
        tabs = pPr.find(qn('w:tabs'))
        if tabs is not None: pPr.remove(tabs)
        tabs = el('w:tabs'); tabs.append(el('w:tab', **{'w:val': 'right', 'w:leader': 'dot', 'w:pos': TEXT_W_TW}))
        pPr.insert(0, tabs)
        out[name] = st
    try:
        st = styles['목차 제목']
    except KeyError:
        st = styles.add_style('목차 제목', WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = styles['Normal']
    style_rpr(st, 16, True); style_ppr(st, WD_ALIGN_PARAGRAPH.CENTER, None, 0, 30)
    out['title'] = st
    return out


def toc_block(doc, anchor, title, instr, entries, st, page_break=True):
    """Insert before `anchor` (an element): title paragraph and a TOC field whose cached result lists entries
    [(style name, text, page label)]."""
    made = []
    p = doc.add_paragraph(title, style=st['title'])
    if page_break: p.paragraph_format.page_break_before = True
    made.append(p)
    for i, (sname, text, page) in enumerate(entries):
        q = doc.add_paragraph(style=st[sname])
        if i == 0:
            for child in (el('w:fldChar', **{'w:fldCharType': 'begin'}),):
                rr = el('w:r'); rr.append(child); q._p.append(rr)
            it = el('w:instrText'); it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); it.text = f' {instr} '
            rr = el('w:r'); rr.append(it); q._p.append(rr)
            rr = el('w:r'); rr.append(el('w:fldChar', **{'w:fldCharType': 'separate'})); q._p.append(rr)
        q.add_run(text); q.add_run('\t' + page)
        if i == len(entries) - 1:
            rr = el('w:r'); rr.append(el('w:fldChar', **{'w:fldCharType': 'end'})); q._p.append(rr)
        made.append(q)
    for m in made:
        anchor.addprevious(m._p)


# =============================================================== 6. sections and page numbers
def sect_pr(kind):
    """kind: 'front' (no number), 'roman', 'arabic', 'pledge' (no number)."""
    s = el('w:sectPr')
    s.append(el('w:type', **{'w:val': 'nextPage'}))
    s.append(el('w:pgSz', **{'w:w': 11906, 'w:h': 16838}))
    mm = lambda x: int(round(x / 25.4 * 1440))
    if kind == 'front':
        mar = dict(top=20, bottom=20, left=12.5, right=12.5, header=10, footer=10)
    elif kind == 'pledge':
        mar = dict(top=30, bottom=25, left=20, right=20, header=10, footer=10)
    else:   # the form's body section
        mar = dict(top=40, bottom=40, left=35, right=35, header=15, footer=15)
    s.append(el('w:pgMar', **{'w:top': mm(mar['top']), 'w:right': mm(mar['right']), 'w:bottom': mm(mar['bottom']),
                              'w:left': mm(mar['left']), 'w:header': mm(mar['header']), 'w:footer': mm(mar['footer']),
                              'w:gutter': 0}))
    if kind == 'roman':
        s.append(el('w:pgNumType', **{'w:fmt': 'lowerRoman', 'w:start': 1}))
    elif kind == 'arabic':
        s.append(el('w:pgNumType', **{'w:fmt': 'decimal', 'w:start': 1}))
    s.append(el('w:cols', **{'w:space': 720}))
    return s


def end_section(p, kind):
    pPr = p._p.get_or_add_pPr()
    pPr.append(sect_pr(kind))


def footers(doc, kinds):
    for sec, kind in zip(doc.sections, kinds):
        ft = sec.footer
        ft.is_linked_to_previous = False
        p = ft.paragraphs[0]
        for r in list(p.runs): r._r.getparent().remove(r._r)
        if kind in ('roman', 'arabic'):
            r = p.add_run('- '); fmt_run(r, 10)
            field_runs(p, 'PAGE', 'i' if kind == 'roman' else '1')
            for r in p.runs: fmt_run(r, 10)
            r = p.add_run(' -'); fmt_run(r, 10)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hd = sec.header
        hd.is_linked_to_previous = False


# =============================================================== 7. the document
def body_styles(doc):
    styles = doc.styles
    # the form defines only its own styles; add the ones pandoc's paragraphs refer to
    for name, base in (('Body Text', 'Normal'), ('First Paragraph', 'Body Text'), ('Compact', 'Body Text'),
                       ('Heading 1', 'Normal'), ('Heading 2', 'Normal')):
        try:
            styles[name]
        except KeyError:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH); st.base_style = styles[base]
            if name.startswith('Heading'):
                st.next_paragraph_style = styles['Body Text']
    try:
        styles['Table']
    except KeyError:
        styles.add_style('Table', WD_STYLE_TYPE.TABLE).base_style = styles['Normal Table']
    # document defaults: the form's fonts (HY신명조, 바탕체 for other Latin characters), black
    rpr_default = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    set_fonts(rpr_default)
    for name in ('Normal', 'Body Text', 'First Paragraph'):
        st = styles[name]; style_rpr(st, BODY_PT)
        style_ppr(st, WD_ALIGN_PARAGRAPH.JUSTIFY, LINE_PT, 0, 0, first_chars=2)
    st = styles['Normal']; style_ppr(st, WD_ALIGN_PARAGRAPH.JUSTIFY, None, 0, 0, first_chars=0)
    st = styles['Compact']; style_rpr(st, 9); style_ppr(st, None, None, 0, 0, first_chars=0)
    h1 = styles['Heading 1']; style_rpr(h1, 16, True)
    style_ppr(h1, WD_ALIGN_PARAGRAPH.CENTER, None, 0, 36, page_break=True, keep=True, outline=0)
    h2 = styles['Heading 2']; style_rpr(h2, 13, True)
    style_ppr(h2, WD_ALIGN_PARAGRAPH.LEFT, None, 18, 10, keep=True, outline=1)
    for name in ('Heading 1', 'Heading 2'):
        rPr = styles[name].element.get_or_add_rPr()
        for tag in ('w:i', 'w:iCs'):
            e = rPr.find(qn(tag))
            if e is not None: rPr.remove(e)


def tc_field(p, text, kind):
    rr = el('w:r'); rr.append(el('w:fldChar', **{'w:fldCharType': 'begin'})); p._p.append(rr)
    it = el('w:instrText'); it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    it.text = ' TC "%s" \\f %s \\l 3 ' % (text.replace('"', "'"), kind)
    rr = el('w:r'); rr.append(it); p._p.append(rr)
    rr = el('w:r'); rr.append(el('w:fldChar', **{'w:fldCharType': 'end'})); p._p.append(rr)


def build(ko_title, en_title, pages):
    tmp = pathlib.Path(tempfile.mkdtemp())
    raw = tmp / 'raw.docx'
    pypandoc.convert_file(str(MD_OUT), 'docx', outputfile=str(raw),
                          extra_args=['--resource-path=.:figures', '--dpi=300', f'--reference-doc={FORM}'])
    doc = Document(str(raw))
    body_styles(doc)
    st = toc_styles(doc)
    # tables as in the journal build (column widths, 9/8.5/8 pt, rules) on the 140 mm text block
    for tbl in doc.tables:
        docx_polish.polish_table(tbl, TEXT_W_TW)
    paras = doc.paragraphs
    heads, caps = [], []
    zone = None
    for i, p in enumerate(paras):
        txt = p.text.strip()
        sname = p.style.name
        if sname == 'Heading 1':
            zone = txt; heads.append((1, txt, p)); continue
        if sname == 'Heading 2':
            heads.append((2, txt, p)); continue
        if p._p.xpath('.//w:drawing'):
            pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=4, first_chars=0, keep=True)
            p.paragraph_format.line_spacing = 1.0
            for shp in p._p.xpath('.//wp:extent'):
                cx, cy = int(shp.get('cx')), int(shp.get('cy'))
                maxw = int(TEXT_W_MM * 36000)
                if cx > maxw:
                    shp.set('cx', str(maxw)); shp.set('cy', str(int(cy * maxw / cx)))
                    for a in p._p.xpath('.//a:ext'):
                        a.set('cx', str(maxw)); a.set('cy', str(int(cy * maxw / cx)))
            continue
        if docx_polish.CAPTION_RE.match(txt):
            is_tab = txt.startswith('표 ') and re.match(r'^표 \d+\.', txt)
            is_fig = re.match(r'^그림 \d+\.', txt)
            docx_polish._set_run_size(p, 10)
            pfmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_pt=15, before=8 if (is_tab or txt.startswith('표 ')) else 0,
                 after=4 if txt.startswith('표 ') else 8, first_chars=0, keep=bool(txt.startswith('표 ')))
            if is_tab or is_fig:
                title = caption_title(txt)
                tc_field(p, title, 't' if is_tab else 'f')
                caps.append(('t' if is_tab else 'f', title, p))
            continue
        if zone == '참고문헌' and txt:
            pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_pt=15, after=3, left_tw=400, hanging_tw=400)
            for r in p.runs: r.font.size = Pt(10)
            continue
        if txt.startswith('주제어:') or txt.startswith('Keywords:'):
            pfmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_pt=LINE_PT, before=12, first_chars=0)
            continue
        if zone == 'ABSTRACT':
            for r in p.runs: fmt_run(r)
            if p.runs and p.runs[0].bold and txt == p.runs[0].text.strip() and 'Convergence' in txt:
                pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_pt=20, before=0, after=18, first_chars=0)
                for r in p.runs: fmt_run(r, 13, True)
            elif txt.startswith('Lee, HyungKyu'):
                pfmt(p, align=WD_ALIGN_PARAGRAPH.RIGHT, line_pt=16, after=24, first_chars=0)
    # tables: cells keep their own (compact) paragraph format
    for tbl in doc.tables:
        for row in tbl.rows:
            for c in row.cells:
                for q in c.paragraphs:
                    ind = q._p.get_or_add_pPr()
                    e = ind.find(qn('w:ind'))
                    if e is not None: ind.remove(e)
                    ind.append(el('w:ind', **{'w:firstLine': 0, 'w:firstLineChars': 0, 'w:left': 0}))
                    if ind.find(qn('w:jc')) is None:
                        ind.append(el('w:jc', **{'w:val': 'left'}))   # the form's default is justified
    body = doc.element.body
    first = body[0]
    # ---- the form's front pages, ending with its section break (no page numbers)
    for e in front_pages(ko_title, en_title):
        first.addprevious(e)
    # ---- lists before 국문요지
    gukmun = next(p for lv, t, p in heads if t == '국문요지')
    gukmun_page = pages.get('h:국문요지', '?')
    toc_entries = [('toc 1' if lv == 1 else 'toc 2', t, pages.get(f'h:{t}', '?')) for lv, t, p in heads]
    t_entries = [('toc 3', t, pages.get(f'c:{t}', '?')) for k, t, p in caps if k == 't']
    f_entries = [('toc 3', t, pages.get(f'c:{t}', '?')) for k, t, p in caps if k == 'f']
    toc_block(doc, gukmun._p, '목     차', 'TOC \\o "1-2" \\h \\z \\u', toc_entries, st, page_break=False)
    toc_block(doc, gukmun._p, '표  목  차', 'TOC \\h \\z \\f t', t_entries, st)
    toc_block(doc, gukmun._p, '그 림 목 차', 'TOC \\h \\z \\f f', f_entries, st)
    # ---- section ends: after 국문요지 (roman) and after ABSTRACT (arabic)
    ch1 = next(p for lv, t, p in heads if t.startswith('제1장'))
    prev = ch1._p.getprevious()
    sp = doc.add_paragraph(); pfmt(sp, line_pt=1); fmt_run(sp.add_run(), 1)
    ch1._p.addprevious(sp._p)
    end_section(sp, 'roman')
    # chapter 1 starts the new section: no extra page break needed there
    ch1.paragraph_format.page_break_before = False
    sp2 = doc.add_paragraph(); pfmt(sp2, line_pt=1); fmt_run(sp2.add_run(), 1)
    end_section(sp2, 'arabic')
    # ---- pledges (last section, no page numbers)
    pledges(doc)
    tiny_par(doc)
    final = body.find(qn('w:sectPr'))
    body.remove(final)
    body.append(sect_pr('pledge'))
    footers(doc, ['front', 'roman', 'arabic', 'pledge'])
    # Word: offer to update the fields (the lists) when the file is opened
    update_fields_on_open(doc)
    # the form's compatibility option "balance SBCS and DBCS characters" sets Latin letters, digits and spaces on a
    # half-width grid, which makes the spacing look uneven in Word (the user's remark on v79); v80+ switch it off
    comp = doc.settings.element.find(qn('w:compat'))
    if comp is not None:
        for e in comp.findall(qn('w:balanceSingleByteDoubleByteWidth')):
            comp.remove(e)
    fix_order(doc.element.body); fix_order(doc.styles.element)
    for sec in doc.sections:
        fix_order(sec.footer._element); fix_order(sec.header._element)
    doc.core_properties.title = ko_title
    doc.core_properties.author = '이형규'
    doc.save(str(DOCX_OUT))
    shutil.rmtree(tmp)
    return heads, caps


# =============================================================== 8. schema order (Word rejects misordered properties)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
ORDER = {
    'pPr': 'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs '
           'suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd '
           'snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment '
           'textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange',
    'rPr': 'rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid '
           'vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em '
           'lang eastAsianLayout specVanish oMath',
    'tblPr': 'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd '
             'tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription',
    'tcPr': 'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark',
    'sectPr': 'headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType '
              'pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings '
              'sectPrChange',
    'tcMar': 'top start left bottom end right', 'tblCellMar': 'top start left bottom end right',
    'tcBorders': 'top start left bottom end right insideH insideV tl2br tr2bl',
    'tblBorders': 'top start left bottom end right insideH insideV',
}
ORDER = {k: {n: i for i, n in enumerate(v.split())} for k, v in ORDER.items()}


def fix_order(root):
    for e in root.iter():
        if not isinstance(e.tag, str) or not e.tag.startswith(W):
            continue
        name = e.tag[len(W):]
        rank = ORDER.get(name)
        if rank is None or len(e) < 2:
            continue
        kids = list(e)
        if name == 'rPr' and e.getparent() is not None and e.getparent().tag == W + 'rPrChange':
            continue
        keyed = sorted(kids, key=lambda k: rank.get(k.tag[len(W):] if isinstance(k.tag, str) and k.tag.startswith(W) else '', 999))
        if keyed != kids:
            for k in kids: e.remove(k)
            for k in keyed: e.append(k)


def update_fields_on_open(doc):
    st = doc.settings.element
    if st.find(qn('w:updateFields')) is not None:
        return
    uf = el('w:updateFields', **{'w:val': 'true'})
    later = ('hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids', 'mathPr', 'attachedSchema',
             'themeFontLang', 'clrSchemeMapping', 'doNotIncludeSubdocsInStats', 'doNotAutoCompressPictures', 'forceUpgrade',
             'captions', 'readModeInkLockDown', 'smartTagType', 'schemaLibrary', 'shapeDefaults', 'doNotEmbedSmartTags',
             'decimalSymbol', 'listSeparator')
    for k in st:
        if isinstance(k.tag, str) and k.tag.split('}')[-1] in later:
            k.addprevious(uf); return
    st.append(uf)


# =============================================================== 9. rendering and page numbers
def render(docx_path, outdir):
    env = dict(os.environ)
    conf = outdir / 'fonts.conf'
    conf.write_text('''<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>
<include ignore_missing="yes">/etc/fonts/fonts.conf</include>
<alias binding="same"><family>HY신명조</family><prefer><family>Noto Serif CJK KR</family></prefer></alias>
<alias binding="same"><family>바탕체</family><prefer><family>Noto Serif CJK KR</family></prefer></alias>
<alias binding="same"><family>바탕</family><prefer><family>Noto Serif CJK KR</family></prefer></alias>
<alias binding="same"><family>Batang</family><prefer><family>Noto Serif CJK KR</family></prefer></alias>
<alias binding="same"><family>Times New Roman</family><prefer><family>Liberation Serif</family></prefer></alias>
</fontconfig>''', encoding='utf-8')
    env['FONTCONFIG_FILE'] = str(conf)
    prof = outdir / 'lo_profile'
    subprocess.run(['soffice', f'-env:UserInstallation=file://{prof}', '--headless', '--convert-to', 'pdf',
                    '--outdir', str(outdir), str(docx_path)], check=True, env=env, capture_output=True, timeout=600)
    return outdir / (docx_path.stem + '.pdf')


def page_map(pdf, heads, caps):
    import pymupdf
    d = pymupdf.open(str(pdf))
    pages = []
    for pg in d:
        lines = [l.strip() for l in pg.get_text().split('\n') if l.strip()]
        label = None
        for l in lines[-3:]:
            m = re.match(r'^-\s*([ivxlc]+|\d+)\s*-$', l)
            if m: label = m.group(1)
        pages.append((label, lines, ''.join(lines).replace(' ', '')))
    out, k = {}, 0
    norm = lambda s: s.replace(' ', '')
    # start after the lists: first page whose first line is 국문요지
    start = next(i for i, (lab, lines, _) in enumerate(pages) if lines and norm(lines[0]) == '국문요지')
    k = start
    for lv, t, p in heads:
        while k < len(pages):
            lab, lines, flat = pages[k]
            if (lv == 1 and lines and norm(lines[0]) == norm(t)) or (lv == 2 and norm(t) in flat):
                break
            k += 1
        if k == len(pages): fail(f'heading not found in the rendering: {t}')
        out[f'h:{t}'] = pages[k][0] or '?'
    for kind, t, p in caps:
        label = t.split(' ', 2)[0] + ' ' + t.split(' ', 2)[1]   # '표 7.'
        hit = [i for i in range(start, len(pages)) if any(norm(l).startswith(norm(label)) for l in pages[i][1])]
        if not hit: fail(f'caption not found: {t}')
        out[f'c:{t}'] = pages[hit[0]][0] or '?'
    return out, len(d), [p[0] for p in pages]


def main():
    ko_title, en_title, md = compose_md()
    pages = {}
    if not RENDER:
        build(ko_title, en_title, pages); print('built without page numbers ->', DOCX_OUT); return
    work = pathlib.Path(tempfile.mkdtemp())
    for it in range(3):
        heads, caps = build(ko_title, en_title, pages)
        pdf = render(DOCX_OUT, work)
        new, n, labels = page_map(pdf, heads, caps)
        print(f'pass {it + 1}: {n} pages; labels', labels[:6], '...', labels[-6:])
        if new == pages:
            break
        pages = new
    else:
        heads, caps = build(ko_title, en_title, pages)
    keep = work / f'thesis_v{VER}_ko_preview.pdf'   # LibreOffice preview with substitute fonts (not part of the thesis)
    shutil.copy(pdf, keep)
    json.dump(pages, open(work / 'pages.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(json.dumps(pages, ensure_ascii=False, indent=0)[:3000])
    print('->', DOCX_OUT, '| preview', keep)


if __name__ == '__main__':
    main()

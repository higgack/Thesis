# -*- coding: utf-8 -*-
"""A1 Research Day poster (v1) built on the MOT A1 template.
Usage: python3 build_poster.py TEMPLATE.pptx ASSET_DIR ICON_DIR OUT.pptx"""
import sys, copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.oxml.ns import qn
from lxml import etree

TPL, ASSET, ICON, OUT = sys.argv[1:5]
FONT = '맑은 고딕'
NAVY, NAVY2 = '13294B', '24406E'
PROC, INTEG, BOND, CORE = '2563C9', '7C3AED', 'EA7A28', 'B5441B'
PROC_L, INTEG_L, BOND_L = 'E3ECFB', 'EEE6FD', 'FDEBDD'
YEL, TEXT, MUTED, PANEL, WHITE = 'FFC000', '1E2532', '5B6475', 'F2F5FA', 'FFFFFF'
TEAL, GOLD = '10A39A', 'E0A100'

def cm(v): return Emu(int(round(v * 360000)))
def rgb(h): return RGBColor.from_string(h)

prs = Presentation(TPL)
s = prs.slides[0]
# drop the empty section placeholders; keep title/authors/affiliation (153/151/150) and the P box
for sh in list(s.shapes):
    if sh.is_placeholder and sh.placeholder_format.idx not in (150, 151, 153):
        sh._element.getparent().remove(sh._element)
P_BOX = [sh for sh in s.shapes if sh.name == '직사각형 1'][0]

def _nostyle(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None: sh._element.remove(st)

def font_run(r, size, color=TEXT, bold=False, italic=False):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', FONT)
    rPr.set('lang', 'ko-KR')

def set_para(p, align=PP_ALIGN.LEFT, space_after=0, line=None, bullet=None, indent=0.0, space_before=0):
    p.alignment = align
    p.space_after = Pt(space_after); p.space_before = Pt(space_before)
    if line: p.line_spacing = line
    pPr = p._p.get_or_add_pPr()
    for t in ('a:buNone', 'a:buChar', 'a:buFont', 'a:buClr'):
        for el in pPr.findall(qn(t)): pPr.remove(el)
    if bullet:
        ch, col = bullet
        pPr.set('marL', str(int(cm(indent)))); pPr.set('indent', str(-int(cm(indent))))
        bc = etree.SubElement(pPr, qn('a:buClr')); sc = etree.SubElement(bc, qn('a:srgbClr')); sc.set('val', col)
        bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', 'Arial')
        b = etree.SubElement(pPr, qn('a:buChar')); b.set('char', ch)
    else:
        pPr.set('marL', '0'); pPr.set('indent', '0')
        etree.SubElement(pPr, qn('a:buNone'))

def text(x, y, w, h, paras, anchor='t', margin=0.0, wrap=True, autofit=False, name=None):
    """paras: list of dicts {runs:[(text,size,color,bold)], align, after, line, bullet, indent}"""
    tb = s.shapes.add_textbox(cm(x), cm(y), cm(w), cm(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = cm(margin)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    bp = tf._txBody.find(qn('a:bodyPr'))
    for el in list(bp): bp.remove(el)
    etree.SubElement(bp, qn('a:noAutofit'))
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_para(p, pd.get('align', PP_ALIGN.LEFT), pd.get('after', 0), pd.get('line'), pd.get('bullet'), pd.get('indent', 0), pd.get('before', 0))
        for rr in pd['runs']:
            t, size, col, bold = norm(rr)
            r = p.add_run(); r.text = t; font_run(r, size, col, bold)
    return tb

def norm(rr):
    rr = tuple(rr)
    return (rr[0], rr[1] if len(rr) > 1 else 24, rr[2] if len(rr) > 2 else TEXT, rr[3] if len(rr) > 3 else False)

def P(*runs, **kw):
    d = {'runs': [r if isinstance(r, tuple) else (r, kw.get('size', 24), kw.get('color', TEXT), kw.get('bold', False)) for r in runs]}
    d.update({k: v for k, v in kw.items() if k not in ('size', 'color', 'bold')}); return d

def shape(kind, x, y, w, h, fill=None, line=None, lw=0.0, radius=None, shadow=False, alpha=None, name=None):
    sh = s.shapes.add_shape(kind, cm(x), cm(y), cm(w), cm(h)); _nostyle(sh)
    if name: sh.name = name
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = sh.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            a = etree.SubElement(clr, qn('a:alpha')); a.set('val', str(int(alpha * 100000)))
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if shadow:
        spPr = sh._element.spPr
        eff = etree.SubElement(spPr, qn('a:effectLst'))
        sd = etree.SubElement(eff, qn('a:outerShdw'))
        for k, v in (('blurRad', '190500'), ('dist', '50800'), ('dir', '5400000'), ('algn', 't'), ('rotWithShape', '0')): sd.set(k, v)
        c = etree.SubElement(sd, qn('a:srgbClr')); c.set('val', '1E2532')
        a = etree.SubElement(c, qn('a:alpha')); a.set('val', '18000')
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = cm(0.15)
    return sh

def shape_text(sh, paras, anchor='m', margin=0.2):
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = cm(margin); tf.margin_top = tf.margin_bottom = cm(0.1)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_para(p, pd.get('align', PP_ALIGN.CENTER), pd.get('after', 0), pd.get('line'), pd.get('bullet'), pd.get('indent', 0))
        for (t, size, col, bold) in map(norm, pd['runs']):
            r = p.add_run(); r.text = t; font_run(r, size, col, bold)
    return sh

def line(x1, y1, x2, y2, color, w=2.0, dash=None, head=None, tail=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cm(x1), cm(y1), cm(x2), cm(y2)); _nostyle(c)
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', dash)
    for tag, kind in (('a:headEnd', head), ('a:tailEnd', tail)):
        if kind:
            e = etree.SubElement(ln, qn(tag)); e.set('type', kind); e.set('w', 'med'); e.set('len', 'med')
    return c

def icon(name, x, y, size, color='FFFFFF'):
    return s.shapes.add_picture(f'{ICON}/{name}_{color}.png', cm(x), cm(y), cm(size), cm(size))

def icon_circle(name, x, y, d, fill, pad=0.26):
    shape(MSO_SHAPE.OVAL, x, y, d, d, fill=fill)
    icon(name, x + d * pad, y + d * pad, d * (1 - 2 * pad))

def section(num, title, x, y, w, h, icon_name, sub=None):
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=PANEL, radius=0.025)
    shape_text(shape(MSO_SHAPE.OVAL, x + 0.7, y + 0.6, 2.3, 2.3, fill=NAVY), [P(('%02d' % num, 30, WHITE, True))], margin=0)
    runs = [(title, 40, NAVY, True)]
    if sub: runs.append(('   ' + sub, 24, MUTED, False))
    text(x + 3.4, y + 0.55, w - 4.2, 2.4, [P(*runs)], anchor='m')

def style_chart(ch, legend=True, size=16):
    ch.font.size = Pt(size); ch.font.name = FONT
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size); ch.legend.font.name = FONT
    for ax in (ch.category_axis, ch.value_axis):
        ax.tick_labels.font.size = Pt(size); ax.tick_labels.font.name = FONT
        ax.format.line.color.rgb = rgb('9AA3B2')
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = rgb('DDE2EA')
    ch.value_axis.format.line.fill.background()
    for d in ch._chartSpace.iter(qn('a:defRPr')):
        if d.find(qn('a:ea')) is None:
            ea = etree.SubElement(d, qn('a:ea')); ea.set('typeface', FONT)

# ---------------------------------------------------------------- layout grid
import json, math
W, H = 59.41, 84.10
X0, X1 = 1.6, 57.8
CW = X1 - X0
COLW = (CW - 1.2) / 2
XL, XR = X0, X0 + COLW + 1.2
GAP = 0.7
HDR_H, KEY_H, A_H, B_H, C_H, D_H = 11.4, 3.4, 17.0, 15.0, 19.6, 10.0
Y_KEY = HDR_H + GAP
Y_A = Y_KEY + KEY_H + GAP
Y_B = Y_A + A_H + GAP
Y_C = Y_B + B_H + GAP
Y_D = Y_C + C_H + GAP
assert Y_D + D_H <= 80.1, Y_D + D_H

# ---------------------------------------------------------------- header band
hdr = shape(MSO_SHAPE.RECTANGLE, 0, 0, W, HDR_H, fill=NAVY, name='Header')
s.shapes._spTree.remove(hdr._element); s.shapes._spTree.insert(2, hdr._element)
hx, hw = 45.6, 12.2                               # stacked dies joined by hybrid-bonded Cu pad layers
dies = ['8FA6CC', '7D96BF', '6B86B2', '5976A5']
yy = 4.0
for k, col in enumerate(dies):
    shape(MSO_SHAPE.RECTANGLE, hx + 0.6, yy, hw - 1.2, 0.95, fill=col)
    yy += 0.95
    shape(MSO_SHAPE.RECTANGLE, hx + 0.6, yy, hw - 1.2, 0.32, fill='BFE3E6')
    for m in range(9):
        shape(MSO_SHAPE.RECTANGLE, hx + 1.0 + m * (hw - 2.0 - 0.55) / 8, yy, 0.55, 0.32, fill=BOND)
    yy += 0.32
shape(MSO_SHAPE.RECTANGLE, hx, yy, hw, 1.05, fill='3E5A8A')
text(hx, yy + 1.15, hw, 0.7, [P(('하이브리드 본딩 적층 구조 (개념도)', 13, '9FB0CC'), align=PP_ALIGN.CENTER)])
s.shapes._spTree.remove(P_BOX._element); s.shapes._spTree.append(P_BOX._element)

def fill_ph(idx, x, y, w, h, paras, anchor='t'):
    ph = [sh for sh in s.placeholders if sh.placeholder_format.idx == idx][0]
    ph.left, ph.top, ph.width, ph.height = cm(x), cm(y), cm(w), cm(h)
    tf = ph.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    bp = tf._txBody.find(qn('a:bodyPr'))
    for el in list(bp): bp.remove(el)
    etree.SubElement(bp, qn('a:noAutofit'))
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_para(p, pd.get('align', PP_ALIGN.LEFT), pd.get('after', 0), pd.get('line'), None, 0)
        for (t, size, col, bold) in map(norm, pd['runs']):
            r = p.add_run(); r.text = t; font_run(r, size, col, bold)

fill_ph(153, X0, 0.9, 43.0, 5.6, [
    P(('반도체 하이브리드 본딩 특허를 통한', 58, WHITE, True), line=0.92),
    P(('공정 단계 간 기술융합', 58, YEL, True), ('의 동태적 분석', 58, WHITE, True), line=0.92)])
fill_ph(151, X0, 6.75, 43.0, 1.2, [
    P(('A Dynamic Analysis of Cross-Stage Technological Convergence through Semiconductor Hybrid Bonding Patents', 19, 'B9C6DD'))])
fill_ph(150, X0, 8.3, 43.0, 2.2, [
    P(('이형규', 34, WHITE, True), ('   한양대학교 기술경영학과', 26, 'DCE3EE'))], anchor='m')

# ---------------------------------------------------------------- key message
shape(MSO_SHAPE.ROUNDED_RECTANGLE, X0, Y_KEY, CW, KEY_H, fill='FFF6D6', radius=0.14)
icon_circle('FaLightbulb', X0 + 0.7, Y_KEY + 0.5, 2.5, YEL, pad=0.24)
text(X0 + 3.9, Y_KEY + 0.15, CW - 4.6, KEY_H - 0.3, [
    P(('하이브리드 본딩 관련 특허출원 928건을 분석한 결과, 공정 단계 간 기술융합은 한 방향으로 확대되지 않았다. ', 27, NAVY, True),
      ('형성기 → 확산기 → 재편기', 27, INTEG, True), ('의 세 국면을 거치며 융합의 중심이 ', 27, NAVY, True),
      ('핵심 접합 기술에서 주변 공정으로', 27, CORE, True), (' 이동하였다.', 27, NAVY, True), line=1.0)], anchor='m')

# ---------------------------------------------------------------- 01 background
section(1, '연구 배경', XL, Y_A, COLW, A_H, None)
bul = ('•', NAVY)
text(XL + 0.9, Y_A + 3.05, COLW - 1.8, 7.8, [
    P(('반도체 산업은 ', 19), ('전공정', 19, PROC, True), ('(웨이퍼 위 소자 형성)과 ', 19), ('후공정', 19, BOND, True),
      ('(절단·조립·패키징)의 분업 구조를 반세기 이상 유지해 왔다.', 19), bullet=bul, indent=0.6, after=4),
    P(('미세화의 한계로 다이를 분할·적층·연결하는 첨단 패키징이 새로운 성능 향상의 원천이 되었다.', 19), bullet=bul, indent=0.6, after=4),
    P(('하이브리드 본딩', 19, TEXT, True), ('은 솔더 범프 없이 구리 패드와 산화막을 직접 접합하여 10 µm 이하의 연결 간격을 구현한다.', 19), bullet=bul, indent=0.6, after=4),
    P(('그 성패를 좌우하는 평탄화(CMP)·표면 활성화·정밀 정렬은 ', 19), ('전공정의 제조공정 역량', 19, PROC, True), ('에 해당한다.', 19), bullet=bul, indent=0.65)])

def die_stack(x, y, w, hybrid):
    die_h, gap = 0.9, (0.1 if hybrid else 0.8)
    shape(MSO_SHAPE.RECTANGLE, x, y, w, die_h, fill='8A97AD')
    by = y + die_h + gap
    if hybrid:
        oy = y + die_h - 0.38
        shape(MSO_SHAPE.RECTANGLE, x, oy, w, 0.38, fill='BFE3E6')
        shape(MSO_SHAPE.RECTANGLE, x, by, w, 0.38, fill='BFE3E6')
        n = 6
        for k in range(n):
            px = x + 0.45 + k * (w - 0.9 - 0.7) / (n - 1)
            shape(MSO_SHAPE.RECTANGLE, px, oy, 0.7, 0.38, fill=BOND)
            shape(MSO_SHAPE.RECTANGLE, px, by, 0.7, 0.38, fill=BOND)
        shape(MSO_SHAPE.RECTANGLE, x, by - gap, w, gap, fill=INTEG)
        by += 0.38
    else:
        n = 5
        for k in range(n):
            px = x + 0.55 + k * (w - 1.1 - 0.85) / (n - 1)
            shape(MSO_SHAPE.OVAL, px, y + die_h - 0.02, 0.85, gap + 0.04, fill='A9AFBA', line='7D8594', lw=0.75)
    shape(MSO_SHAPE.RECTANGLE, x, by, w, 0.9, fill='5E6C85')

sy, sh_ = Y_A + 11.0, A_H - 11.0 - 0.6
cw2 = (COLW - 1.8 - 1.2) / 2
for k, hyb in enumerate((False, True)):
    cx = XL + 0.9 + k * (cw2 + 1.2)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, sy, cw2, sh_, fill=WHITE, radius=0.06, shadow=True)
    text(cx + 0.3, sy + 0.2, cw2 - 0.6, 0.9, [P((('하이브리드 본딩' if hyb else '기존: 솔더 범프 접합'), 20, (INTEG if hyb else MUTED), True), align=PP_ALIGN.CENTER)])
    die_stack(cx + 0.9, sy + 1.1, cw2 - 1.8, hyb)
    text(cx + 0.3, sy + sh_ - 1.45, cw2 - 0.6, 1.3, [
        P((('구리–구리 + 산화막–산화막 직접 접합' if hyb else '다이 사이를 솔더 범프로 연결'), 16, TEXT, True), align=PP_ALIGN.CENTER),
        P((('범프 없이 10 µm 이하 간격' if hyb else '10 µm 이하 간격은 구현하기 어려움'), 16, MUTED), align=PP_ALIGN.CENTER)])
line(XL + 0.9 + cw2 + 0.15, sy + sh_ / 2, XL + 0.9 + cw2 + 1.05, sy + sh_ / 2, INTEG, w=4, tail='triangle')

# ---------------------------------------------------------------- 02 research questions
section(2, '연구 질문', XR, Y_A, COLW, A_H, None)
text(XR + 0.9, Y_A + 3.1, COLW - 1.8, 3.4, [
    P(('기존 기술융합 연구는 주로 ', 20), ('서로 다른 산업·기술 분야 사이', 20, TEXT, True),
      ('의 융합(ICT–가전, 식품–제약, 나노–바이오)을 다루었다. 반면 ', 20),
      ('한 산업 안에서 공정 단계를 가로지르는 융합', 20, INTEG, True), ('과 ', 20),
      ('융합 내용의 시간적 변화', 20, INTEG, True), ('는 실증 연구가 부족하다.', 20))])
rqs = [('RQ1', '수준과 유형', '제조공정과 접합 기술 간의 융합은 어느 정도로, 어떠한 형태로 나타나는가?'),
       ('RQ2', '시간적 변화', '융합의 수준과 형태는 시간에 따라 어떻게 변화하는가?'),
       ('RQ3', '구조적 중심 변화', '융합의 중심은 핵심 기술과 주변 기술 사이에서 어떻게 이동하는가?')]
rh, rg = 3.1, 0.3
ry = Y_A + A_H - 0.6 - 3 * rh - 2 * rg
for i, (tag, ttl, q) in enumerate(rqs):
    yy = ry + i * (rh + rg)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, XR + 0.9, yy, COLW - 1.8, rh, fill=WHITE, radius=0.1, shadow=True)
    shape_text(shape(MSO_SHAPE.ROUNDED_RECTANGLE, XR + 1.25, yy + 0.4, 3.4, rh - 0.8, fill=NAVY, radius=0.18),
               [P((tag, 28, YEL, True))], margin=0)
    text(XR + 5.1, yy + 0.15, COLW - 6.4, rh - 0.3, [P((ttl, 22, NAVY, True), after=1), P((q, 19, TEXT))], anchor='m')

# ---------------------------------------------------------------- 03 data & method
section(3, '자료와 방법', X0, Y_B, CW, B_H, None)
steps = [('FaDatabase', '자료 수집', ['EPO PATSTAT Global, CPC H01L21·H01L24', '+ 제목 키워드 → 관련 출원 928건 (1968–2024)']),
         ('FaCodeMerge', '공동분류 유형화', ['두 분류의 공동분류 = 융합의 대리측정치', '접합 단독형 · 통합형 · 공정 단독형']),
         ('FaCircleNodes', '핵심 기술 식별', ['기술 클러스터 · 기술 네트워크', '→ 핵심 접합 기호 24/03·05·08·80']),
         ('FaChartLine', '통계 분석', ['다항 로짓(RQ1·RQ2), 이항 로짓(RQ3)', 'Jaccard 계수 · 매개 중심성(RQ3)'])]
c_top, c_h = Y_B + 3.1, B_H - 3.1 - 0.6
stw, sth = 19.6, (c_h - 3 * 0.25) / 4
for i, (ic, ttl, ls) in enumerate(steps):
    yy = c_top + i * (sth + 0.25)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, X0 + 0.9, yy, stw, sth, fill=WHITE, radius=0.12, shadow=True)
    icon_circle(ic, X0 + 1.25, yy + (sth - 1.6) / 2, 1.6, NAVY2)
    text(X0 + 3.2, yy + 0.1, stw - 3.5, sth - 0.2, [P(('STEP %d  ' % (i + 1), 15, MUTED, True), (ttl, 20, NAVY, True), after=1)] +
         [P((l, 16, TEXT)) for l in ls], anchor='m')
vx = X0 + 0.9 + stw + 0.9
vw = 17.8
shape(MSO_SHAPE.ROUNDED_RECTANGLE, vx, c_top, vw, c_h, fill=WHITE, radius=0.04, shadow=True)
text(vx, c_top + 0.25, vw, 0.9, [P(('CPC 공동분류에 따른 출원 유형 (N = 928)', 19, NAVY, True), align=PP_ALIGN.CENTER)])
img_w = vw - 1.2
s.shapes.add_picture(f'{ASSET}/venn.png', cm(vx + 0.6), cm(c_top + 1.3), cm(img_w))
text(vx + 0.3, c_top + c_h - 1.15, vw - 0.6, 1.0, [P(('통합형', 16, INTEG, True), (' = 두 분류가 함께 부여된, 공정 단계의 경계를 넘는 발명', 16), align=PP_ALIGN.CENTER)])
nx = vx + vw + 0.9
nw_ = X1 - 0.9 - nx
shape(MSO_SHAPE.ROUNDED_RECTANGLE, nx, c_top, nw_, c_h, fill=WHITE, radius=0.04, shadow=True)
text(nx, c_top + 0.25, nw_, 0.9, [P(('기술 네트워크와 핵심 접합 기술', 19, NAVY, True), align=PP_ALIGN.CENTER)])
ih = c_h - 1.2 - 2.3
iw = ih * 1621 / 1196
if iw > nw_ - 0.6:
    iw = nw_ - 0.6; ih = iw * 1196 / 1621
s.shapes.add_picture(f'{ASSET}/network_nolegend.png', cm(nx + (nw_ - iw) / 2), cm(c_top + 1.15), cm(iw), cm(ih))
lg = [(CORE, '핵심 기술'), ('EFB26A', '기타 접합 기호'), ('1F3A5F', '제조공정 기호')]
lgx = nx + 0.7
for k, (col, lab) in enumerate(lg):
    lx = lgx + k * (nw_ - 1.4) / 3
    shape(MSO_SHAPE.OVAL, lx, c_top + c_h - 2.05, 0.45, 0.45, fill=col)
    text(lx + 0.6, c_top + c_h - 2.25, (nw_ - 1.4) / 3 - 0.6, 0.8, [P((lab, 14, TEXT))], anchor='m')
text(nx + 0.3, c_top + c_h - 1.3, nw_ - 0.6, 1.1, [P(('핵심 접합 기호 4개를 접합 출원의 81.3%가 보유', 15, TEXT), align=PP_ALIGN.CENTER)], anchor='m')

# ---------------------------------------------------------------- 04 results
section(4, '분석 결과', X0, Y_C, CW, C_H, None)
rw = (CW - 1.8 - 2 * 0.9) / 3
rcols = [X0 + 0.9 + i * (rw + 0.9) for i in range(3)]
heads = [('RQ1', '공정 단계 간 융합은 예외가 아니다'), ('RQ2', '형성 → 확산 → 재편의 세 국면'), ('RQ3', '융합의 중심: 핵심에서 주변으로')]
r_top, r_h = Y_C + 3.1, C_H - 3.1 - 0.6
tk_h = 3.3
tk_y = r_top + r_h - tk_h - 0.2
for x, (tag, hd) in zip(rcols, heads):
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, r_top, rw, r_h, fill=WHITE, radius=0.04, shadow=True)
    shape_text(shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.5, r_top + 0.45, 2.5, 1.4, fill=NAVY, radius=0.2), [P((tag, 22, YEL, True))], margin=0)
    text(x + 3.3, r_top + 0.4, rw - 3.7, 1.5, [P((hd, 23, NAVY, True))], anchor='m')
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.4, tk_y, rw - 0.8, tk_h, fill=INTEG_L, radius=0.08)

def takeaways(x, items):
    text(x + 0.7, tk_y + 0.15, rw - 1.3, tk_h - 0.3, [P(*runs, bullet=('●', INTEG), indent=0.65, after=3) for runs in items], anchor='m')

# RQ1
x = rcols[0]
text(x + 0.5, r_top + 2.0, 7.4, 3.2, [P(('45.8%', 58, INTEG, True), line=0.85), P(('통합형 = 최다 유형', 17, TEXT, True))])
text(x + 8.0, r_top + 2.2, rw - 8.4, 2.6, [P(('37.2%', 36, '9A72E8', True), line=0.9), P(('패키징 단계 공정을 제외한', 15, MUTED)), P(('좁은 정의에서도', 15, MUTED))])
pbj = json.load(open(f'{ASSET}/pred_breadth.json'))
bs = list(range(1, 16))
cd = CategoryChartData(); cd.categories = [str(b) for b in bs]
cd.add_series('통합형', [round(pbj[str(b)][1], 3) for b in bs])
cd.add_series('접합 단독형', [round(pbj[str(b)][0], 3) for b in bs])
cd.add_series('공정 단독형', [round(pbj[str(b)][2], 3) for b in bs])
gy = r_top + 5.35
text(x + 0.5, gy, rw - 1.0, 0.8, [P(('기술범위(CPC 서브그룹 수)별 예측확률', 16, NAVY, True))])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE, cm(x + 0.3), cm(gy + 0.75), cm(rw - 0.6), cm(tk_y - gy - 0.95), cd); ch = gf.chart
style_chart(ch, size=14)
for ser, col in zip(ch.plots[0].series, (INTEG, BOND, PROC)):
    ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(4.5); ser.smooth = False
    ser.marker.style = XL_MARKER_STYLE.NONE
ch.value_axis.maximum_scale = 1.0; ch.value_axis.minimum_scale = 0; ch.value_axis.major_unit = 0.2
ch.value_axis.tick_labels.number_format = '0.0'; ch.value_axis.tick_labels.number_format_is_linked = False
takeaways(x, [[('세부 기술이 다양한 출원일수록 통합형', 18, TEXT, True)],
              [('기술범위 1 증가 → 상대위험 1.40배', 17), (' (RRR 1.403***)', 15, MUTED)],
              [('기술범위 3 → 통합형 확률 0.26, 8 → 0.68', 17)]])

# RQ2
x = rcols[1]
cd = CategoryChartData(); cd.categories = ['형성기 –2014', '확산기 2015–19', '재편기 2020–24']
cd.add_series('공정 단독형', (43.6, 14.8, 10.4)); cd.add_series('통합형', (31.3, 62.2, 44.3)); cd.add_series('접합 단독형', (25.1, 23.0, 45.4))
text(x + 0.5, r_top + 2.05, rw - 1.0, 0.8, [P(('시기별 공동분류 유형 구성비 (%)', 16, NAVY, True))])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, cm(x + 0.3), cm(r_top + 2.8), cm(rw - 0.6), cm(tk_y - r_top - 3.0), cd); ch = gf.chart
style_chart(ch, size=15)
pl = ch.plots[0]; pl.gap_width = 40; pl.overlap = 100
for ser, col in zip(pl.series, (PROC, INTEG, BOND)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
    dl = ser.data_labels
    dl.show_value = True; dl.number_format = '0.0'; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.CENTER
    dl.font.size = Pt(17); dl.font.bold = True; dl.font.color.rgb = rgb(WHITE); dl.font.name = FONT
ch.value_axis.maximum_scale = 100; ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
takeaways(x, [[('통합형 31.3% → 62.2% → 44.3%', 18, TEXT, True)],
              [('공정 단독형 상대위험 매년 8.7% 감소', 17), (' (RRR 0.913***)', 15, MUTED)],
              [('출원의 무게중심: 제조공정 → 접합', 17, TEXT, True)]])

# RQ3: odds-ratio dot plot (shapes) + Jaccard slope chart
x = rcols[2]
text(x + 0.5, r_top + 2.05, rw - 1.0, 0.8, [P(('핵심 기술 보유 → 통합형 승산비 (95% CI)', 16, NAVY, True))])
ax0, ax1 = x + 3.6, x + rw - 0.8
lo, hi = math.log2(0.25), math.log2(32)
def xp(v): return ax0 + (math.log2(v) - lo) / (hi - lo) * (ax1 - ax0)
rows = [('형성기', 11.071, 4.211, 29.107, '11.07***', PROC), ('확산기', 2.452, 1.158, 5.194, '2.45**', INTEG), ('재편기', 0.658, 0.374, 1.160, '0.66 (비유의)', BOND)]
oy0, rs = r_top + 2.95, 0.95
line(xp(1), oy0, xp(1), oy0 + 3 * rs, '9AA3B2', w=1.5, dash='dash')
for i, (nm, orv, l, h, lab, col) in enumerate(rows):
    yy = oy0 + 0.5 + i * rs
    text(x + 0.5, yy - 0.45, 3.0, 0.9, [P((nm, 17, col, True))], anchor='m')
    line(xp(l), yy, xp(h), yy, col, w=3.5)
    shape(MSO_SHAPE.OVAL, xp(orv) - 0.3, yy - 0.3, 0.6, 0.6, fill=col)
    if i == 0:
        text(xp(orv) - 3.4, yy - 0.95, 3.0, 0.6, [P((lab, 14, col, True), align=PP_ALIGN.RIGHT)])
    else:
        text(xp(h) + 0.2, yy - 0.32, 3.4, 0.6, [P((lab, 14, col, True))], anchor='m')
axy = oy0 + 3 * rs + 0.1
line(ax0, axy, ax1, axy, '9AA3B2', w=1.25)
for v in (0.25, 0.5, 1, 2, 4, 8, 16, 32):
    line(xp(v), axy, xp(v), axy + 0.18, '9AA3B2', w=1.25)
    text(xp(v) - 0.8, axy + 0.18, 1.6, 0.6, [P((('%g' % v), 13, MUTED), align=PP_ALIGN.CENTER)])
jy = axy + 0.85
text(x + 0.5, jy, rw - 1.0, 0.8, [P(('제조공정 영역별 접합 기호와의 결합 강도 (Jaccard)', 16, NAVY, True))])
cd = CategoryChartData(); cd.categories = ['형성기', '확산기', '재편기']
cd.add_series('웨이퍼 제조공정', (0.306, 0.598, 0.159)); cd.add_series('공정 장치', (0.046, 0.059, 0.176)); cd.add_series('패키징 단계 공정', (0.194, 0.234, 0.268))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, cm(x + 0.3), cm(jy + 0.7), cm(rw - 0.6), cm(tk_y - jy - 0.85), cd); ch = gf.chart
style_chart(ch, size=13)
for ser, col in zip(ch.plots[0].series, (PROC, TEAL, GOLD)):
    ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(4); ser.smooth = False
    ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = 10
    ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col); ser.marker.format.line.color.rgb = rgb(col)
ch.value_axis.maximum_scale = 0.6; ch.value_axis.minimum_scale = 0; ch.value_axis.major_unit = 0.2
ch.value_axis.tick_labels.number_format = '0.0'; ch.value_axis.tick_labels.number_format_is_linked = False
takeaways(x, [[('재편기에는 핵심 기술의 통합형 우위가 확인되지 않음', 18, TEXT, True)],
              [('접합 출원 중 핵심 기술 53.5% → 88.0%', 17)],
              [('결합 중심: 웨이퍼 제조공정 → 공정 장치·패키징 단계 공정', 17, TEXT, True)]])

# ---------------------------------------------------------------- 05 conclusion
section(5, '결론 및 시사점', X0, Y_D, CW, D_H, None)
phw = 12.0
phases = [('형성기', '–2014', PROC, ['공정 단독형 우세 (43.6%)', '핵심 기술 출원은 제조공정 지식을 함께 포함 (OR 11.1)']),
          ('확산기', '2015–2019', INTEG, ['통합형 정점 (62.2%)', '웨이퍼 제조공정과의 결합이 가장 강함']),
          ('재편기', '2020–2024', BOND, ['접합 단독형 증가 (45.4%)', '핵심 설계 고착화 → 공정 장치·패키징 단계 공정으로 이동'])]
px0, py = X0 + 0.9, Y_D + 3.0
step_x = phw - 0.3
for i, (nm, yrs_, col, ls) in enumerate(phases):
    kind = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
    sh = shape(kind, px0 + i * step_x, py, phw, 1.8, fill=col)
    shape_text(sh, [P((nm + '  ', 22, WHITE, True), (yrs_, 17, WHITE))], margin=0.9)
    text(px0 + i * step_x + 0.5, py + 1.95, phw - 1.5, 2.4, [P((ls[0], 17, col, True), after=1), P((ls[1], 15, TEXT))])
cwid = 3 * step_x + 0.3
shape(MSO_SHAPE.ROUNDED_RECTANGLE, px0, py + 4.45, cwid, D_H - 3.0 - 4.45 - 0.6, fill=NAVY, radius=0.12)
text(px0 + 0.6, py + 4.45, cwid - 1.2, D_H - 3.0 - 4.45 - 0.6, [
    P(('공정 단계 간 기술융합은 단선적 확대가 아니라 융합의 내용과 중심이 국면에 따라 재편되는 ', 19, WHITE, True),
      ('동태적 과정', 19, YEL, True), ('이며, 그 배경에 핵심 기술 설계의 고착화가 있음을 시사한다.', 19, WHITE, True), line=1.0)], anchor='m')
ccx = px0 + cwid + 0.9
iw2 = X1 - 0.9 - ccx
imp_h = (D_H - 3.0 - 0.6 - 0.3) / 2
for k, (ic, ttl, body) in enumerate([
        ('FaGraduationCap', '학술적 시사점', '기술융합 연구의 범위를 공정 단계 간 융합으로 확장하고, 융합을 동태적 과정으로 실증'),
        ('FaIndustry', '실무적 시사점', 'HBM 등 하이브리드 본딩 경쟁력은 접합 원리 하나에 좌우되지 않을 수 있음 → 평탄화·표면 화학·정렬 계측과 얇은 다이를 다루고 적층하는 공정 역량까지 필요')]):
    yy = Y_D + 3.0 + k * (imp_h + 0.3)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, ccx, yy, iw2, imp_h, fill=WHITE, radius=0.08, shadow=True)
    icon_circle(ic, ccx + 0.35, yy + 0.3, 1.3, NAVY2)
    text(ccx + 1.95, yy + 0.25, iw2 - 2.2, 1.4, [P((ttl, 18, NAVY, True))], anchor='m')
    text(ccx + 0.4, yy + 1.65, iw2 - 0.75, imp_h - 1.75, [P((body, 14, TEXT), line=1.0)])

# ---------------------------------------------------------------- footer
refs = ['Anderson & Tushman (1990), Administrative Science Quarterly 35(4)',
        'Curran & Leker (2011), Technological Forecasting and Social Change 78(2)',
        'Hacklin, Marxt & Fahrni (2010), Int. J. Technology Management 49(1–3)',
        'Lau (2022), IEEE Trans. Components, Packaging and Manufacturing Tech. 12(2)',
        'Murmann & Frenken (2006), Research Policy 35(7)']
text(X0, 80.45, 21.0, 3.4, [P(('주요 참고문헌', 12, NAVY, True))] + [P((r, 10.5, MUTED)) for r in refs], anchor='t')
text(37.8, 80.45, X1 - 37.8, 3.4, [
    P(('Contact', 12, NAVY, True), align=PP_ALIGN.RIGHT),
    P(('이형규 · 한양대학교 기술경영학과', 12, TEXT), align=PP_ALIGN.RIGHT),
    P(('higgack@gmail.com', 12, TEXT), align=PP_ALIGN.RIGHT),
    P(('자료: EPO PATSTAT Global · 분석 단위: 특허출원', 10.5, MUTED), align=PP_ALIGN.RIGHT)], anchor='t')

prs.save(OUT); print('saved', OUT)

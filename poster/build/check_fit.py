# -*- coding: utf-8 -*-
"""Check every text frame of a poster PPTX for overflow as PowerPoint would lay it out with Malgun Gothic.

PowerPoint (맑은 고딕) sets Korean text about 6-13 % wider than the Noto Sans CJK used for the PDF preview
(calibrated on the user's screenshots of v5, 10-12) and wraps Korean text only at spaces (a word such as
"절단·조립·패키징" is never split). This script estimates line breaks with Noto metrics x WSCALE, breaking at spaces
only, and compares the text height (exact line pitch + paragraph spacing) with the frame height.

Usage: python3 check_fit.py deck.pptx [WSCALE=1.13]
Prints one line per frame that overflows or has a word wider than the frame; exit 1 if any.
"""
import sys
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn
from PIL import ImageFont

FONT = {False: ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc', 1000, index=1),
        True: ImageFont.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc', 1000, index=1)}
PT_CM = 0.03528


def runs_of(p):
    out = []
    for r in p.runs:
        sz = r.font.size.pt if r.font.size else 18.0
        out.append((r.text, sz, bool(r.font.bold)))
    return out


def para_lines(p, avail_cm, scale):
    """number of lines PowerPoint needs for paragraph p (wrap at spaces only)"""
    toks = []          # (width_cm, is_space)
    for t, sz, b in runs_of(p):
        for part in t.replace('\v', '\n').split('\n'):
            pass
        cur = ''
        for ch in t:
            if ch == ' ':
                if cur: toks.append((FONT[b].getlength(cur) / 1000 * sz * PT_CM * scale, False)); cur = ''
                toks.append((FONT[b].getlength(' ') / 1000 * sz * PT_CM * scale, True))
            else:
                cur += ch
        if cur: toks.append((FONT[b].getlength(cur) / 1000 * sz * PT_CM * scale, False))
    # merge consecutive non-space tokens (a word may span runs)
    words, cur, spaces = [], 0.0, []
    for w, sp in toks:
        if sp:
            if cur: words.append(cur); cur = 0.0
            words.append(-w)            # negative = space width
        else:
            cur += w
    if cur: words.append(cur)
    lines, x, widest = 1, 0.0, 0.0
    pending_space = 0.0
    for w in words:
        if w < 0:
            pending_space += -w; continue
        widest = max(widest, w)
        if x > 0 and x + pending_space + w > avail_cm:
            lines += 1; x = w
        else:
            x += (pending_space if x > 0 else 0) + w
        pending_space = 0.0
    return lines, widest


def pitch_cm(p):
    pPr = p._p.find(qn('a:pPr'))
    sizes = [sz for _, sz, _ in runs_of(p)] or [18.0]
    lnspc = None
    if pPr is not None:
        ls = pPr.find(qn('a:lnSpc'))
        if ls is not None:
            pts = ls.find(qn('a:spcPts'))
            pct = ls.find(qn('a:spcPct'))
            if pts is not None: lnspc = int(pts.get('val')) / 100 * PT_CM
            elif pct is not None: lnspc = max(sizes) * 1.2 * int(pct.get('val')) / 100000 * PT_CM
    return lnspc if lnspc else max(sizes) * 1.2 * PT_CM


def spacing_cm(p):
    pPr = p._p.find(qn('a:pPr')); total = 0.0
    if pPr is not None:
        for tag in ('a:spcBef', 'a:spcAft'):
            el = pPr.find(qn(tag))
            if el is not None and el.find(qn('a:spcPts')) is not None:
                total += int(el.find(qn('a:spcPts')).get('val')) / 100 * PT_CM
    return total


def indent_cm(p):
    pPr = p._p.find(qn('a:pPr'))
    if pPr is not None and pPr.get('marL'):
        return Emu(int(pPr.get('marL'))).cm
    return 0.0


def check_frame(tf, w_cm, h_cm, scale, label):
    bp = tf._txBody.find(qn('a:bodyPr'))
    ins = {k: Emu(int(bp.get(k))).cm if bp is not None and bp.get(k) else d for k, d in (('lIns', 0.254), ('rIns', 0.254), ('tIns', 0.127), ('bIns', 0.127))}
    avail = w_cm - ins['lIns'] - ins['rIns']
    need, widest_over, nl = 0.0, [], 0
    paras = [p for p in tf.paragraphs]
    for i, p in enumerate(paras):
        if not ''.join(r.text for r in p.runs).strip():
            need += pitch_cm(p) if i < len(paras) - 1 else 0; continue
        a = avail - indent_cm(p)
        lines, widest = para_lines(p, a, scale)
        nl += lines
        need += lines * pitch_cm(p) + (spacing_cm(p) if i < len(paras) - 1 else 0)
        if widest > a + 0.01:
            widest_over.append((''.join(r.text for r in p.runs)[:30], round(widest, 2), round(a, 2)))
    room = h_cm - ins['tIns'] - ins['bIns']
    return need, room, nl, widest_over


def main(path, scale=1.13):
    prs = Presentation(path); s = prs.slides[0]; bad = 0
    for sh in s.shapes:
        frames = []
        if sh.has_text_frame and sh.text_frame.text.strip():
            frames.append((sh.text_frame, Emu(sh.width).cm, Emu(sh.height).cm, sh.text_frame.text.strip().replace('\n', ' / ')[:40]))
        if getattr(sh, 'has_table', False) and sh.has_table:
            tbl = sh.table
            for r_i, row in enumerate(tbl.rows):
                for c_i, cell in enumerate(row.cells):
                    if cell.text.strip():
                        frames.append((cell.text_frame, Emu(tbl.columns[c_i].width).cm, Emu(row.height).cm, 'cell: ' + cell.text.strip()[:30]))
        for tf, w, h, lab in frames:
            need, room, nl, wide = check_frame(tf, w, h, scale, lab)
            top = Emu(sh.top).cm
            if need > room + 0.02 or wide:
                bad += 1
                print(f'OVER y={top:5.1f} w={w:5.1f} h={h:4.2f} need={need:4.2f} lines={nl}  {lab}' + (f'  WIDE {wide}' if wide else ''))
    print('frames over:', bad, '(scale', scale, ')')
    return bad


if __name__ == '__main__':
    sc = float(sys.argv[2]) if len(sys.argv) > 2 else 1.13
    sys.exit(1 if main(sys.argv[1], sc) else 0)

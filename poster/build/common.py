# -*- coding: utf-8 -*-
"""Shared helpers and thesis content for the Research Day poster variants (B, C, D).
All numbers come from the thesis final version v59."""
import json, math
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.oxml.ns import qn
from lxml import etree

FONT = '맑은 고딕'
NAVY, NAVY2 = '13294B', '24406E'
PROC, INTEG, BOND, CORE = '2563C9', '7C3AED', 'EA7A28', 'B5441B'
PROC_L, INTEG_L, BOND_L = 'E3ECFB', 'EEE6FD', 'FDEBDD'
YEL, TEXT, MUTED, PANEL, WHITE = 'FFC000', '1E2532', '5B6475', 'F2F5FA', 'FFFFFF'
TEAL, GOLD = '10A39A', 'E0A100'
W, H = 59.41, 84.10
L, C, R = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT

prs = s = P_BOX = None
ASSET = ICON = None

def cm(v): return Emu(int(round(v * 360000)))
def rgb(h): return RGBColor.from_string(h)

def init(tpl, asset, icon, keep=(150, 151, 153)):
    global prs, s, P_BOX, ASSET, ICON
    ASSET, ICON = asset, icon
    prs = Presentation(tpl); s = prs.slides[0]
    for sh in list(s.shapes):
        if sh.is_placeholder and sh.placeholder_format.idx not in keep:
            sh._element.getparent().remove(sh._element)
    P_BOX = [sh for sh in s.shapes if sh.name == '직사각형 1'][0]
    return s

def to_back(sh):
    s.shapes._spTree.remove(sh._element); s.shapes._spTree.insert(2, sh._element)

def to_front(sh):
    s.shapes._spTree.remove(sh._element); s.shapes._spTree.append(sh._element)

def _nostyle(sh):
    st = sh._element.find(qn('p:style'))
    if st is not None: sh._element.remove(st)

def font_run(r, size, color=TEXT, bold=False, italic=False):
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = rgb(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', FONT)
    rPr.set('lang', 'ko-KR')

def set_para(p, align=L, space_after=0, line=None, bullet=None, indent=0.0, space_before=0):
    p.alignment = align
    p.space_after = Pt(space_after); p.space_before = Pt(space_before)
    if line: p.line_spacing = line if not isinstance(line, (int, float)) or line > 5 else line
    pPr = p._p.get_or_add_pPr()
    for t in ('a:buNone', 'a:buChar', 'a:buFont', 'a:buClr', 'a:buSzPct'):
        for el in pPr.findall(qn(t)): pPr.remove(el)
    if bullet:
        ch, col = bullet[:2]
        pPr.set('marL', str(int(cm(indent)))); pPr.set('indent', str(-int(cm(indent))))
        bc = etree.SubElement(pPr, qn('a:buClr')); sc = etree.SubElement(bc, qn('a:srgbClr')); sc.set('val', col)
        if len(bullet) > 2:
            bs = etree.SubElement(pPr, qn('a:buSzPct')); bs.set('val', str(int(bullet[2] * 1000)))
        bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', 'Arial')
        b = etree.SubElement(pPr, qn('a:buChar')); b.set('char', ch)
    else:
        pPr.set('marL', '0'); pPr.set('indent', '0')
        etree.SubElement(pPr, qn('a:buNone'))

def norm(rr):
    rr = tuple(rr)
    return (rr[0], rr[1] if len(rr) > 1 else 24, rr[2] if len(rr) > 2 else TEXT, rr[3] if len(rr) > 3 else False,
            rr[4] if len(rr) > 4 else False)

def P(*runs, **kw):
    d = {'runs': [r if isinstance(r, tuple) else (r, kw.get('size', 24), kw.get('color', TEXT), kw.get('bold', False)) for r in runs]}
    d.update({k: v for k, v in kw.items() if k not in ('size', 'color', 'bold')}); return d

LINE_K = 1.2   # exact line pitch = LINE_K x largest font size (same in PowerPoint and LibreOffice)

def _pitch(pd):
    mx = max(norm(r)[1] for r in pd['runs'])
    return Pt(round(mx * LINE_K * (pd.get('line') or 1.0), 1))

def _fill_tf(tf, paras, default_align=L):
    for i, pd in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_para(p, pd.get('align', default_align), pd.get('after', 0), _pitch(pd), pd.get('bullet'), pd.get('indent', 0), pd.get('before', 0))
        for (t, size, col, bold, ital) in map(norm, pd['runs']):
            r = p.add_run(); r.text = t; font_run(r, size, col, bold, ital)

def _anchor(a): return {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[a]

def text(x, y, w, h, paras, anchor='t', margin=0.0, name=None):
    tb = s.shapes.add_textbox(cm(x), cm(y), cm(w), cm(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = cm(margin)
    tf.vertical_anchor = _anchor(anchor)
    bp = tf._txBody.find(qn('a:bodyPr'))
    for el in list(bp): bp.remove(el)
    etree.SubElement(bp, qn('a:noAutofit'))
    _fill_tf(tf, paras)
    return tb

def shape(kind, x, y, w, h, fill=None, line=None, lw=0.0, radius=None, shadow=False, alpha=None, name=None, dash=None):
    sh = s.shapes.add_shape(kind, cm(x), cm(y), cm(w), cm(h)); _nostyle(sh)
    if name: sh.name = name
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = sh.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            a = etree.SubElement(clr, qn('a:alpha')); a.set('val', str(int(alpha * 100000)))
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
        if dash:
            ln = sh.line._get_or_add_ln(); d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', dash)
    else:
        sh.line.fill.background()
    if radius is not None and kind in (MSO_SHAPE.ROUNDED_RECTANGLE,):
        sh.adjustments[0] = radius
    if shadow:
        spPr = sh._element.spPr
        eff = etree.SubElement(spPr, qn('a:effectLst'))
        sd = etree.SubElement(eff, qn('a:outerShdw'))
        for k, v in (('blurRad', '190500'), ('dist', '50800'), ('dir', '5400000'), ('algn', 't'), ('rotWithShape', '0')): sd.set(k, v)
        c = etree.SubElement(sd, qn('a:srgbClr')); c.set('val', '1E2532')
        a = etree.SubElement(c, qn('a:alpha')); a.set('val', '18000')
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = cm(0.15)
    return sh

def shape_text(sh, paras, anchor='m', margin=0.2):
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = cm(margin); tf.margin_top = tf.margin_bottom = cm(0.1)
    tf.vertical_anchor = _anchor(anchor)
    _fill_tf(tf, paras, default_align=C)
    return sh

def line(x1, y1, x2, y2, color, w=2.0, dash=None, head=None, tail=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cm(x1), cm(y1), cm(x2), cm(y2)); _nostyle(c)
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', dash)
    for tag, kind in (('a:headEnd', head), ('a:tailEnd', tail)):
        if kind:
            e = etree.SubElement(ln, qn(tag)); e.set('type', kind); e.set('w', 'med'); e.set('len', 'med')
    return c

def picture(path, x, y, w=None, h=None):
    return s.shapes.add_picture(path, cm(x), cm(y), cm(w) if w else None, cm(h) if h else None)

def icon(name, x, y, size, color='FFFFFF'):
    return picture(f'{ICON}/{name}_{color}.png', x, y, size, size)

def icon_circle(name, x, y, d, fill, pad=0.26, color='FFFFFF'):
    shape(MSO_SHAPE.OVAL, x, y, d, d, fill=fill)
    icon(name, x + d * pad, y + d * pad, d * (1 - 2 * pad), color)

def fill_ph(idx, x, y, w, h, paras, anchor='t'):
    ph = [sh for sh in s.placeholders if sh.placeholder_format.idx == idx][0]
    ph.left, ph.top, ph.width, ph.height = cm(x), cm(y), cm(w), cm(h)
    tf = ph.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _anchor(anchor)
    bp = tf._txBody.find(qn('a:bodyPr'))
    for el in list(bp): bp.remove(el)
    etree.SubElement(bp, qn('a:noAutofit'))
    _fill_tf(tf, paras)
    return ph

# ---------------------------------------------------------------- tables
NO_GRID = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'

def _cell_border(cell, side, color, w):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {'L': 'a:lnL', 'R': 'a:lnR', 'T': 'a:lnT', 'B': 'a:lnB'}[side]
    for el in tcPr.findall(qn(tag)): tcPr.remove(el)
    ln = etree.Element(qn(tag)); ln.set('w', str(int(Pt(w)))); ln.set('cap', 'flat'); ln.set('cmpd', 'sng'); ln.set('algn', 'ctr')
    if color:
        sf = etree.SubElement(ln, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', color)
    else:
        etree.SubElement(ln, qn('a:noFill'))
    etree.SubElement(ln, qn('a:prstDash')).set('val', 'solid')
    # borders must precede the cell fill in tcPr
    tcPr.insert(0, ln)

def table(x, y, colw, rows, rowh, head_fill=NAVY, head_color=WHITE, size=16, head_size=None, body_fill=WHITE,
          alt_fill='F4F7FB', rule='C9D2E0', bold_cols=(), aligns=None, cell_fill=None, cell_color=None, cell_bold=None):
    """rows: list of lists of strings (first row = header)."""
    nR, nC = len(rows), len(rows[0])
    gf = s.shapes.add_table(nR, nC, cm(x), cm(y), cm(sum(colw)), cm(sum(rowh) if isinstance(rowh, list) else rowh * nR))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None: sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
    sid.text = NO_GRID
    tbl.first_row = True; tbl.horz_banding = False
    for j, w_ in enumerate(colw): tbl.columns[j].width = cm(w_)
    for i in range(nR): tbl.rows[i].height = cm(rowh[i] if isinstance(rowh, list) else rowh)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            fill = head_fill if i == 0 else (alt_fill if (i % 2 == 0 and alt_fill) else body_fill)
            if cell_fill and (i, j) in cell_fill: fill = cell_fill[(i, j)]
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fill)
            cell.margin_left = cell.margin_right = cm(0.15); cell.margin_top = cell.margin_bottom = cm(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            col = head_color if i == 0 else TEXT
            if cell_color and (i, j) in cell_color: col = cell_color[(i, j)]
            bold = (i == 0) or (j in bold_cols) or bool(cell_bold and (i, j) in cell_bold)
            al = C if i == 0 else (aligns[j] if aligns else C)
            lines_ = str(val).split('\n')
            _fill_tf(tf, [P((ln_, (head_size or size) if i == 0 else size, col, bold), align=al) for ln_ in lines_])
            for side in 'LR': _cell_border(cell, side, None, 0)
            _cell_border(cell, 'T', None if i > 0 else None, 0)
            _cell_border(cell, 'B', rule if i < nR - 1 else head_fill, 1.0 if i < nR - 1 else 2.0)
    return gf

# ---------------------------------------------------------------- charts
def style_chart(ch, legend=True, size=16, dark=False):
    fg = 'D6DEEB' if dark else TEXT
    axis = '5A6B88' if dark else '9AA3B2'
    grid = '263652' if dark else 'DDE2EA'
    ch.font.size = Pt(size); ch.font.name = FONT; ch.font.color.rgb = rgb(fg)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size); ch.legend.font.name = FONT; ch.legend.font.color.rgb = rgb(fg)
    for ax in (ch.category_axis, ch.value_axis):
        ax.tick_labels.font.size = Pt(size); ax.tick_labels.font.name = FONT; ax.tick_labels.font.color.rgb = rgb(fg)
        ax.format.line.color.rgb = rgb(axis)
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = rgb(grid)
    ch.value_axis.format.line.fill.background()
    for d in ch._chartSpace.iter(qn('a:defRPr')):
        if d.find(qn('a:ea')) is None:
            ea = etree.SubElement(d, qn('a:ea')); ea.set('typeface', FONT)

def chart_pred(x, y, w, h, size=14, dark=False, legend=True):
    pbj = json.load(open(f'{ASSET}/pred_breadth.json'))
    bs = list(range(1, 16))
    cd = CategoryChartData(); cd.categories = [str(b) for b in bs]
    cd.add_series('통합형', [round(pbj[str(b)][1], 3) for b in bs])
    cd.add_series('접합 단독형', [round(pbj[str(b)][0], 3) for b in bs])
    cd.add_series('공정 단독형', [round(pbj[str(b)][2], 3) for b in bs])
    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE, cm(x), cm(y), cm(w), cm(h), cd).chart
    style_chart(ch, legend=legend, size=size, dark=dark)
    cols = ('A78BFA', 'FB9A4B', '5B9BFF') if dark else (INTEG, BOND, PROC)
    for ser, col in zip(ch.plots[0].series, cols):
        ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(4.5); ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.NONE
    ch.value_axis.maximum_scale = 1.0; ch.value_axis.minimum_scale = 0; ch.value_axis.major_unit = 0.2
    ch.value_axis.tick_labels.number_format = '0.0'; ch.value_axis.tick_labels.number_format_is_linked = False
    return ch

def chart_period(x, y, w, h, size=15, label_size=17, dark=False, cats=('형성기 –2014', '확산기 2015–19', '재편기 2020–24'), legend=True):
    cd = CategoryChartData(); cd.categories = list(cats)
    cd.add_series('공정 단독형', (43.6, 14.8, 10.4)); cd.add_series('통합형', (31.3, 62.2, 44.3)); cd.add_series('접합 단독형', (25.1, 23.0, 45.4))
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, cm(x), cm(y), cm(w), cm(h), cd).chart
    style_chart(ch, legend=legend, size=size, dark=dark)
    pl = ch.plots[0]; pl.gap_width = 40; pl.overlap = 100
    cols = ('3B7BF0', '8B5CF6', 'F08A3C') if dark else (PROC, INTEG, BOND)
    for ser, col in zip(pl.series, cols):
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
        dl = ser.data_labels
        dl.show_value = True; dl.number_format = '0.0'; dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.CENTER
        dl.font.size = Pt(label_size); dl.font.bold = True; dl.font.color.rgb = rgb(WHITE); dl.font.name = FONT
    ch.value_axis.maximum_scale = 100; ch.value_axis.visible = False; ch.value_axis.has_major_gridlines = False
    return ch

def chart_jaccard(x, y, w, h, size=13, dark=False, legend=True):
    cd = CategoryChartData(); cd.categories = ['형성기', '확산기', '재편기']
    cd.add_series('웨이퍼 제조공정', (0.306, 0.598, 0.159)); cd.add_series('공정 장치', (0.046, 0.059, 0.176)); cd.add_series('패키징 단계 공정', (0.194, 0.234, 0.268))
    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, cm(x), cm(y), cm(w), cm(h), cd).chart
    style_chart(ch, legend=legend, size=size, dark=dark)
    cols = ('5B9BFF', '2DD4BF', 'FACC15') if dark else (PROC, TEAL, GOLD)
    for ser, col in zip(ch.plots[0].series, cols):
        ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(4); ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.CIRCLE; ser.marker.size = 10
        ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col); ser.marker.format.line.color.rgb = rgb(col)
    ch.value_axis.maximum_scale = 0.6; ch.value_axis.minimum_scale = 0; ch.value_axis.major_unit = 0.2
    ch.value_axis.tick_labels.number_format = '0.0'; ch.value_axis.tick_labels.number_format_is_linked = False
    return ch

def or_plot(x, y, w, label_w=3.0, rs=0.95, dark=False, size=17, lab_size=14, ref_note=None):
    """Odds ratios of core technology by period on a log2 axis (shapes). Returns bottom y."""
    fg = 'D6DEEB' if dark else MUTED
    axis = '5A6B88' if dark else '9AA3B2'
    cols = ('5B9BFF', 'A78BFA', 'FB9A4B') if dark else (PROC, INTEG, BOND)
    ax0, ax1 = x + label_w + 0.6, x + w - 0.6
    lo, hi = math.log2(0.25), math.log2(32)
    def xp(v): return ax0 + (math.log2(v) - lo) / (hi - lo) * (ax1 - ax0)
    rows = [('형성기', 11.071, 4.211, 29.107, '11.07***'), ('확산기', 2.452, 1.158, 5.194, '2.45**'), ('재편기', 0.658, 0.374, 1.160, '0.66 (비유의)')]
    line(xp(1), y, xp(1), y + 3 * rs, axis, w=1.5, dash='dash')
    for i, ((nm, orv, l, h, lab), col) in enumerate(zip(rows, cols)):
        yy = y + 0.5 + i * rs
        text(x, yy - 0.45, label_w, 0.9, [P((nm, size, col, True))], anchor='m')
        line(xp(l), yy, xp(h), yy, col, w=3.5)
        shape(MSO_SHAPE.OVAL, xp(orv) - 0.3, yy - 0.3, 0.6, 0.6, fill=col)
        if i == 0:
            text(xp(orv) - 3.4, yy - 0.95, 3.0, 0.6, [P((lab, lab_size, col, True), align=R)])
        else:
            text(xp(h) + 0.2, yy - 0.32, 3.6, 0.6, [P((lab, lab_size, col, True))], anchor='m')
    axy = y + 3 * rs + 0.1
    line(ax0, axy, ax1, axy, axis, w=1.25)
    for v in (0.25, 0.5, 1, 2, 4, 8, 16, 32):
        line(xp(v), axy, xp(v), axy + 0.18, axis, w=1.25)
        text(xp(v) - 0.8, axy + 0.18, 1.6, 0.6, [P((('%g' % v), 13, fg), align=C)])
    if ref_note:
        text(xp(1) - 3.0, axy + 0.75, 6.0, 0.6, [P((ref_note, 12, fg), align=C)])
        return axy + 1.4
    return axy + 0.8

def die_stack(x, y, w, hybrid, top='8A97AD', bot='5E6C85', bump='A9AFBA', oxide='BFE3E6'):
    die_h, gap = 0.9, (0.1 if hybrid else 0.8)
    shape(MSO_SHAPE.RECTANGLE, x, y, w, die_h, fill=top)
    by = y + die_h + gap
    if hybrid:
        oy = y + die_h - 0.38
        shape(MSO_SHAPE.RECTANGLE, x, oy, w, 0.38, fill=oxide)
        shape(MSO_SHAPE.RECTANGLE, x, by, w, 0.38, fill=oxide)
        n = 6
        for k in range(n):
            px = x + 0.45 + k * (w - 0.9 - 0.7) / (n - 1)
            shape(MSO_SHAPE.RECTANGLE, px, oy, 0.7, 0.38, fill=BOND)
            shape(MSO_SHAPE.RECTANGLE, px, by, 0.7, 0.38, fill=BOND)
        shape(MSO_SHAPE.RECTANGLE, x, by - gap, w, gap, fill=INTEG)
        by += 0.38
    else:
        n = 5
        for k in range(n):
            px = x + 0.55 + k * (w - 1.1 - 0.85) / (n - 1)
            shape(MSO_SHAPE.OVAL, px, y + die_h - 0.02, 0.85, gap + 0.04, fill=bump, line='7D8594', lw=0.75)
    shape(MSO_SHAPE.RECTANGLE, x, by, w, 0.9, fill=bot)
    return by + 0.9

def bump_vs_hybrid(x, y, w, h, card_fill=WHITE, title_muted=MUTED, fg=TEXT, sub=MUTED, shadow=True, gap=1.2, size=(20, 16)):
    cw2 = (w - gap) / 2
    for k, hyb in enumerate((False, True)):
        cx = x + k * (cw2 + gap)
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, y, cw2, h, fill=card_fill, radius=0.06, shadow=shadow)
        text(cx + 0.3, y + 0.2, cw2 - 0.6, 0.9, [P((('하이브리드 본딩' if hyb else '기존: 솔더 범프 접합'), size[0], (INTEG if hyb else title_muted), True), align=C)])
        die_stack(cx + 0.9, y + 1.1, cw2 - 1.8, hyb)
        text(cx + 0.3, y + h - 1.45, cw2 - 0.6, 1.3, [
            P((('구리–구리 + 산화막–산화막 직접 접합' if hyb else '다이 사이를 솔더 범프로 연결'), size[1], fg, True), align=C),
            P((('범프 없이 10 µm 이하 간격' if hyb else '10 µm 이하 간격은 구현하기 어려움'), size[1], sub), align=C)])
    line(x + cw2 + 0.15, y + h / 2, x + cw2 + gap - 0.15, y + h / 2, INTEG, w=4, tail='triangle')

# ---------------------------------------------------------------- content (thesis v59)
TITLE1 = '반도체 하이브리드 본딩 특허를 통한'
TITLE2a, TITLE2b = '공정 단계 간 기술융합', '의 동태적 분석'
TITLE_EN = 'A Dynamic Analysis of Cross-Stage Technological Convergence through Semiconductor Hybrid Bonding Patents'
AUTHOR, AFFIL, EMAIL = '이형규', '한양대학교 기술경영학과', 'higgack@naver.com'
EVENT = '2026 한양MOT 리서치데이'

def background(sz, hl=True):
    b = lambda t, c: (t, sz, c, True) if hl else (t, sz)
    return [
        [('반도체 산업은 ', sz), b('전공정', PROC), ('(웨이퍼 위 소자 형성)과 ', sz), b('후공정', BOND), ('(절단·조립·패키징)의 분업 구조를 반세기 이상 유지해 왔다.', sz)],
        [('미세화의 한계로 다이를 분할·적층·연결하는 첨단 패키징이 새로운 성능 향상의 원천이 되었다.', sz)],
        [('하이브리드 본딩', sz, TEXT, True), ('은 솔더 범프 없이 구리 패드와 산화막을 직접 접합하여 10 µm 이하의 연결 간격을 구현한다.', sz)],
        [('그 성패를 좌우하는 평탄화(CMP)·표면 활성화·정밀 정렬은 ', sz), b('전공정의 제조공정 역량', PROC), ('에 해당한다.', sz)]]

GAP = ['기존 기술융합 연구는 주로 서로 다른 산업·기술 분야 사이의 융합(ICT–가전, 식품–제약, 나노–바이오)을 다루었다.',
       '한 산업 안에서 공정 단계를 가로지르는 융합은 충분히 다루어지지 않았고, 이를 특허 단위에서 검정한 연구도 부족하다.',
       '융합을 단계적 과정으로 보는 이론에 비해, 융합 내용의 시간적 변화를 분석한 실증 연구는 부족하다.']
RQS = [('RQ1', '수준과 유형', '제조공정과 접합 기술 간의 융합은 어느 정도로, 어떠한 형태로 나타나는가?'),
       ('RQ2', '시간적 변화', '융합의 수준과 형태는 시간에 따라 어떻게 변화하는가?'),
       ('RQ3', '구조적 중심 변화', '융합의 중심은 핵심 기술과 주변 기술 사이에서 어떻게 이동하는가?')]
STEPS = [('FaDatabase', '자료 수집', ['EPO PATSTAT Global, CPC H01L21·H01L24', '+ 제목 키워드 → 관련 출원 928건 (1968–2024)']),
         ('FaCodeMerge', '공동분류 유형화', ['두 분류의 공동분류 = 융합의 대리측정치', '접합 단독형 · 통합형 · 공정 단독형']),
         ('FaCircleNodes', '핵심 기술 식별', ['기술 클러스터 · 기술 네트워크', '→ 핵심 접합 기호 24/03·05·08·80']),
         ('FaChartLine', '통계 분석', ['다항 로짓(RQ1·RQ2), 이항 로짓(RQ3)', 'Jaccard 계수 · 매개 중심성(RQ3)'])]
PHASES = [('형성기', '–2014', PROC, '공정 단독형 우세 (43.6%)', '핵심 기술 출원은 제조공정 지식을 함께 포함 (OR\u00a011.1)'),
          ('확산기', '2015–2019', INTEG, '통합형 정점 (62.2%)', '웨이퍼 제조공정과의 결합이 가장 강함'),
          ('재편기', '2020–2024', BOND, '접합 단독형 증가 (45.4%)', '핵심 설계 고착화 → 공정 장치·패키징 단계 공정으로 이동')]
CONCL = ('공정 단계 간 기술융합은 단선적 확대가 아니라 융합의 내용과 중심이 국면에 따라 재편되는 ', '동태적 과정',
         '이며, 그 배경에 핵심 기술 설계의 고착화가 있음을 시사한다.')
IMPL_ACAD = '기술융합 연구의 범위를 공정 단계 간 융합으로 확장하고, 융합을 동태적 과정으로 실증'
IMPL_PRAC = 'HBM 등 하이브리드 본딩 경쟁력은 접합 원리 하나에 좌우되지 않을 수 있음 → 평탄화·표면 화학·정렬 계측과 얇은 다이를 다루고 적층하는 공정 역량까지 필요'
REFS = ['Anderson & Tushman (1990), Administrative Science Quarterly 35(4)',
        'Curran & Leker (2011), Technological Forecasting and Social Change 78(2)',
        'Hacklin, Marxt & Fahrni (2010), Int. J. Technology Management 49(1–3)',
        'Lau (2022), IEEE Trans. Components, Packaging and Manufacturing Tech. 12(2)',
        'Murmann & Frenken (2006), Research Policy 35(7)']

def footer(fg=TEXT, head=NAVY, muted=MUTED, y=80.45):
    text(1.6, y, 21.0, 3.4, [P(('주요 참고문헌', 12, head, True))] + [P((r, 10.5, muted)) for r in REFS], anchor='t')
    text(37.8, y, 57.8 - 37.8, 3.4, [
        P(('Contact', 12, head, True), align=R),
        P((f'{AUTHOR} · {AFFIL}', 12, fg), align=R),
        P((EMAIL, 12, fg), align=R),
        P(('자료: EPO PATSTAT Global · 분석 단위: 특허출원', 10.5, muted), align=R)], anchor='t')

def save(out):
    prs.save(out); print('saved', out)

# ---------------------------------------------------------------- rough text metrics (for stacking blocks)
def _cw(ch):
    o = ord(ch)
    if 0xAC00 <= o <= 0xD7A3 or 0x3000 <= o <= 0x9FFF: return 1.0
    if ch == ' ': return 0.3
    if ch.isalnum(): return 0.58
    return 0.45

from PIL import ImageFont as _IF
_NOTO = _IF.truetype('/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc', 1000, index=1)

def est_lines(txt, size, width, indent=0.0):
    em = size * 0.03528
    avail = width - indent
    lines, cur = 1, 0.0
    for word in txt.split(' '):
        wl = _NOTO.getlength(word) / 1000 * em
        add = wl if cur == 0 else wl + 0.224 * em
        if cur + add > avail and cur > 0:
            lines += 1; cur = wl
        else:
            cur += add
    return lines

def est_height(items, size, width, indent=0.0, after_pt=0, line=1.0):
    """items: list of plain strings (one per paragraph)."""
    lh = size * 0.03528 * LINE_K * line
    n = sum(est_lines(t, size, width, indent) for t in items)
    return n * lh + len(items) * after_pt * 0.03528

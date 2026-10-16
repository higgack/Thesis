"""Capstone project report (한양대학교 기술경영전문대학원 석사학위 캡스톤 프로젝트 보고서) in Word, after the school's HWP form.

Usage: python3 capstone/tools/make_capstone.py 1            (add --no-render to skip the LibreOffice page-number passes)
  in : capstone/sources/capstone_v1.md (the report text; {+ +} added or changed, {- -} removed against the user's PDF),
       capstone/figures/*.png (the figures and table images taken from the user's PDF at their own resolution)
  out: capstone/capstone_report_v1.docx         the report to submit (marks resolved, nothing highlighted)
       capstone/capstone_report_v1_marked.docx  the same pages with yellow = added or changed text and fields the
                                                student fills in, red strike = text removed from the PDF

The layout copies the two HWP files the user received (parsed record by record, see capstone/README_ko.md):
  order      개요표 (19-row table + 결과물 사진 box) - 신규고용 창출 - 표지 - 인준지 - 목차 - 표 목차 - 그림 목차 - 제1-4장
             - 참고문헌
  page       A4; left/right 15 mm; body from 27 mm (위쪽 2 + 머리말 25) to 30 mm above the bottom (아래쪽 25 + 꼬리말 5)
  header     "캡스톤 프로젝트 보고서" left, page number right, 9 pt, from 목차 on; front pages lower-roman counted from
             표지 (the form's 목차 lists 표 목차 iv, 그림 목차 vi), body arabic from 1 at 제1장 (the form's 새 번호)
  fonts      한양신명조 -> HY신명조, 한양중고딕 -> HY중고딕 (the Word equivalents of the HWP fonts), header 맑은 고딕
  body       11 pt, line pitch 200 % (22 pt), justified, first line indented one character
  headings   제 N 장 16 pt bold, new page; N.M. 12 pt bold indented 10 pt; N.M.K. 11 pt bold; (1) 11 pt bold (not listed)
  captions   [표 N-k] above the table, 11 pt, centred; [그림 N-k] below the figure, 10 pt, centred; numbered by chapter
  references 10 pt, line 160 %, blank line between entries, Korean (가나다) then English (alphabetical)
  표지        rule above the 25 pt bold Korean and English titles; 한양대학교 기술경영전문대학원 / 석사학위 캡스톤 프로젝트
             보고서 / date, 13 pt; signature block (student, 지도교수) on the right
  인준지       the separate HWP 인준지: titles, "이 보고서를 …의 석사학위 보고서로 인준함.", date, 지도교수 / 학과장 / 원장
             (인), 한양대학교 기술경영전문대학원
The three lists are Word TOC fields (Word offers to update them when the file is opened); their cached entries carry the
page numbers of a LibreOffice rendering (repeated until stable) so they read right before updating.
"""
import os, pathlib, re, shutil, subprocess, sys, tempfile
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, Emu, RGBColor

VER = sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith('-') else '1'
RENDER = '--no-render' not in sys.argv
ROOT = pathlib.Path(__file__).resolve().parents[1]
SRC = ROOT / 'sources' / f'capstone_v{VER}.md'
FIG = ROOT / 'figures'
OUT = ROOT / f'capstone_report_v{VER}.docx'
OUT_MARKED = ROOT / f'capstone_report_v{VER}_marked.docx'

SERIF, GOTHIC, HEADF = 'HY신명조', 'HY중고딕', '맑은 고딕'
TEXT_W = 180.0                       # mm, 210 - 2 x 15
TEXT_W_TW = int(TEXT_W / 25.4 * 1440)
BODY_PT, LINE_PT = 11, 22.0           # 200 % of 11 pt, as measured in the form (7.75 mm per line)

# ---- report data. Segments: (text, mark) with mark None, 'add' (changed against the PDF) or 'fill' (to confirm)
TITLE_KO = [('특허 데이터', None), ('를', 'add'), (' 활용한 차세대 반도체 패키징 기술 생태계의 진화 메커니즘 연구: 하이브리드 ', None),
            ('본딩 기술', 'add'), ('을 중심으로', None)]
TITLE_EN = [('A Study on the Evolutionary Mechanism of the Next-Generation Semiconductor Packaging Technology '
             'Ecosystem Using Patent Data: Focusing on Hybrid Bonding Technology', None)]
NAME, NAME_SP = '이형규', '이 형 규'
ADVISOR, ADVISOR_SP = '강민구', '강 민 구'
STUDENT_ID = '2025155651'            # from the PDF's file name; the student confirms it
DATE = '2027 년  2 월'                # the degree date used on the thesis; the student confirms it
SCHOOL = '한양대학교 기술경영전문대학원'
TOPIC = '특허 데이터 기반 하이브리드 본딩 기술 생태계 분석'
CONTENT = [                           # one line each, as the form's four lines
    ('배경 및 목표', '하이브리드 본딩 중심의 차세대 패키징 기술 생태계 진화 점검'),
    ('수행내용 및 방법', 'PATSTAT 특허 기반 공정–구조–제품 프레임워크 분석'),
    ('주요 결과물', '출원 성장성, IP 기업의 원천 특허 장벽, 기업별 기술 전략 차이'),
    ('활용방안 및 기대 효과', '턴키 전략, 소부장·학계 협력 등 기술 전략 수립에 활용'),
]
ROLE = '연구 설계, 특허 데이터 구축·분석, 보고서 작성'


def fail(msg):
    print('FAIL:', msg); sys.exit(1)


# =============================================================== 1. source
SEG_RE = re.compile(r'\{\+(.*?)\+\}|\{-(.*?)-\}', re.S)


def segments(text):
    out, pos = [], 0
    for m in SEG_RE.finditer(text):
        if m.start() > pos: out.append((text[pos:m.start()], None))
        out.append((m.group(1), 'add') if m.group(1) is not None else (m.group(2), 'del'))
        pos = m.end()
    if pos < len(text): out.append((text[pos:], None))
    return out


def plain(segs, marked=False):
    return ''.join(t for t, k in segs if k != 'del')


def parse():
    t = SRC.read_text(encoding='utf-8')
    t = re.sub(r'<!--.*?-->', '', t, flags=re.S)
    if t.count('{+') != t.count('+}') or t.count('{-') != t.count('-}'): fail('unbalanced marks')
    blocks, chap = [], 0
    counters = {}
    for raw in re.split(r'\n\s*\n', t.strip()):
        s = raw.strip()
        if not s: continue
        if '\n' in s: fail('line break inside a paragraph: ' + s[:60])
        m = re.match(r'^(#{1,4}) (.+)$', s)
        if m:
            lv = len(m.group(1)); txt = m.group(2)
            if lv == 1:
                mm = re.match(r'^제 (\d) 장 ', txt)
                if mm: chap = int(mm.group(1))
            blocks.append(('h%d' % lv, segments(txt))); continue
        m = re.match(r'^!(fig|tab)\[([^|\]]+)\|(\d+)\] (.+)$', s)
        if m:
            kind, fn, w, cap = m.groups()
            if not (FIG / fn).exists(): fail('missing figure ' + fn)
            key = (kind, chap); counters[key] = counters.get(key, 0) + 1
            label = '[%s %d-%d]' % ('그림' if kind == 'fig' else '표', chap, counters[key])
            blocks.append((kind, fn, int(w), label, segments(cap))); continue
        if s.startswith('- '):
            blocks.append(('bullet', segments(s))); continue
        blocks.append(('p', segments(s)))
    # audits: every cross-reference names an existing table or figure; style words of the user's rules
    labels = {b[3] for b in blocks if b[0] in ('fig', 'tab')}
    body = '\n'.join(plain(b[-1]) if b[0] in ('p', 'bullet') else '' for b in blocks)
    for ref in re.findall(r'\[(?:그림|표) \d-\d+\]', body):
        if ref not in labels: fail('reference to a missing ' + ref)
    for word in ('동시에', '예컨대', '요컨대'):
        if word in body: fail('style word left: ' + word)
    return blocks


# =============================================================== 2. XML helpers
def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_fonts(rPr, latin=SERIF, ea=SERIF):
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = el('w:rFonts'); rPr.insert(0, rf)
    for k in list(rf.attrib): del rf.attrib[k]
    rf.set(qn('w:ascii'), latin); rf.set(qn('w:hAnsi'), latin); rf.set(qn('w:cs'), latin); rf.set(qn('w:eastAsia'), ea)


class Builder:
    def __init__(self, marked):
        self.marked = marked
        self.doc = Document()
        self.heads, self.caps = [], []

    # ---------------------------------------------------------------- runs and paragraphs
    def run(self, p, text, size=BODY_PT, bold=False, font=SERIF, mark=None, scale=None, spacing=None, color=None,
            italic=False):
        if mark == 'del' and not self.marked:
            return None
        r = p.add_run(text)
        rPr = r._r.get_or_add_rPr()
        set_fonts(rPr, font, font)
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        if color: r.font.color.rgb = RGBColor.from_string(color)
        if spacing is not None:
            rPr.append(el('w:spacing', **{'w:val': int(round(spacing * 20))}))
        if scale is not None:
            rPr.append(el('w:w', **{'w:val': scale}))
        if self.marked and mark in ('add', 'fill'):
            rPr.append(el('w:highlight', **{'w:val': 'yellow'}))
        if self.marked and mark == 'del':
            rPr.append(el('w:strike')); r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        return r

    def runs(self, p, segs, **kw):
        for t, k in segs:
            if t: self.run(p, t, mark=k, **kw)

    def pfmt(self, p, align=None, line=None, rule='exact', before=0, after=0, first_chars=None, left=None, hanging=None,
             right=None, keep=False, page_break=False, outline=None):
        pf = p.paragraph_format
        if align is not None: pf.alignment = align
        pf.space_before = Pt(before); pf.space_after = Pt(after)
        pPr = p._p.get_or_add_pPr()
        if line:
            pf.line_spacing = Pt(line)
            pPr.find(qn('w:spacing')).set(qn('w:lineRule'), rule)
        else:
            pf.line_spacing = 1.0
        if keep: pf.keep_with_next = True
        if page_break: pf.page_break_before = True
        ind = el('w:ind', **{'w:left': int((left or 0) * 20), 'w:right': int((right or 0) * 20)})
        if hanging:
            ind.set(qn('w:hanging'), str(int(hanging * 20)))
        else:
            fc = first_chars or 0
            ind.set(qn('w:firstLineChars'), str(int(fc * 100))); ind.set(qn('w:firstLine'), str(int(fc * BODY_PT * 20)))
        old = pPr.find(qn('w:ind'))
        if old is not None: pPr.remove(old)
        pPr.append(ind)
        if outline is not None:
            pPr.append(el('w:outlineLvl', **{'w:val': outline}))
        return p

    def para(self, segs=None, **kw):
        p = self.doc.add_paragraph()
        fmt = {k: kw.pop(k) for k in list(kw) if k in ('align', 'line', 'rule', 'before', 'after', 'first_chars', 'left',
                                                       'hanging', 'right', 'keep', 'page_break', 'outline')}
        self.pfmt(p, **fmt)
        if segs: self.runs(p, segs, **kw)
        return p

    def field(self, p, instr, cached=None, size=9, font=HEADF):
        def r(child):
            rr = el('w:r'); rr.append(child); p._p.append(rr); return rr
        r(el('w:fldChar', **{'w:fldCharType': 'begin'}))
        it = el('w:instrText'); it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); it.text = f' {instr} '
        r(it)
        if cached is not None:
            r(el('w:fldChar', **{'w:fldCharType': 'separate'}))
            self.run(p, cached, size=size, font=font)
        r(el('w:fldChar', **{'w:fldCharType': 'end'}))

    # ---------------------------------------------------------------- styles
    def styles(self):
        doc = self.doc
        rpr = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
        set_fonts(rpr)
        lang = rpr.find(qn('w:lang'))
        if lang is None:
            lang = el('w:lang'); rpr.append(lang)
        lang.set(qn('w:val'), 'en-US'); lang.set(qn('w:eastAsia'), 'ko-KR')
        ppr = doc.styles.element.find(qn('w:docDefaults')).find(qn('w:pPrDefault'))
        if ppr is not None:
            sp = ppr.find('.//' + qn('w:spacing'))
            if sp is not None: sp.getparent().remove(sp)
        st = doc.styles['Normal']
        set_fonts(st.element.get_or_add_rPr()); st.font.size = Pt(BODY_PT)
        st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.0
        for name, size, lvl in (('Heading 1', 16, 0), ('Heading 2', 12, 1), ('Heading 3', 11, 2), ('Heading 4', 11, None)):
            h = doc.styles[name]
            rPr = h.element.get_or_add_rPr()
            for tag in ('w:color', 'w:i', 'w:iCs', 'w:sz', 'w:szCs', 'w:rFonts'):
                for e in rPr.findall(qn(tag)): rPr.remove(e)
            set_fonts(rPr); h.font.size = Pt(size); h.font.bold = True; h.font.italic = False
            h.font.color.rgb = RGBColor(0, 0, 0)
            pPr = h.element.get_or_add_pPr()
            for e in pPr.findall(qn('w:outlineLvl')): pPr.remove(e)
            if lvl is not None: pPr.append(el('w:outlineLvl', **{'w:val': lvl}))
            h.paragraph_format.keep_with_next = True
        # list title (표 목차, 그림 목차): listed in 목차 like a chapter
        lt = doc.styles.add_style('목록 제목', WD_STYLE_TYPE.PARAGRAPH)
        lt.base_style = doc.styles['Normal']
        set_fonts(lt.element.get_or_add_rPr()); lt.font.size = Pt(16); lt.font.bold = True
        lt.element.get_or_add_pPr().append(el('w:outlineLvl', **{'w:val': 0}))
        # TOC styles
        for name, size, bold, left in (('toc 1', 13, True, 0), ('toc 2', 11, False, 1), ('toc 3', 11, False, 2),
                                       ('toc 4', 11, False, 0)):
            try:
                s = doc.styles[name]
            except KeyError:
                s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = doc.styles['Normal']
            set_fonts(s.element.get_or_add_rPr()); s.font.size = Pt(size); s.font.bold = bold
            pf = s.paragraph_format
            pf.space_before = Pt(6 if name == 'toc 1' else 0); pf.space_after = Pt(0)
            pf.line_spacing = Pt(26 if name == 'toc 1' else 22)
            pPr = s.element.get_or_add_pPr()
            pPr.find(qn('w:spacing')).set(qn('w:lineRule'), 'exact')
            tabs = el('w:tabs'); tabs.append(el('w:tab', **{'w:val': 'right', 'w:leader': 'dot', 'w:pos': TEXT_W_TW}))
            pPr.insert(0, tabs)
            pPr.append(el('w:ind', **{'w:left': int(left * BODY_PT * 20), 'w:right': 0}))

    # ---------------------------------------------------------------- page set-up
    def page(self, sec):
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(15)
        sec.top_margin, sec.bottom_margin = Mm(27), Mm(30)
        sec.header_distance, sec.footer_distance = Mm(16), Mm(10)

    def numbering(self, sec, fmt, start=None):
        sp = sec._sectPr
        for e in sp.findall(qn('w:pgNumType')): sp.remove(e)
        n = el('w:pgNumType', **{'w:fmt': fmt})
        if start is not None: n.set(qn('w:start'), str(start))
        sp.append(n)

    def header(self, sec, text):
        hd = sec.header
        hd.is_linked_to_previous = False
        p = hd.paragraphs[0]
        p.style = self.doc.styles['Normal']          # the template's Header style carries centre and right tabs
        for r in list(p.runs): r._r.getparent().remove(r._r)
        if text:
            self.pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=13.5)
            pPr = p._p.get_or_add_pPr()
            tabs = el('w:tabs'); tabs.append(el('w:tab', **{'w:val': 'right', 'w:pos': TEXT_W_TW})); pPr.insert(0, tabs)
            self.run(p, text + '\t', size=9, font=HEADF)
            self.field(p, 'PAGE', '1', size=9, font=HEADF)
        ft = sec.footer
        ft.is_linked_to_previous = False

    # ---------------------------------------------------------------- tables
    @staticmethod
    def borders(tbl, outer=9, inner=3):
        tblPr = tbl._tbl.tblPr
        for e in tblPr.findall(qn('w:tblBorders')): tblPr.remove(e)
        b = el('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            b.append(el(f'w:{edge}', **{'w:val': 'single' if outer else 'nil', 'w:sz': outer, 'w:space': 0, 'w:color': '000000'}))
        for edge in ('insideH', 'insideV'):
            b.append(el(f'w:{edge}', **{'w:val': 'single' if inner else 'nil', 'w:sz': inner, 'w:space': 0, 'w:color': '000000'}))
        tblPr.append(b)

    @staticmethod
    def fixed(tbl, widths_mm, align=WD_TABLE_ALIGNMENT.CENTER, margin_mm=1.0):
        tblPr = tbl._tbl.tblPr
        tblPr.append(el('w:tblLayout', **{'w:type': 'fixed'}))
        tw = [int(round(w / 25.4 * 1440)) for w in widths_mm]
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is None:
            tblW = el('w:tblW'); tblPr.append(tblW)
        tblW.set(qn('w:w'), str(sum(tw))); tblW.set(qn('w:type'), 'dxa')
        tbl.alignment = align
        for gc, w in zip(tbl._tbl.tblGrid.findall(qn('w:gridCol')), tw):
            gc.set(qn('w:w'), str(w))
        for row in tbl.rows:
            for c, w in zip(row.cells, tw):
                c.width = Emu(int(w * 635))
        mar = el('w:tblCellMar')
        m = int(round(margin_mm / 25.4 * 1440))
        for side in ('top', 'left', 'bottom', 'right'):
            mar.append(el(f'w:{side}', **{'w:w': m if side in ('left', 'right') else 0, 'w:type': 'dxa'}))
        tblPr.append(mar)

    @staticmethod
    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': fill}))

    @staticmethod
    def valign(cell, v='center'):
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(el('w:vAlign', **{'w:val': v}))

    def cell_text(self, cell, lines, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, gray=None, font=GOTHIC):
        """lines: list of segment lists (one paragraph each)."""
        p = cell.paragraphs[0]
        for i, segs in enumerate(lines):
            q = p if i == 0 else cell.add_paragraph()
            self.pfmt(q, align=align, line=size * 1.6)
            for t, k in segs:
                if k == 'b':
                    self.run(q, t, size=size, bold=True, font=font, scale=95)
                else:
                    self.run(q, t, size=size, bold=bold, font=font, scale=95, mark=k)
        if gray: self.shade(cell, gray)
        self.valign(cell)

    def fill_or_blank(self):
        return [('(작성 필요)', 'fill')] if self.marked else [('', None)]

    # ---------------------------------------------------------------- 개요표, 결과물 사진, 신규고용
    def overview(self):
        doc = self.doc
        G = 'D8D8D8'
        widths = [19.7, 28.7, 50.4, 1.1, 22.6, 1.0, 46.5]
        tbl = doc.add_table(rows=19, cols=7)
        self.borders(tbl, 9, 3); self.fixed(tbl, widths, margin_mm=0.8)
        heights = [8.5, 8.5, 8.5, 8.5, 8.5, 8.5, 10.7, 10.7, 8.5, 8.5, 8.5, 8.5, 8.5, 22.0, 10.7, 10.7, 8.5, 22.0, 8.5]
        for row, h in zip(tbl.rows, heights):
            row.height = Mm(h); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        C = lambda r, c: tbl.cell(r, c)
        M = lambda r1, c1, r2, c2: C(r1, c1).merge(C(r2, c2))
        H = lambda cell, *lines: self.cell_text(cell, [[(l, None)] for l in lines], bold=True, gray=G)
        V = lambda cell, segs, align=WD_ALIGN_PARAGRAPH.CENTER: self.cell_text(cell, [segs], align=align)
        blank = self.fill_or_blank
        title = [(t, k) for t, k in TITLE_KO]
        H(M(0, 0, 0, 1), '프로젝트명'); V(M(0, 2, 0, 6), title)
        H(M(1, 0, 6, 0), '컨소시엄', '기업')
        H(M(1, 1, 1, 6), '학생 소속기업   /   컨소시엄(패밀리 기업)      (중복체크 가능)')
        for r, lab in ((2, '기업명'), (3, '주소'), (4, '대표자'), (5, '홈페이지'), (6, '주요 생산품')):
            H(C(r, 1), lab)
        H(M(2, 3, 2, 4), '설립일'); H(M(4, 3, 4, 4), '연락처'); H(M(5, 3, 5, 4), 'FAX'); H(M(6, 3, 6, 4), '매출현황', '(3년치)')
        V(C(2, 2), blank()); V(M(2, 5, 2, 6), blank()); V(M(3, 2, 3, 6), blank())
        V(C(4, 2), blank()); V(M(4, 5, 4, 6), blank()); V(C(5, 2), blank()); V(M(5, 5, 5, 6), blank())
        V(C(6, 2), blank()); V(M(6, 5, 6, 6), blank())
        H(M(7, 0, 15, 0), '산학', '프로젝트', '개요')
        H(C(7, 1), '책임교수'); V(C(7, 2), [(ADVISOR, 'fill'), (' / 기술경영전문대학원 교수', None)])
        H(M(7, 3, 7, 4), '연락처'); V(M(7, 5, 7, 6), blank())
        H(C(8, 1), '실무담당자'); V(C(8, 2), blank()); H(M(8, 3, 8, 4), 'e-mail'); V(M(8, 5, 8, 6), blank())
        H(C(9, 1), '참여학생명'); V(M(9, 2, 9, 6), [(NAME + ' / ', None), (STUDENT_ID, 'fill'), (' / 석사과정', None)])
        H(C(10, 1), '참여학생의 역할'); V(M(10, 2, 10, 6), [(ROLE, 'fill')])
        H(C(11, 1), '전체 참석자'); V(M(11, 2, 11, 6), blank())
        H(C(12, 1), '내부 참석자'); V(M(12, 2, 12, 3), blank()); H(M(12, 4, 12, 5), '외부 참석자'); V(C(12, 6), blank())
        H(C(13, 1), '구분')
        c = M(13, 2, 13, 6)
        self.cell_text(c, [[('직접 / 간접', 'fill')]])
        for line in ('* 직접 : 우리사업의 예산을 통해 지원·참여·운영되고 성과가 발생하여 사사표기가 가능한 경우',
                     '   간접 : 우리사업의 수혜인원이 참여한 타사업과의 연계 성과'):
            q = c.add_paragraph(); self.pfmt(q, align=WD_ALIGN_PARAGRAPH.LEFT, line=15, left=3, hanging=0)
            self.run(q, line, size=9, font=GOTHIC, scale=95)
        H(C(14, 1), '프로젝트 투입비용', '(정부출연금, 천원)')
        V(M(14, 2, 14, 6), [('정부출연금 :          / 민간부담금 : 현금          현물          / 합계 :          ', 'fill')],
          align=WD_ALIGN_PARAGRAPH.LEFT)
        H(C(15, 1), '프로젝트', '총수행기간'); V(M(15, 2, 15, 6), blank())
        H(M(16, 0, 16, 1), '프로젝트 주제'); V(M(16, 2, 16, 6), [(TOPIC, 'fill')])
        H(M(17, 0, 17, 1), '프로젝트 내용')
        c = M(17, 2, 17, 6)
        self.cell_text(c, [[(lab + ': ', 'b'), (txt, 'fill')] for lab, txt in CONTENT])
        H(M(18, 0, 18, 1), '산학프로젝트 결과물'); V(M(18, 2, 18, 6), [('보고서 1부', None)])
        # 결과물 사진 box (left empty; the student may add photos)
        sp = self.para(line=1, after=0); self.run(sp, '', size=1)
        self.pfmt(sp, line=11, after=0)
        box = doc.add_table(rows=1, cols=1)
        self.borders(box, 6, 0); self.fixed(box, [170])
        box.rows[0].height = Mm(25); box.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        p = box.cell(0, 0).paragraphs[0]; self.pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=16)
        if self.marked:
            self.run(p, '프로젝트 결과물 사진 등(해당 시)', size=10, font=GOTHIC, mark='fill')
        self.valign(box.cell(0, 0))
        # 신규고용 창출
        p = self.para(page_break=True, align=WD_ALIGN_PARAGRAPH.LEFT, line=17.6, after=6)
        self.run(p, '□ 신규고용 창출(※ 프로젝트 연계 취업 기술)', size=11, font=GOTHIC)
        t2 = doc.add_table(rows=3, cols=4)
        self.borders(t2, 9, 3); self.fixed(t2, [48.4, 40.5, 40.5, 40.5])
        for i, lab in enumerate(('순번', '성명', '입사연월일', '비고')):
            self.cell_text(t2.cell(0, i), [[(lab, None)]], bold=True, gray='E5E5E5')
        for r in (1, 2):
            for i in range(4):
                self.cell_text(t2.cell(r, i), [[('', None)]])
        for row, h in zip(t2.rows, (6.9, 8.9, 8.9)):
            row.height = Mm(h); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    # ---------------------------------------------------------------- 표지, 인준지
    def rule(self, page_break=False):
        p = self.para(line=2, after=0, page_break=page_break)
        pPr = p._p.get_or_add_pPr()
        b = el('w:pBdr'); b.append(el('w:bottom', **{'w:val': 'single', 'w:sz': 6, 'w:space': 1, 'w:color': '000000'}))
        pPr.append(b)
        self.run(p, '', size=1)

    def titles(self, before=12):
        p = self.para(align=WD_ALIGN_PARAGRAPH.LEFT, line=40, before=before)
        self.runs(p, TITLE_KO, size=25, bold=True)
        p = self.para(align=WD_ALIGN_PARAGRAPH.LEFT, line=40, before=6)
        self.runs(p, TITLE_EN, size=25, bold=True)

    def cover(self):
        self.rule(); self.titles()
        lines = [([(SCHOOL, None)], True), ([('석사학위 캡스톤 프로젝트 보고서', None)], False), ([('', None)], False),
                 ([(DATE, 'fill')], True)]
        for i, (segs, bold) in enumerate(lines):
            p = self.para(align=WD_ALIGN_PARAGRAPH.LEFT, line=20.8, before=48 if i == 0 else 0)
            self.runs(p, segs, size=13, bold=bold, spacing=0.65)
        p = self.para(line=10, before=0); self.run(p, '', size=2)
        sig = self.doc.add_table(rows=4, cols=2)
        self.borders(sig, 0, 0); self.fixed(sig, [44, 24], align=WD_TABLE_ALIGNMENT.RIGHT, margin_mm=0)
        rows = [([(NAME_SP, None)], True, '(서명/인)'), ([(SCHOOL, None)], False, None), ([('지도교수', None)], True, None),
                ([(ADVISOR_SP, 'fill')], False, '(서명/인)')]
        for (segs, bold, tail), row, h in zip(rows, sig.rows, (6.5, 11.0, 6.5, 11.0)):
            row.height = Mm(h); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            cell = row.cells[0] if tail else row.cells[0].merge(row.cells[1])
            for q in cell.paragraphs[1:]: q._p.getparent().remove(q._p)
            p = cell.paragraphs[0]; self.pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=17.6, before=3, after=3)
            self.runs(p, segs, size=11, bold=bold)
            if tail:
                q = row.cells[1].paragraphs[0]; self.pfmt(q, align=WD_ALIGN_PARAGRAPH.RIGHT, line=17.6, before=3, after=3)
                self.run(q, tail, size=10, font=GOTHIC)

    def approval(self):
        self.rule(page_break=True); self.titles()
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, line=20.8, before=36)
        self.run(p, '이 보고서를 %s의 석사학위 보고서로 인준함.' % NAME, size=13, bold=True, spacing=0.65)
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, line=20.8, before=26)
        self.run(p, DATE, size=13, bold=True, spacing=0.65, mark='fill')
        p = self.para(line=14, before=0); self.run(p, '', size=2)
        tbl = self.doc.add_table(rows=3, cols=3)
        self.borders(tbl, 0, 0); self.fixed(tbl, [62, 34, 18], align=WD_TABLE_ALIGNMENT.RIGHT, margin_mm=0)
        who = [('지도교수', ADVISOR_SP, 'fill'), ('기술경영학과 학과장', '(성명)' if self.marked else '', 'fill'),
               ('기술경영전문대학원 원장', '(성명)' if self.marked else '', 'fill')]
        for (role, name, k), row in zip(who, tbl.rows):
            row.height = Mm(15); row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            for i, (txt, al, mk) in enumerate(((role, WD_ALIGN_PARAGRAPH.RIGHT, None), (name, WD_ALIGN_PARAGRAPH.CENTER, k),
                                               ('(인)', WD_ALIGN_PARAGRAPH.RIGHT, None))):
                c = row.cells[i]; q = c.paragraphs[0]; self.pfmt(q, align=al, line=20.8)
                self.run(q, txt, size=13, bold=True, spacing=0.65, mark=mk)
                self.valign(c, 'center')
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, line=24, before=30)
        self.run(p, SCHOOL, size=15, bold=True, spacing=0.75)

    # ---------------------------------------------------------------- lists
    def toc_block(self, title, instr, entries, page_break=False, list_title=False):
        if list_title:
            p = self.doc.add_paragraph(style='목록 제목')
            self.pfmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, line=25.6, after=18, page_break=page_break, keep=True)
            # "표 목차" listed as is, shown evenly spaced like the form's "표 목 차"
            for t, k in [(title, None)]:
                self.run(p, t, size=16, bold=True, spacing=6)
        else:
            p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, line=25.6, after=18, page_break=page_break)
            self.run(p, title, size=16, bold=True)
        if not entries:
            entries = [('toc 4', '(항목 없음)', '')]
        for i, (sname, text, page) in enumerate(entries):
            q = self.doc.add_paragraph(style=sname)
            if i == 0:
                for child in (el('w:fldChar', **{'w:fldCharType': 'begin'}),):
                    rr = el('w:r'); rr.append(child); q._p.append(rr)
                it = el('w:instrText'); it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); it.text = f' {instr} '
                rr = el('w:r'); rr.append(it); q._p.append(rr)
                rr = el('w:r'); rr.append(el('w:fldChar', **{'w:fldCharType': 'separate'})); q._p.append(rr)
            q.add_run(text + '\t' + page)
            if i == len(entries) - 1:
                rr = el('w:r'); rr.append(el('w:fldChar', **{'w:fldCharType': 'end'})); q._p.append(rr)

    @staticmethod
    def list_text(label, segs):
        cap = plain(segs)
        cap = re.sub(r'\s+\([^()]*\)$', '', cap)          # the figure's source stays in the caption only
        return f'{label} {cap}'

    def tc(self, p, text, kind):
        rr = el('w:r'); rr.append(el('w:fldChar', **{'w:fldCharType': 'begin'})); p._p.append(rr)
        it = el('w:instrText'); it.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        it.text = ' TC "%s" \\f %s \\l 4 ' % (text.replace('"', "'"), kind)
        rr = el('w:r'); rr.append(it); p._p.append(rr)
        rr = el('w:r'); rr.append(el('w:fldChar', **{'w:fldCharType': 'end'})); p._p.append(rr)

    # ---------------------------------------------------------------- body
    def body(self, blocks):
        first = True
        in_refs = False
        for b in blocks:
            kind = b[0]
            if kind == 'h1':
                txt = plain(b[1])
                p = self.doc.add_paragraph(style='Heading 1')
                self.pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=25.6, after=22, page_break=not first, keep=True)
                self.runs(p, b[1], size=16, bold=True)
                self.heads.append((1, txt, p)); first = False; in_refs = txt == '참고문헌'
            elif kind == 'h2':
                p = self.doc.add_paragraph(style='Heading 2')
                self.pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=19.2, before=12, after=10, left=10, keep=True)
                self.runs(p, b[1], size=12, bold=True)
                self.heads.append((2, plain(b[1]), p))
            elif kind == 'h3':
                p = self.doc.add_paragraph(style='Heading 3')
                self.pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=17.6, before=8, after=4, left=10, keep=True)
                self.runs(p, b[1], size=11, bold=True)
                self.heads.append((3, plain(b[1]), p))
            elif kind == 'h4':
                p = self.doc.add_paragraph(style='Heading 4')
                self.pfmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=17.6, before=6, after=2, left=10, keep=True)
                self.runs(p, b[1], size=11, bold=True)
            elif kind in ('fig', 'tab'):
                _, fn, w, label, cap = b
                w = min(w, 172)
                if kind == 'tab':
                    p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, line=17.6, before=10, after=4, keep=True)
                    self.run(p, label + ' ', size=11); self.runs(p, cap, size=11)
                    self.tc(p, self.list_text(label, cap), 't'); self.caps.append(('t', self.list_text(label, cap), p))
                    q = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
                    q.add_run().add_picture(str(FIG / fn), width=Mm(w))
                else:
                    q = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, keep=True)
                    q.add_run().add_picture(str(FIG / fn), width=Mm(w))
                    p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, line=16, after=12)
                    self.run(p, label + ' ', size=10); self.runs(p, cap, size=10)
                    self.tc(p, self.list_text(label, cap), 'f'); self.caps.append(('f', self.list_text(label, cap), p))
                # pictures taller than the page are scaled down
                for ext in q._p.xpath('.//wp:extent'):
                    cx, cy = int(ext.get('cx')), int(ext.get('cy'))
                    maxh = int(190 * 36000)
                    if cy > maxh:
                        ext.set('cx', str(int(cx * maxh / cy))); ext.set('cy', str(maxh))
                        for a in q._p.xpath('.//a:ext'):
                            a.set('cx', str(int(cx * maxh / cy))); a.set('cy', str(maxh))
            elif kind == 'bullet':
                self.para(b[1], align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=LINE_PT, left=10, hanging=10)
            elif in_refs:
                self.para(b[1], align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=16, after=10, size=10)
            else:
                self.para(b[1], align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=LINE_PT, first_chars=1)

    # ---------------------------------------------------------------- whole document
    def build(self, blocks, pages):
        doc = self.doc
        self.styles()
        self.page(doc.sections[0])
        self.overview()
        doc.add_section(WD_SECTION.NEW_PAGE)
        self.cover(); self.approval()
        doc.add_section(WD_SECTION.NEW_PAGE)
        # lists are written after the body (they need its headings and captions); keep a marker paragraph here
        marker = self.para(line=1); self.run(marker, '', size=1)
        doc.add_section(WD_SECTION.NEW_PAGE)
        self.body(blocks)
        # 목차 / 표 목차 / 그림 목차 before the marker
        n_before = len(doc.element.body) - 1        # new paragraphs go in before the body-level sectPr
        toc = [('toc 1', '표 목차', pages.get('h:표 목차', '?')), ('toc 1', '그림 목차', pages.get('h:그림 목차', '?'))]
        toc += [('toc %d' % lv, t, pages.get('h:' + t, '?')) for lv, t, p in self.heads]
        self.toc_block('목  차', 'TOC \\o "1-3" \\h \\z \\u', toc)
        self.toc_block('표 목차', 'TOC \\h \\z \\f t', [('toc 4', t, pages.get('c:' + t, '?')) for k, t, p in self.caps
                                                    if k == 't'], page_break=True, list_title=True)
        self.toc_block('그림 목차', 'TOC \\h \\z \\f f', [('toc 4', t, pages.get('c:' + t, '?')) for k, t, p in self.caps
                                                     if k == 'f'], page_break=True, list_title=True)
        made = list(doc.element.body)[n_before:-1]
        if any(e.tag == qn('w:sectPr') for e in made) or doc.element.body[-1].tag != qn('w:sectPr'): fail('list move')
        for e in made:
            marker._p.addprevious(e)
        # sections: page set-up, numbering, headers (python-docx keeps the last section's sectPr at the body level, so
        # the sections are taken from the finished document)
        s1, s2, s3, s4 = doc.sections
        for s in doc.sections: self.page(s)
        self.numbering(s2, 'lowerRoman', 1); self.numbering(s3, 'lowerRoman'); self.numbering(s4, 'decimal', 1)
        self.header(s1, None); self.header(s2, None)
        self.header(s3, '캡스톤 프로젝트 보고서'); self.header(s4, '캡스톤 프로젝트 보고서')
        update_fields_on_open(doc)
        fix_order(doc.element.body); fix_order(doc.styles.element)
        for s in doc.sections:
            fix_order(s.header._element); fix_order(s.footer._element)
        doc.core_properties.title = plain(TITLE_KO)
        doc.core_properties.author = NAME
        return doc


# =============================================================== schema order (Word rejects misordered properties)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
ORDER = {
    'pPr': 'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs '
           'suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd '
           'snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment '
           'textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange',
    'rPr': 'rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid '
           'vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em '
           'lang eastAsianLayout specVanish oMath',
    'tblPr': 'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd '
             'tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription',
    'tcPr': 'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark',
    'sectPr': 'headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType '
              'pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings '
              'sectPrChange',
    'tcMar': 'top start left bottom end right', 'tblCellMar': 'top start left bottom end right',
    'tcBorders': 'top start left bottom end right insideH insideV tl2br tr2bl',
    'tblBorders': 'top start left bottom end right insideH insideV',
    'pBdr': 'top left bottom right between bar',
}
ORDER = {k: {n: i for i, n in enumerate(v.split())} for k, v in ORDER.items()}


def fix_order(root):
    for e in root.iter():
        if not isinstance(e.tag, str) or not e.tag.startswith(W):
            continue
        name = e.tag[len(W):]
        rank = ORDER.get(name)
        if rank is None or len(e) < 2:
            continue
        kids = list(e)
        keyed = sorted(kids, key=lambda k: rank.get(k.tag[len(W):] if isinstance(k.tag, str) and k.tag.startswith(W) else '', 999))
        if keyed != kids:
            for k in kids: e.remove(k)
            for k in keyed: e.append(k)


def update_fields_on_open(doc):
    st = doc.settings.element
    if st.find(qn('w:updateFields')) is not None:
        return
    uf = el('w:updateFields', **{'w:val': 'true'})
    later = ('hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids', 'mathPr', 'attachedSchema',
             'themeFontLang', 'clrSchemeMapping', 'doNotIncludeSubdocsInStats', 'doNotAutoCompressPictures', 'forceUpgrade',
             'captions', 'readModeInkLockDown', 'smartTagType', 'schemaLibrary', 'shapeDefaults', 'doNotEmbedSmartTags',
             'decimalSymbol', 'listSeparator')
    for k in st:
        if isinstance(k.tag, str) and k.tag.split('}')[-1] in later:
            k.addprevious(uf); return
    st.append(uf)


# =============================================================== rendering and page numbers
def render(docx_path, outdir):
    env = dict(os.environ)
    conf = outdir / 'fonts.conf'
    conf.write_text('''<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>
<include ignore_missing="yes">/etc/fonts/fonts.conf</include>
<alias binding="same"><family>HY신명조</family><prefer><family>Noto Serif CJK KR</family></prefer></alias>
<alias binding="same"><family>HY중고딕</family><prefer><family>Noto Sans CJK KR</family></prefer></alias>
<alias binding="same"><family>맑은 고딕</family><prefer><family>Noto Sans CJK KR</family></prefer></alias>
</fontconfig>''', encoding='utf-8')
    env['FONTCONFIG_FILE'] = str(conf)
    prof = outdir / 'lo_profile'
    subprocess.run(['soffice', f'-env:UserInstallation=file://{prof}', '--headless', '--convert-to', 'pdf',
                    '--outdir', str(outdir), str(docx_path)], check=True, env=env, capture_output=True, timeout=600)
    return outdir / (docx_path.stem + '.pdf')


def page_map(pdf, heads, caps):
    import pymupdf
    d = pymupdf.open(str(pdf))
    norm = lambda s: re.sub(r'\s+', '', s)
    pages = []
    for pg in d:
        top = [w[4] for w in pg.get_text('words') if w[3] < 25 / 25.4 * 72]
        label = None
        if top and '캡스톤' in ''.join(top) and re.match(r'^([ivxlc]+|\d+)$', top[-1]):
            label = top[-1]
        pages.append((label, norm(pg.get_text())))
    out = {}
    for name in ('표 목차', '그림 목차'):
        hit = [i for i, (lab, flat) in enumerate(pages) if lab and flat.startswith('캡스톤프로젝트보고서' + lab + norm(name))]
        if not hit: fail('list page not found: ' + name)
        out['h:' + name] = pages[hit[0]][0]
    body_start = next(i for i, (lab, flat) in enumerate(pages) if lab == '1')
    k = body_start
    for lv, t, p in heads:
        key = norm(t)
        while k < len(pages) and key not in pages[k][1]:
            k += 1
        if k == len(pages): fail('heading not found in the rendering: ' + t)
        out['h:' + t] = pages[k][0] or '?'
    for kind, t, p in caps:
        key = norm(t)[:18]
        hit = [i for i in range(body_start, len(pages)) if key in pages[i][1]]
        if not hit: fail('caption not found: ' + t)
        out['c:' + t] = pages[hit[0]][0] or '?'
    return out, len(d), [p[0] for p in pages]


def build_one(marked, pages):
    b = Builder(marked)
    doc = b.build(parse(), pages)
    path = OUT_MARKED if marked else OUT
    doc.save(str(path))
    return b, path


def main():
    pages = {}
    if not RENDER:
        build_one(False, pages); build_one(True, pages); print('built without page numbers'); return
    work = pathlib.Path(tempfile.mkdtemp())
    for it in range(4):
        b, path = build_one(False, pages)
        pdf = render(path, work)
        new, n, labels = page_map(pdf, b.heads, b.caps)
        print(f'pass {it + 1}: {n} pages; labels', labels)
        if new == pages:
            break
        pages = new
    build_one(True, pages)
    keep = ROOT / 'preview'
    keep.mkdir(exist_ok=True)
    shutil.copy(pdf, keep / f'capstone_report_v{VER}_preview.pdf')
    pdf2 = render(OUT_MARKED, work)
    shutil.copy(pdf2, keep / f'capstone_report_v{VER}_marked_preview.pdf')
    print('->', OUT, OUT_MARKED, '| previews in', keep)


if __name__ == '__main__':
    main()
